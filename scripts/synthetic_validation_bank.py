"""Cinco DOCX inventados y expectativas definidas sin consultar el extractor."""

from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement


def expected(kind, value, unit, scale, normalized, confidence, usable):
    return dict(tipo=kind, valor_decimal=value, unidad_original=unit, multiplicador=scale,
                valor_normalizado=normalized, confidence=confidence, usable_for_comparison=usable)


EXPECTED = {
    "cifras_correctas": [expected("importe", "125.50", "pesos", "1", "125.50", "high", True),
                         expected("porcentaje", "12.5", "%", "1", "12.5", "high", True)],
    "cifras_ambiguas": [expected("cantidad", "125", None, "1", "125", "low", False),
                        expected("año", "2024", None, "1", "2024", "low", False),
                        expected("otro", None, None, "1", None, "low", False)],
    "tablas": [expected("importe", "25", "pesos", "1", "25", "high", True),
               expected("importe", "10", "miles de pesos", "1000", "10000", "high", True),
               expected("importe", "3", "millones de pesos", "1000000", "3000000", "high", True),
               expected("importe", "4", "miles de pesos", "1000", "4000", "medium", False)],
    "revisiones": [expected("importe", "200", "pesos", "1", "200", "high", True)],
    "no_importes": [expected("año", "2024", None, "1", "2024", "medium", False),
                    expected("fecha", None, None, "1", None, "medium", False),
                    expected("artículo", "33", None, "1", "33", "medium", False),
                    expected("fracción", "4", None, "1", "4", "medium", False),
                    expected("folio", None, None, "1", None, "medium", False),
                    expected("numero_de_auditoria", "1234", None, "1", None, "medium", False),
                    expected("cantidad", "15", "personas", "1", "15", "high", True)],
}


def generate(directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    docs = {key: Document() for key in EXPECTED}
    for text in ("125.50 pesos", "12.5%"):
        docs["cifras_correctas"].add_paragraph(text)
    for text in ("125", "2024", "12-14"):
        docs["cifras_ambiguas"].add_paragraph(text)
    d = docs["tablas"]
    t = d.add_table(rows=2, cols=1)
    t.cell(0, 0).text = "Importe (pesos)"
    t.cell(1, 0).text = "25"
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Importe (miles de pesos)"
    t.cell(0, 1).text = "10"
    d.add_paragraph("Importes en millones de pesos")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "3"
    d.add_paragraph("")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "4"
    d.add_paragraph("Nota: cifras en miles de pesos.")
    p = docs["revisiones"].add_paragraph("Importe: ")
    for tag, text in (("del", "100"), ("ins", "200")):
        change = OxmlElement("w:" + tag)
        run = OxmlElement("w:r")
        value = OxmlElement("w:delText" if tag == "del" else "w:t")
        value.text = text
        run.append(value)
        change.append(run)
        p._p.append(change)
    p.add_run(" pesos")
    for text in ("Ejercicio 2024", "31/12/2024", "artículo 33", "fracción IV", "Folio AB-001", "Auditoría 001234", "15 personas"):
        docs["no_importes"].add_paragraph(text)
    paths = {}
    for key, document in docs.items():
        paths[key] = directory / (key + ".docx")
        document.save(paths[key])
    return paths


if __name__ == "__main__":
    paths = generate(Path(__file__).resolve().parents[1] / "data" / "synthetic_validation_bank")
    print({"docx_sinteticos_generados": len(paths)})
