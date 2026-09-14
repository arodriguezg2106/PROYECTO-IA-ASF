"""Paquete de 30 decisiones humanas, con selección reproducible sin reemplazo."""

import hashlib
import json
import random
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from workbook_helpers import reduced_context, write_text, style_header

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "results" / "inventario_cifras" / "final"
OUTPUT = ROOT / "results" / "validacion_humana_reducida"
SEED = 20260909
QUOTAS = {"tablas": 5, "high": 10, "ambiguas": 10, "no_clasificadas": 5}
COLUMNS = ["documento_id", "registro_id", "estrato", "elemento_id", "tipo_elemento", "tabla", "fila", "columna",
           "numero_original", "clasificacion_automatica", "tipo_candidato", "unidad_automatica", "multiplicador",
           "valor_normalizado", "origen_unidad", "confidence", "usable_for_arithmetic", "contexto_reducido",
           "clasificacion_correcta", "unidad_correcta", "valor_correcto", "puede_compararse", "observaciones"]


def select_sample(records):
    rng = random.Random(SEED)
    remaining = sorted(records, key=lambda r: (r["documento_id"], r["id"]))
    rng.shuffle(remaining)
    selected, covered = [], set()
    filters = {"tablas": lambda r: bool(r["tabla"]), "high": lambda r: r["confidence"] == "high",
               "ambiguas": lambda r: r["classification"] == "ambiguous" and not r["no_clasificada"],
               "no_clasificadas": lambda r: r["no_clasificada"]}
    for group, quota in QUOTAS.items():
        for _ in range(quota):
            candidates = [i for i, r in enumerate(remaining) if filters[group](r)]
            if not candidates:
                raise ValueError("Registros insuficientes para las cuotas de 30.")
            i = max(candidates, key=lambda i: remaining[i]["documento_id"] not in covered)
            record = remaining.pop(i)
            covered.add(record["documento_id"])
            selected.append({**record, "estrato": group})
    rng.shuffle(selected)
    return selected


def main():
    records, hashes = [], {}
    for path in sorted(SOURCE.glob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        if "cifras" not in data:
            continue
        if not data.get("included_in_aggregates", True):
            continue
        assert data["reading_mode"] == "final"
        hashes[path.name] = hashlib.sha256(path.read_bytes()).hexdigest()
        records.extend({**n, "documento_id": data["source"]["file"]} for n in data["cifras"])
    sample = select_sample(records)
    assert Counter(r["estrato"] for r in sample) == QUOTAS
    assert len({(r["documento_id"], r["id"]) for r in sample}) == 30
    assert sample == select_sample(records)
    wb = Workbook()
    ws = wb.active
    ws.title = "Muestra"
    ws.append(COLUMNS)
    for i, n in enumerate(sample, 2):
        values = [n["documento_id"], n["id"], n["estrato"], n["elemento_id"], "tabla" if n["tabla"] else "párrafo",
                  n["tabla"], n["fila"], n["columna"], n["texto_original"], n["classification"], n["tipo"],
                  n["unidad_original"], n["multiplicador"], n["valor_normalizado"], n["unit_source"], n["confidence"],
                  "sí" if n["usable_for_arithmetic"] else "no", reduced_context(n)]
        for j, value in enumerate(values, 1):
            write_text(ws, i, j, value)
        for j in range(1, 24):
            ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
            if j >= 19:
                ws.cell(i, j).fill = PatternFill("solid", fgColor="E2F0D9")
                ws.cell(i, j).number_format = "@"
        ws.row_dimensions[i].height = 58
    for col in range(1, 24):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 24 if col != 18 else 48
    style_header(ws, 1, 23)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = "A1:W31"
    for col in ("S", "T", "U", "V"):
        validation = DataValidation(type="list", formula1='"sí,no"', allow_blank=True)
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        validation.error = "Selecciona sí o no."
        ws.add_data_validation(validation)
        validation.add(f"{col}2:{col}31")
    summary = wb.create_sheet("Resumen")
    summary.append(["Dimensión", "Categoría", "Registros"])
    for field in ("estrato", "documento_id", "confidence", "tipo", "unit_source"):
        for key, count in sorted(Counter(n[field] for n in sample).items()):
            summary.append([field, key, count])
    style_header(summary, 1, 3)
    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 70
    summary.column_dimensions["C"].width = 14
    instructions = wb.create_sheet("Instrucciones")
    instructions.append(["Tema", "Instrucción"])
    guide = [
        ("Trabajo manual", "Sólo 30 registros; completar las cinco columnas verdes S:W."),
        ("Selección", "Semilla 20260909. Se reservan primero 5 registros de tablas; después 10 high, 10 ambiguos clasificados y 5 no clasificados, sin repetir registros. Las cuotas son grupos de selección, no totales globales por propiedad."),
        ("Solapamiento de propiedades", "Los cinco del grupo tablas pueden ser high, ambiguos o no clasificados; otros grupos también pueden contener tablas. No hay registros duplicados."),
        ("Clasificación correcta", "Confirma sí/no la clasificación automática; tipo_candidato conserva la hipótesis previa cuando classification=ambiguous. Si falta evidencia, deja la casilla vacía y explica la duda."),
        ("Unidad y valor", "Confirma sí/no unidad_automatica y valor_normalizado. Si no aplica a fecha o identificador, deja vacío y escribe no aplica en observaciones. Porcentajes en puntos porcentuales."),
        ("Puede compararse", "Opinión humana sí/no sobre elegibilidad futura; no ejecuta ninguna comparación ni implica compatibilidad con otra cifra."),
        ("Contexto", "Sólo indicios léxicos y etiquetas abreviadas; no se copian párrafos. Si falta evidencia, deja la respuesta vacía y registra una observación."),
        ("Confianza", "high: magnitud normalizable, tipo claro, unidad explícita o encabezado/título inequívoco. medium: evidencia indirecta o referencia no aritmética. low: ambigüedad, sin unidad o varias unidades posibles."),
        ("Usos diferenciados", "La hoja Elegibilidad distingue aritmética, referencias cruzadas, temporal e identificadores. usable_for_arithmetic puede habilitarse con medium dentro de una columna de Decimal válidos con escala común. Especifica el uso al registrar tu opinión humana."),
        ("Copias documentales", "Se seleccionan representantes únicos; las copias idénticas permanecen en el reporte por documento, pero no en esta muestra."),
        ("Precisión", "Los valores se guardan como texto para conservar ceros y precisión decimal. No conviertas automáticamente los identificadores a números."),
        ("Trazabilidad", "documento_id, registro_id y elemento_id permiten localizar la cifra; el manifiesto registra semilla, hashes e IDs seleccionados."),
    ]
    for row in guide:
        instructions.append(row)
    eligibility = wb.create_sheet("Elegibilidad")
    eligibility.append(["documento_id", "registro_id", "usable_for_arithmetic", "usable_for_cross_reference", "usable_for_temporal_analysis", "usable_for_identifier_matching", "motivo"])
    for i, n in enumerate(sample, 2):
        values = [n["documento_id"], n["id"], *[n[k] for k in ("usable_for_arithmetic", "usable_for_cross_reference", "usable_for_temporal_analysis", "usable_for_identifier_matching")], n["reason"]]
        for j, value in enumerate(values, 1):
            write_text(eligibility, i, j, value)
    style_header(eligibility, 1, 7)
    for j in range(1, 8):
        eligibility.column_dimensions[eligibility.cell(1, j).column_letter].width = 32
    style_header(instructions, 1, 2)
    instructions.column_dimensions["A"].width = 30
    instructions.column_dimensions["B"].width = 110
    for row in instructions.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True)
        instructions.row_dimensions[row[0].row].height = 55
    OUTPUT.mkdir(parents=True, exist_ok=True)
    target = OUTPUT / "validacion_final_30.xlsx"
    wb.save(target)
    check = load_workbook(target)
    assert check["Muestra"].max_row == 31
    assert all(check["Muestra"].cell(i, j).value is None for i in range(2, 32) for j in range(19, 24))
    assert all(c.data_type != "f" for sheet in check for row in sheet for c in row)
    check.close()
    assert all(hashlib.sha256((SOURCE / name).read_bytes()).hexdigest() == digest for name, digest in hashes.items())
    manifest = {"semilla": SEED, "algoritmo": "reserve_tables_then_confidence_v1", "cuotas": QUOTAS,
                "fuentes_sha256": hashes, "xlsx_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
                "seleccion": [{"documento_id": n["documento_id"], "registro_id": n["id"], "estrato": n["estrato"]} for n in sample]}
    (OUTPUT / "manifiesto.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"registros": 30, "cuotas": dict(Counter(n["estrato"] for n in sample)),
                      "documentos": len({n["documento_id"] for n in sample}), "duplicados": 0, "respuestas_humanas_vacias": True}))


if __name__ == "__main__":
    main()
