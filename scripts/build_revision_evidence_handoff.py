"""Extrae evidencia editorial real de comentarios y control de cambios de DOCX."""
from __future__ import annotations

import argparse
import difflib
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from revisor_crf.editorial import verify_protected


W="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS={"w":W}; Q=lambda name:f"{{{W}}}{name}"
CONTENT_PART=re.compile(r"^word/(?:document|footnotes|endnotes|header\d+|footer\d+)\.xml$")
CATEGORIES=("Coherencia o claridad","Cifras o datos","Ortografía o gramática","Formato o cuadro",
            "Criterio técnico o normativo","Instrucción administrativa","Otro","No puedo determinarlo")
OPTIONS={
    "usable":("Sí","No","No puedo determinarlo"),
    "category":CATEGORIES,
    "interpretation":("Sí","No","No aplica","No puedo determinarlo"),
    "improves":("Sí","No","No hay cambio asociado","No puedo determinarlo"),
    "protected":("Sí","No","No cambiaron datos protegidos","No puedo determinarlo"),
}


def _plain(text):
    value=unicodedata.normalize("NFD",str(text or ""))
    return "".join(char for char in value if unicodedata.category(char)!="Mn").casefold()


def _digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _value(node):
    if node.tag in {Q("t"),Q("delText")}:
        return node.text or ""
    if node.tag==Q("tab"):return "\t"
    if node.tag in {Q("br"),Q("cr")}:return "\n"
    return ""


def _revision_state(node):
    tags={ancestor.tag for ancestor in node.iterancestors()}
    if Q("del") in tags or Q("moveFrom") in tags:return "original"
    if Q("ins") in tags or Q("moveTo") in tags:return "final"
    return "both"


def _node_text(node,mode):
    state=_revision_state(node)
    return _value(node) if state=="both" or state==mode else ""


def _paragraph_text(paragraph,mode):
    return "".join(_node_text(node,mode) for node in paragraph.iter()
                   if node.tag in {Q("t"),Q("delText"),Q("tab"),Q("br"),Q("cr")}).strip()


def _category(text):
    value=_plain(text)
    patterns=(
        ("Cifras o datos",("cifra","monto","importe","porcentaje","total","dato","coincid","pesos","decimal")),
        ("Coherencia o claridad",("coher","clar","no se entiende","redaccion","redacción","precisar","explicar","sentido","conclusion","conclusión","antecedente","idea","sujeto","secuencia","congru")),
        ("Ortografía o gramática",("ortograf","tilde","acento","coma","punto","mayus","minus","verbo","gerund","concordancia","gramat")),
        ("Formato o cuadro",("formato","fuente","letra","tamano","tamaño","aline","margen","interline","tabla","cuadro","renglon","renglón")),
        ("Criterio técnico o normativo",("fundamento","articulo","artículo","normativa","ley ","criterio","tecnico","técnico")),
        ("Instrucción administrativa",("atender","pendiente","revisar con","confirmar con","validar con","favor de")),
    )
    return next((name for name,terms in patterns if any(term in value for term in terms)),"Otro")


def _result_number(text):
    match=re.search(r"RESULTADO\s*:?\s*(\d+)",_plain(text).upper())
    return match.group(1) if match else None


def _comment_map(archive):
    if "word/comments.xml" not in archive.namelist():return {}
    root=etree.fromstring(archive.read("word/comments.xml")); result={}
    for comment in root.xpath(".//w:comment",namespaces=NS):
        identifier=comment.get(Q("id")); text="\n".join(
            "".join(p.xpath(".//w:t/text()",namespaces=NS)).strip()
            for p in comment.xpath("./w:p",namespaces=NS)).strip()
        result[identifier]={"comment_id":identifier,"comment":text,
                            "comment_date":comment.get(Q("date"))}
    return result


def _part_evidence(root,part,comments,document_sha,document_name):
    tree=root.getroottree(); paragraph_rows=[]; result=None
    for paragraph in root.xpath(".//w:p",namespaces=NS):
        original=_paragraph_text(paragraph,"original"); final=_paragraph_text(paragraph,"final")
        candidate=_result_number(final or original)
        if candidate:result=candidate
        paragraph_rows.append({"node":paragraph,"path":tree.getpath(paragraph),"original":original,"final":final,
                               "result_number":result,"has_revision":original!=final})
    by_path={row["path"]:row for row in paragraph_rows}
    anchors=defaultdict(lambda:{"original":[],"final":[],"paragraph_paths":[]})
    active=set()
    for node in root.iter():
        if node.tag==Q("commentRangeStart"):
            active.add(node.get(Q("id")))
        if active and node.tag in {Q("t"),Q("delText"),Q("tab"),Q("br"),Q("cr")}:
            paragraph=next((ancestor for ancestor in node.iterancestors() if ancestor.tag==Q("p")),None)
            path=tree.getpath(paragraph) if paragraph is not None else None
            for identifier in active:
                anchors[identifier]["original"].append(_node_text(node,"original"))
                anchors[identifier]["final"].append(_node_text(node,"final"))
                if path and path not in anchors[identifier]["paragraph_paths"]:anchors[identifier]["paragraph_paths"].append(path)
        if node.tag==Q("commentRangeEnd"):
            active.discard(node.get(Q("id")))
        if node.tag==Q("commentReference"):
            identifier=node.get(Q("id")); paragraph=next((a for a in node.iterancestors() if a.tag==Q("p")),None)
            path=tree.getpath(paragraph) if paragraph is not None else None
            if path and path not in anchors[identifier]["paragraph_paths"]:anchors[identifier]["paragraph_paths"].append(path)
    records=[]; commented_paths=set()
    for identifier,anchor in anchors.items():
        if identifier not in comments:continue
        paths=anchor["paragraph_paths"]; commented_paths.update(paths)
        rows=[by_path[path] for path in paths if path in by_path]
        before="\n".join(row["original"] for row in rows if row["original"]).strip()
        after="\n".join(row["final"] for row in rows if row["final"]).strip()
        comment=comments[identifier]["comment"]
        raw=f"{document_sha}|{part}|comment|{identifier}|{'|'.join(paths)}"
        records.append({"id":"REV-"+hashlib.sha256(raw.encode()).hexdigest()[:16],"evidence_type":"comment",
                        "document":document_name,"document_sha256":document_sha,"part":part,"xml_paths":paths,
                        "result_number":next((row["result_number"] for row in rows if row["result_number"]),None),
                        "comment_id":identifier,"comment":comment,"anchor_original":"".join(anchor["original"]).strip(),
                        "anchor_final":"".join(anchor["final"]).strip(),"before":before,"after":after,
                        "has_associated_revision":before!=after,"suggested_category":_category(comment),
                        "protection_check":verify_protected(before,after) if before or after else {"safe":True,"removed":[],"added":[]}})
    for row in paragraph_rows:
        if not row["has_revision"] or row["path"] in commented_paths or not (row["original"] or row["final"]):continue
        raw=f"{document_sha}|{part}|revision|{row['path']}"
        records.append({"id":"REV-"+hashlib.sha256(raw.encode()).hexdigest()[:16],"evidence_type":"revision_only",
                        "document":document_name,"document_sha256":document_sha,"part":part,"xml_paths":[row["path"]],
                        "result_number":row["result_number"],"comment_id":None,"comment":None,
                        "anchor_original":None,"anchor_final":None,"before":row["original"],"after":row["final"],
                        "has_associated_revision":True,"suggested_category":"Otro",
                        "protection_check":verify_protected(row["original"],row["final"])})
    return records


def extract_evidence(path):
    path=Path(path); document_sha=_digest(path); records=[]
    try:
        with ZipFile(path) as archive:
            comments=_comment_map(archive)
            for part in archive.namelist():
                if CONTENT_PART.match(part):
                    root=etree.fromstring(archive.read(part))
                    records.extend(_part_evidence(root,part,comments,document_sha,path.name))
    except (BadZipFile,etree.XMLSyntaxError) as exc:
        raise ValueError(f"DOCX inválido: {path.name}") from exc
    return records


def select_sample(records,maximum=40,comment_target=32):
    comments=[row for row in records if row["evidence_type"]=="comment" and row["comment"].strip()
              and (row["before"].strip() or row["after"].strip() or row["anchor_original"].strip() or row["anchor_final"].strip())]
    revisions=[row for row in records if row["evidence_type"]=="revision_only" and row["before"]!=row["after"]]
    def stable(row):return hashlib.sha256(row["id"].encode()).hexdigest()
    categories=defaultdict(list)
    for row in comments:categories[row["suggested_category"]].append(row)
    for rows in categories.values():rows.sort(key=lambda row:(not row["has_associated_revision"],stable(row)))
    selected=[]; used_documents=set()
    while len(selected)<min(comment_target,maximum):
        progressed=False
        for category in CATEGORIES:
            choice=next((row for row in categories.get(category,[]) if row["document_sha256"] not in used_documents),None)
            if choice:
                categories[category].remove(choice); selected.append(choice); used_documents.add(choice["document_sha256"]); progressed=True
                if len(selected)>=min(comment_target,maximum):break
        if not progressed:break
    remaining_comments=sorted((row for rows in categories.values() for row in rows),key=stable)
    for row in remaining_comments:
        if len(selected)>=min(comment_target,maximum):break
        if row["document_sha256"] not in used_documents:
            selected.append(row); used_documents.add(row["document_sha256"])
    for row in sorted(revisions,key=stable):
        if len(selected)>=maximum:break
        if row["document_sha256"] not in used_documents:
            selected.append(row); used_documents.add(row["document_sha256"])
    for row in [*remaining_comments,*sorted(revisions,key=stable)]:
        if len(selected)>=maximum:break
        if row["id"] not in {item["id"] for item in selected}:selected.append(row)
    return selected


def _diff(before,after,side):
    first=re.findall(r"\s+|[^\s]+",before or ""); second=re.findall(r"\s+|[^\s]+",after or "")
    output=[]
    for operation,a1,a2,b1,b2 in difflib.SequenceMatcher(None,first,second).get_opcodes():
        values=first[a1:a2] if side=="before" else second[b1:b2]
        content=escape("".join(values))
        if operation=="equal":output.append(content)
        elif side=="before" and operation in {"delete","replace"}:output.append(f'<del>{content}</del>')
        elif side=="after" and operation in {"insert","replace"}:output.append(f'<ins>{content}</ins>')
    return "".join(output) or "No disponible"


def _case_html(row,index,total):
    previous=f'caso-{index-1:02d}.html' if index>1 else '../INICIAR_AQUI.html'; following=f'caso-{index+1:02d}.html' if index<total else '../INICIAR_AQUI.html'
    comment=(f'<h2>Comentario del revisor</h2><blockquote>{escape(row["comment"])}</blockquote><p><strong>Fragmento anclado:</strong> {escape(row["anchor_original"] or row["anchor_final"] or "No disponible")}</p>' if row["comment"] else '<h2>Control de cambios sin comentario asociado</h2><p>Determine si la modificación constituye una corrección editorial útil.</p>')
    protected="No" if row["protection_check"]["safe"] else "Sí; requiere validación"
    return f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Evidencia editorial {index}</title><style>body{{font:17px/1.55 system-ui;margin:auto;max-width:980px;padding:28px;background:#f2f5f6;color:#17232d}}main{{background:white;padding:28px;border:1px solid #d6dfe5;border-radius:12px}}.task{{background:#e4f3f5;border-left:5px solid #087187;padding:14px}}blockquote{{font-size:20px;background:#fff6dc;border-left:5px solid #d49b00;margin:0;padding:18px}}.versions{{display:grid;grid-template-columns:1fr 1fr;gap:16px}}.version{{background:#f7f8f9;padding:17px;border-radius:8px;white-space:pre-wrap}}del{{background:#ffd8d8}}ins{{background:#d8f3df;text-decoration:none}}nav{{display:flex;justify-content:space-between;margin-top:24px}}a{{color:#075f73;font-weight:700}}@media(max-width:700px){{.versions{{grid-template-columns:1fr}}}}</style></head><body><main><p>Caso {index} de {total}</p><h1>Comentario y cambio editorial real</h1><p class="task">Confirme en la fila {index+1} del Excel si este caso sirve como referencia y si interpretamos correctamente su categoría.</p>{comment}<h2>Antes y después</h2><div class="versions"><div><h3>Texto anterior</h3><div class="version">{_diff(row["before"],row["after"],"before")}</div></div><div><h3>Texto final propuesto</h3><div class="version">{_diff(row["before"],row["after"],"after")}</div></div></div><h2>Datos para validar</h2><ul><li>Categoría sugerida por el sistema: <strong>{escape(row['suggested_category'])}</strong></li><li>¿Cambian cifras, fechas, folios, siglas o nombres detectables?: <strong>{protected}</strong></li><li>Documento: {escape(row['document'])}</li><li>Resultado: {escape(str(row['result_number'] or 'No identificado'))}</li></ul><nav><a href="{previous}">← Anterior</a><a href="../INICIAR_AQUI.html">Lista</a><a href="{following}">Siguiente →</a></nav></main></body></html>'''


def _index_html(rows):
    items="".join(f'<tr><td>{i}</td><td>{escape(row["document"])}</td><td>{escape(str(row["result_number"] or "—"))}</td><td>{"Comentario" if row["evidence_type"]=="comment" else "Control de cambios"}</td><td>{escape(row["suggested_category"])}</td><td><a href="casos/caso-{i:02d}.html">Abrir</a></td></tr>' for i,row in enumerate(rows,1))
    return f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Validación de cambios editoriales</title><style>body{{font:16px/1.5 system-ui;margin:auto;max-width:1150px;padding:28px;color:#17232d}}.start{{background:#e4f3f5;border-left:5px solid #087187;padding:16px}}table{{width:100%;border-collapse:collapse}}th,td{{padding:10px;border-bottom:1px solid #d6dfe5;text-align:left}}th{{background:#eef2f4}}a{{color:#075f73;font-weight:700}}</style></head><body><h1>Validación de comentarios y control de cambios</h1><div class="start">Cada caso muestra evidencia editorial real. Abra el caso, compare antes/después y complete la misma fila en <code>validacion_cambios_editoriales.xlsx</code>.</div><table><thead><tr><th>Caso</th><th>Documento</th><th>Resultado</th><th>Evidencia</th><th>Categoría sugerida</th><th></th></tr></thead><tbody>{items}</tbody></table></body></html>'''


def _validation(sheet,column,options,last_row):
    validation=DataValidation(type="list",formula1='"'+','.join(options)+'"',allow_blank=False)
    validation.error="Seleccione una opción."; validation.errorTitle="Valor no permitido"
    sheet.add_data_validation(validation); letter=sheet.cell(1,column).column_letter; validation.add(f"{letter}2:{letter}{last_row}")


def _workbook(rows,path):
    wb=Workbook(); ws=wb.active; ws.title="Casos"
    headers=("Caso","Abrir","Documento","Resultado","Tipo de evidencia","Comentario","Texto anterior","Texto final",
             "Categoría sugerida","1. ¿Sirve como referencia?","2. Categoría correcta",
             "3. ¿Interpretación correcta?","4. ¿El texto final mejora?","5. ¿Cambio de datos aceptable?",
             "Explicación breve","Nombre del revisor","Fecha")
    ws.append(headers)
    for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78"); cell.alignment=Alignment(wrap_text=True,vertical="top")
    for index,row in enumerate(rows,1):
        ws.append((index,"ABRIR",row["document"],row["result_number"] or "No identificado",row["evidence_type"],row["comment"] or "Sin comentario",row["before"],row["after"],row["suggested_category"],"","","","","","","",""))
        target=ws.max_row; ws.cell(target,2).hyperlink=f"casos/caso-{index:02d}.html"; ws.cell(target,2).style="Hyperlink"
        for column in range(10,18):ws.cell(target,column).fill=PatternFill("solid",fgColor="FFF2CC")
        ws.row_dimensions[target].height=100
    for column,key in zip(range(10,15),("usable","category","interpretation","improves","protected")):_validation(ws,column,OPTIONS[key],len(rows)+1)
    widths=(8,12,40,12,20,50,65,65,26,24,28,28,26,30,48,24,16)
    for column,width in enumerate(widths,1):ws.column_dimensions[ws.cell(1,column).column_letter].width=width
    ws.freeze_panes="J2"; ws.auto_filter.ref=f"A1:Q{len(rows)+1}"
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical="top")
    guide=wb.create_sheet("Instrucciones"); guide.append(("Paso","Qué hacer"))
    guide.append((1,"Abra cada caso mediante ABRIR o INICIAR_AQUI.html.")); guide.append((2,"Lea el comentario y compare el texto anterior con el final."))
    guide.append((3,"Confirme si el ejemplo sirve, la categoría y si el cambio mejora el texto.")); guide.append((4,"Si cambian cifras, fechas, nombres, folios o siglas, indique si el cambio es aceptable."))
    guide.append((5,"Devuelva solamente este Excel completado."))
    for cell in guide[1]:cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78")
    guide.column_dimensions["A"].width=18; guide.column_dimensions["B"].width=105
    for row in guide.iter_rows():
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical="top")
    wb.save(path)


def build(input_root,corpus_root,package_root,maximum=40):
    input_root=Path(input_root).resolve(); corpus_root=Path(corpus_root).resolve(); package_root=Path(package_root).resolve()
    files=sorted(input_root.rglob("*.docx"),key=lambda path:path.as_posix().casefold())
    if not files:raise ValueError(f"No se encontraron DOCX en {input_root}")
    records=[]; failures=[]
    for path in files:
        try:records.extend(extract_evidence(path))
        except (OSError,ValueError) as exc:failures.append({"file":path.name,"error":str(exc)[:300]})
    corpus_root.mkdir(parents=True,exist_ok=True)
    summary={"documents":len(files),"records":len(records),"comments":sum(x["evidence_type"]=="comment" for x in records),
             "revision_only":sum(x["evidence_type"]=="revision_only" for x in records),
             "comments_with_revision":sum(x["evidence_type"]=="comment" and x["has_associated_revision"] for x in records),
             "by_suggested_category":dict(Counter(x["suggested_category"] for x in records if x["evidence_type"]=="comment")),"failures":failures}
    corpus={"schema_version":"editorial-revision-evidence-1.0","generated_at":datetime.now(timezone.utc).isoformat(),
            "contains_raw_institutional_text":True,"intended_use":"local_evaluation_not_training","summary":summary,"records":records}
    (corpus_root/"corpus.json").write_text(json.dumps(corpus,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    sample=select_sample(records,maximum); package_root.mkdir(parents=True,exist_ok=True); cases=package_root/"casos"; cases.mkdir(exist_ok=True)
    created=[]
    for index,row in enumerate(sample,1):
        target=cases/f"caso-{index:02d}.html"; target.write_text(_case_html(row,index,len(sample)),encoding="utf-8"); created.append(target)
    index=package_root/"INICIAR_AQUI.html"; index.write_text(_index_html(sample),encoding="utf-8")
    workbook=package_root/"validacion_cambios_editoriales.xlsx"; _workbook(sample,workbook)
    instructions=package_root/"INSTRUCCIONES.txt"; instructions.write_text("Abra INICIAR_AQUI.html, revise cada caso y complete validacion_cambios_editoriales.xlsx. Devuelva solamente el Excel.\n",encoding="utf-8")
    manifest={"schema_version":"editorial-revision-handoff-1.0","generated_at":datetime.now(timezone.utc).isoformat(),
              "source_corpus":str(corpus_root/"corpus.json"),"cases":len(sample),"contains_original_docx":False,
              "sample_summary":{"comments":sum(x["evidence_type"]=="comment" for x in sample),"revision_only":sum(x["evidence_type"]=="revision_only" for x in sample),
                                "comments_with_revision":sum(x["evidence_type"]=="comment" and x["has_associated_revision"] for x in sample),
                                "protected_elements_changed":sum(not x["protection_check"]["safe"] for x in sample),
                                "documents":len({x["document_sha256"] for x in sample}),"by_category":dict(Counter(x["suggested_category"] for x in sample))}}
    outputs=[index,workbook,instructions,*created]; manifest["files"]=[{"path":p.relative_to(package_root).as_posix(),"sha256":_digest(p),"bytes":p.stat().st_size} for p in outputs]
    (package_root/"MANIFEST.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return corpus,manifest


def main(argv=None):
    parser=argparse.ArgumentParser(description="Extrae comentarios/revisiones y prepara validación humana.")
    parser.add_argument("input"); parser.add_argument("--corpus",default="results/editorial_revision_corpus")
    parser.add_argument("--output",default="results/paquete_validacion_cambios_editoriales"); parser.add_argument("--cases",type=int,default=40)
    args=parser.parse_args(argv)
    try:corpus,manifest=build(args.input,args.corpus,args.output,args.cases)
    except (OSError,ValueError) as exc:print(f"Error: {exc}"); return 1
    summary=corpus["summary"]; print(f"Corpus: {summary['comments']} comentarios y {summary['revision_only']} pares sin comentario. Paquete: {manifest['cases']} casos.")
    return 0


if __name__=="__main__":raise SystemExit(main())
