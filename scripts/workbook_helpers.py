"""Formato XLSX y contexto abreviado sin copiar p?rrafos."""

import re
from openpyxl.styles import Alignment, Font, PatternFill


def reduced_context(record):
    # Word-level evidence only: no contiguous sentence or full paragraph copied.
    text = record.get("contexto_inmediato", "")
    patterns = [r"auditor[ií]a", r"folio", r"expediente", r"contrato", r"art[ií]culo", r"fracci[oó]n",
                r"ejercicio", r"año", r"fecha", r"importe", r"monto", r"presupuesto", r"pagad[oa]s?",
                r"ejercid[oa]s?", r"ministrad[oa]s?", r"porcentaje", r"personas", r"obras",
                r"municipios", r"beneficiarios", r"cantidad", r"total"]
    words = []
    for pattern in patterns:
        found = re.search(r"\b" + pattern + r"\b", text, re.I)
        if found:
            words.append(found.group().casefold())
    hints = ["indicios=[" + " | ".join(words[:5]) + "]"] if words else []
    for field, label in (("encabezado_columna", "columna"), ("encabezado_fila", "fila")):
        tokens = re.findall(r"[^\W\d_]+", record.get(field) or "", re.UNICODE)
        if tokens:
            hints.append(label + "=[" + " | ".join(tokens[:3]) + "]")
    return "; ".join(hints)[:140] or "Sin indicios textuales adicionales en el contexto disponible."



def write_text(ws, row, col, value):
    if value is None:
        return
    target = ws.cell(row, col, str(value))
    target.data_type = "s"  # Preserve leading zeros and prevent formula execution.
    target.number_format = "@"


def style_header(ws, row, width):
    for cell in ws[row][:width]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 34

