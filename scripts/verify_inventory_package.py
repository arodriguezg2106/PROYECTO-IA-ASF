"""Verifica integridad del paquete local sin modificar la muestra humana."""

import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from human_validation_small import SEED, QUOTAS, select_sample
from revisor_crf.units import signature

ROOT = Path(__file__).resolve().parents[1]


def check(condition, message):
    if not condition:
        raise ValueError(message)


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify(root=ROOT):
    folder = root / "results/inventario_cifras/final"
    sample_folder = root / "results/validacion_humana_reducida"
    coverage = read(folder / "cobertura_confianza.json")
    records, source_hashes = [], {}
    for summary in coverage["documentos"]:
        original = root / "tests/CRF-sinteticas" / summary["documento"]
        check(digest(original) == summary["sha256_original"], "Cambió el hash de un documento original.")
        check(summary["cobertura_xml_texto"] == 100 and not summary["errores_tecnicos"]
              and not summary["duplicados"] and not summary["errores_orden_origen"], "Verificación XML incompleta.")
        path = folder / (original.stem + ".json")
        document = read(path)
        check(document["reading_mode"] == "final", "El inventario no corresponde al modo final.")
        check(len(document["cifras"]) == summary["detectadas"], "Conteo por documento desactualizado.")
        check(document["included_in_aggregates"] == summary["included_in_aggregates"], "Deduplicación incoherente.")
        if document["included_in_aggregates"]:
            source_hashes[path.name] = digest(path)
            records.extend({**n, "documento_id": original.name} for n in document["cifras"])
    check(len(records) == coverage["total_cifras"], "Conteo agregado desactualizado.")
    check(dict(Counter(n["confidence"] for n in records)) == coverage["confidence"], "Confianza agregada desactualizada.")
    check(dict(Counter(n["unit_source"] for n in records)) == coverage["por_origen_unidad"], "Orígenes agregados desactualizados.")
    for n in records:
        if n["confidence"] == "high":
            candidates = [e for e in n["unit_candidates"] if e.get("applicable", True)]
            check(n["unit_unambiguous"] and len({signature(e["unit"]) for e in candidates}) == 1,
                  "Registro high sin unidad inequívoca.")
            check(n["valor_normalizado"] is not None and Decimal(n["valor_normalizado"]).is_finite(),
                  "Registro high sin Decimal normalizado válido.")
        if n["tipo"] in {"artículo", "fracción", "numero_de_auditoria", "folio", "año", "fecha"}:
            check(not n["usable_for_arithmetic"], "Referencia marcada como aritmética.")
    crosses = read(folder / "tablas_cruzadas.json")["tablas"]
    for rows in crosses.values():
        check(sum(r["registros"] for r in rows) == len(records), "Tabla cruzada desactualizada.")
    manifest = read(sample_folder / "manifiesto.json")
    check(manifest["semilla"] == SEED and manifest["cuotas"] == QUOTAS, "Configuración de muestra desactualizada.")
    check(manifest["fuentes_sha256"] == source_hashes, "La muestra no corresponde al inventario vigente.")
    selection = [{"documento_id": n["documento_id"], "registro_id": n["id"], "estrato": n["estrato"]}
                 for n in select_sample(records)]
    check(selection == manifest["seleccion"], "La selección no se reproduce con la semilla registrada.")
    xlsx = sample_folder / "validacion_final_30.xlsx"
    with xlsx.open("rb") as stream:
        wb = load_workbook(stream)
        try:
            sheet = wb["Muestra"]
            actual = [{"documento_id": row[0], "registro_id": row[1], "estrato": row[2]}
                      for row in sheet.iter_rows(min_row=2, values_only=True)]
            check(actual == selection, "Las filas de la muestra difieren del manifiesto.")
            answered = sum(sheet.cell(i, j).value is not None for i in range(2, 32) for j in range(19, 24))
        finally:
            wb.close()
    return {"fecha_utc": datetime.now(timezone.utc).isoformat(), "cifras": len(records),
            "documentos_verificados": len(coverage["documentos"]), "documentos_unicos": len(source_hashes),
            "copias_excluidas": coverage["copias_excluidas"], "cobertura_xml_reportada_pct": 100,
            "originales_intactos": True, "tablas_cruzadas": len(crosses), "muestra_reproducible": len(selection),
            "xlsx_sin_cambios_desde_generacion": digest(xlsx) == manifest["xlsx_sha256"],
            "campos_humanos_completados": answered,
            "alcance": "Integridad de artefactos y consistencia de metadatos; no valida las cifras ni repite la extracción XML."}


if __name__ == "__main__":
    report = verify()
    target = ROOT / "results/auditoria_focalizada/cierre_integridad.json"
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True))
