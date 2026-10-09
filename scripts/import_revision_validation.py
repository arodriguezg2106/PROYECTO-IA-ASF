"""Importa una validación humana de comentarios y control de cambios.

El resultado contiene texto institucional y debe permanecer en ``results/``.
Las columnas de evidencia se contrastan con el paquete original; cualquier
edición se registra, pero no reemplaza la evidencia canónica del corpus.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from build_revision_evidence_handoff import OPTIONS, select_sample


HEADERS = (
    "Caso", "Abrir", "Documento", "Resultado", "Tipo de evidencia", "Comentario",
    "Texto anterior", "Texto final", "Categoría sugerida",
    "1. ¿Sirve como referencia?", "2. Categoría correcta",
    "3. ¿Interpretación correcta?", "4. ¿El texto final mejora?",
    "5. ¿Cambio de datos aceptable?", "Explicación breve", "Nombre del revisor", "Fecha",
)
DECISION_KEYS = ("usable", "category", "interpretation", "improves", "protected")


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _text_digest(value) -> str:
    return hashlib.sha256(str(value or "").encode("utf-8")).hexdigest()


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _same_cell(first, second):
    """Excel representa de igual forma una cadena vacía y una celda vacía."""
    if first in (None, "") and second in (None, ""):
        return True
    return first == second


def _canonical_row(index, record):
    return (
        index,
        "ABRIR",
        record["document"],
        record.get("result_number") or "No identificado",
        record["evidence_type"],
        record.get("comment") or "Sin comentario",
        record.get("before"),
        record.get("after"),
        record["suggested_category"],
    )


def import_validation(submitted_path, corpus_path, output_path, template_path=None):
    submitted_path = Path(submitted_path).resolve()
    corpus_path = Path(corpus_path).resolve()
    output_path = Path(output_path).resolve()
    template_path = Path(template_path).resolve() if template_path else None

    workbook = load_workbook(submitted_path, read_only=True, data_only=True)
    if "Casos" not in workbook.sheetnames:
        raise ValueError("El libro no contiene la hoja Casos.")
    sheet = workbook["Casos"]
    headers = tuple(cell.value for cell in sheet[1])
    if headers != HEADERS:
        raise ValueError("Las columnas no corresponden al paquete de validación esperado.")
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if not rows:
        raise ValueError("El libro no contiene casos para importar.")

    corpus = json.loads(corpus_path.read_text(encoding="utf-8"))
    if corpus.get("schema_version") != "editorial-revision-evidence-1.0":
        raise ValueError("La versión del corpus de evidencia no es compatible.")
    records = select_sample(corpus["records"], len(rows))
    if len(records) != len(rows):
        raise ValueError("El número de filas no coincide con la muestra reconstruida del corpus.")

    template_hash = None
    template_rows = None
    if template_path:
        template = load_workbook(template_path, read_only=True, data_only=True)
        if "Casos" not in template.sheetnames:
            raise ValueError("La plantilla no contiene la hoja Casos.")
        template_headers = tuple(cell.value for cell in template["Casos"][1])
        if template_headers != HEADERS:
            raise ValueError("Las columnas de la plantilla no son compatibles.")
        template_rows = list(template["Casos"].iter_rows(min_row=2, values_only=True))
        if len(template_rows) != len(rows):
            raise ValueError("La plantilla y el libro contestado tienen distinto número de casos.")
        template_hash = _digest(template_path)

    fixed_changes = []
    imported = []
    invalid = []
    for position, (values, record) in enumerate(zip(rows, records), 1):
        canonical = _canonical_row(position, record)
        reference = template_rows[position - 1][:9] if template_rows else canonical
        if template_rows and any(not _same_cell(expected, actual) for expected, actual in zip(canonical, reference)):
            raise ValueError(f"La plantilla no corresponde al corpus reconstruido en el caso {position}.")
        for column, (expected, submitted) in enumerate(zip(reference, values[:9]), 1):
            if not _same_cell(expected, submitted):
                fixed_changes.append({
                    "case": position,
                    "column": HEADERS[column - 1],
                    "canonical_sha256": _text_digest(expected),
                    "submitted_sha256": _text_digest(submitted),
                    "canonical_value": _json_value(expected),
                    "submitted_value": _json_value(submitted),
                })
        decisions = dict(zip(DECISION_KEYS, values[9:14]))
        for key, value in decisions.items():
            if value not in OPTIONS[key]:
                invalid.append({"case": position, "field": key, "value": _json_value(value)})
        eligible = (
            decisions["usable"] == "Sí"
            and decisions["interpretation"] == "Sí"
            and decisions["improves"] == "Sí"
            and decisions["protected"] != "No"
        )
        imported.append({
            "case": position,
            "evidence_id": record["id"],
            "document": record["document"],
            "document_sha256": record["document_sha256"],
            "result_number": record.get("result_number"),
            "evidence_type": record["evidence_type"],
            "canonical_category": record["suggested_category"],
            "human_decision": decisions,
            "explanation": values[14],
            "review_date": _json_value(values[16]),
            "eligible_for_evaluation_reference": eligible,
        })
    if invalid:
        details = ", ".join(f"caso {x['case']} ({x['field']})" for x in invalid)
        raise ValueError(f"Hay decisiones vacías o no permitidas: {details}.")

    category_pairs = Counter(
        (item["canonical_category"], item["human_decision"]["category"])
        for item in imported
    )
    exact_category = sum(
        item["canonical_category"] == item["human_decision"]["category"]
        for item in imported
    )
    summary = {
        "cases": len(imported),
        "usable_yes": sum(x["human_decision"]["usable"] == "Sí" for x in imported),
        "usable_no": sum(x["human_decision"]["usable"] == "No" for x in imported),
        "usable_undetermined": sum(x["human_decision"]["usable"] == "No puedo determinarlo" for x in imported),
        "eligible_evaluation_references": sum(x["eligible_for_evaluation_reference"] for x in imported),
        "category_exact_matches": exact_category,
        "category_exact_agreement": round(exact_category / len(imported), 4),
        "immutable_field_changes": len(fixed_changes),
        "by_human_category": dict(Counter(x["human_decision"]["category"] for x in imported)),
        "category_pairs": [
            {"suggested": pair[0], "human": pair[1], "count": count}
            for pair, count in category_pairs.most_common()
        ],
    }
    payload = {
        "schema_version": "editorial-revision-validation-1.0",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "privacy": {
            "contains_raw_institutional_text": True,
            "reviewer_names_stored": False,
            "intended_use": "local_evaluation_not_automatic_training",
        },
        "source": {
            "submitted_sha256": _digest(submitted_path),
            "corpus_sha256": _digest(corpus_path),
            "template_sha256": template_hash,
        },
        "summary": summary,
        "immutable_field_changes": fixed_changes,
        "cases": imported,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return payload


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", help="Excel contestado")
    parser.add_argument("--corpus", default="results/editorial_revision_corpus/corpus.json")
    parser.add_argument("--template", default="results/paquete_validacion_cambios_editoriales/validacion_cambios_editoriales.xlsx")
    parser.add_argument("--output", default="results/editorial_revision_validation/validated_cases.json")
    args = parser.parse_args(argv)
    try:
        result = import_validation(args.workbook, args.corpus, args.output, args.template)
    except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
        print(f"Error: {exc}")
        return 1
    summary = result["summary"]
    print(
        f"Importados {summary['cases']} casos: "
        f"{summary['eligible_evaluation_references']} referencias elegibles, "
        f"{summary['immutable_field_changes']} cambios en campos fijos."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
