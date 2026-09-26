"""Genera exclusivamente DOCX sintéticos y un manifiesto de expectativas."""
from __future__ import annotations
import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

CASES = [
 {"case_id":"percentage_space","text":"El avance fue de 10.0 %.","expected_findings":[{"rule_id":"ASF-ED-II-07-10","problem_span":"0.0 %"}]},
 {"case_id":"percentage_integer","text":"El avance fue de 10%.","expected_findings":[{"rule_id":"ASF-ED-II-07-12","problem_span":"10%"}]},
 {"case_id":"apostrophe_separator","text":"Se determinó un importe de $3'131,988.77.","expected_findings":[{"rule_id":"ASF-ED-II-07-05","problem_span":"3'1"}]},
 {"case_id":"mixed_amount","text":"Se construyeron 5 mil 277 espacios.","expected_findings":[{"rule_id":"ASF-ED-II-07-06","problem_span":"5 mil 277"}]},
 {"case_id":"fraction","text":"La participación fue de 1/10.","expected_findings":[{"rule_id":"ASF-ED-II-07-08","problem_span":"1/10"}]},
 {"case_id":"acronym_plural","text":"Las UAA's entregaron la información.","expected_findings":[{"rule_id":"ASF-ED-II-08-04C","problem_span":"UAA's"}]},
 {"case_id":"acronym_dotted","text":"La S.H.C.P. emitió el oficio.","expected_findings":[{"rule_id":"ASF-ED-II-08-04A","problem_span":"S.H.C.P."}]},
 {"case_id":"acronym_defined","text":"La Secretaría de Hacienda y Crédito Público (SHCP) informó el resultado.","expected_findings":[]},
 {"case_id":"acronym_undefined","text":"La ASF informó el resultado de la revisión.","expected_findings":[{"rule_id":"ASF-ED-II-08-04B","problem_span":"ASF"}]},
 {"case_id":"unit_plural","text":"La distancia medida fue de 10 kms.","expected_findings":[{"rule_id":"ASF-ED-II-08-03","problem_span":"10 kms."}]},
 {"case_id":"wording","text":"La entidad decidió aperturar una cuenta.","expected_findings":[{"rule_id":"ASF-ED-I-06-06","problem_span":"aperturar"}]},
 {"case_id":"conjunction","text":"Se solicitó el oficio y/o el contrato.","expected_findings":[{"rule_id":"ASF-ED-I-05-YO","problem_span":"y/o"}]},
 {"case_id":"relative","text":"Se firmó el convenio, mismo que entró en vigor.","expected_findings":[{"rule_id":"ASF-ED-I-05-MISMO","problem_span":"mismo que"}]},
 {"case_id":"negative_clean","text":"La entidad abrió una cuenta y entregó el contrato.","expected_findings":[]},
 {"case_id":"negative_lowercase_words","text":"Las entidades recorrieron varios espacios y kilómetros.","expected_findings":[]},
 {"case_id":"gerund_consequence","text":"La entidad omitió los controles, provocando una diferencia.","expected_findings":[{"rule_id":"ASF-ED-I-04-GER","problem_span":", provocando una diferencia"}]},
 {"case_id":"gerund_valid","text":"La persona ingresó caminando lentamente.","expected_findings":[]},
 {"case_id":"omitted_que","text":"Se informó a la entidad debía presentar la documentación.","expected_findings":[{"rule_id":"ASF-ED-I-04-QUE","problem_span":"Se informó a la entidad debía"}]},
 {"case_id":"que_valid","text":"Se informó a la entidad que debía presentar la documentación.","expected_findings":[]},
 {"case_id":"redundancy","text":"La entidad volvió a repetir la solicitud.","expected_findings":[{"rule_id":"ASF-ED-I-04-RED","problem_span":"volvió a repetir"}]},
 {"case_id":"redundancy_valid","text":"La entidad repitió la solicitud.","expected_findings":[]},
 {"case_id":"repetition_three","text":"El programa presentó deficiencias en el programa porque el programa careció de controles.","expected_findings":[{"rule_id":"ASF-ED-I-04-REP","problem_span":"programa"}]},
 {"case_id":"repetition_valid","text":"El programa presentó deficiencias porque careció de controles.","expected_findings":[]},
 {"case_id":"nonconclusive","text":"La auditoría estaría identificando deficiencias.","expected_findings":[{"rule_id":"ASF-ED-I-04-TENSE","problem_span":"estaría identificando"}]},
 {"case_id":"conclusive_valid","text":"La auditoría identificó deficiencias.","expected_findings":[]},
 {"case_id":"subjunctive_conclusion","text":"Se concluyó que la entidad pudiera incumplir.","expected_findings":[{"rule_id":"ASF-ED-I-04-SUBJ","problem_span":"Se concluyó que la entidad pudiera"}]},
 {"case_id":"indicative_valid","text":"Se concluyó que la entidad incumplió.","expected_findings":[]},
 {"case_id":"omitted_se","text":"En el expediente identificaron tres contratos.","expected_findings":[{"rule_id":"ASF-ED-I-04-SE","problem_span":"En el expediente identificaron"}]},
 {"case_id":"se_valid","text":"En el expediente se identificaron tres contratos.","expected_findings":[]},
 {"case_id":"vague_without_antecedent","text":"Lo anterior permitió atender la situación señalada.","expected_findings":[{"rule_id":"ASF-ED-I-04-REF","problem_span":"Lo anterior"}]},
 {"case_id":"specific_reference","text":"La conciliación permitió corregir la diferencia.","expected_findings":[]},
 {"case_id":"percentage_table_heading","text":"10 por ciento","table":True,"expected_findings":[]},
 {"case_id":"format_direct_font","text":"Este párrafo sintético permite comprobar el tamaño directo incorrecto sin información institucional.","format_case":"font","expected_findings":[{"rule_id":"ASF-ED-III-03-02","problem_span":"Este párrafo sintético permite comprobar el tamaño directo incorrecto sin información institucional."}]},
 {"case_id":"format_alignment","text":"Este párrafo sintético permite comprobar una alineación directa distinta de la justificada.","format_case":"alignment","expected_findings":[{"rule_id":"ASF-ED-III-03-04","problem_span":"Este párrafo sintético permite comprobar una alineación directa distinta de la justificada."}]},
 {"case_id":"format_vertical_border","text":"Dato","format_case":"border","expected_findings":[{"rule_id":"ASF-ED-II-09-02","problem_span":"Bordes de cuadro"}]},
]

def _institutional_document():
    doc=Document(); normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(11); normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing=1.5; normal.paragraph_format.space_before=Pt(3); normal.paragraph_format.space_after=Pt(3)
    section=doc.sections[0]; section.top_margin=Cm(2); section.bottom_margin=Cm(3); section.left_margin=Cm(4); section.right_margin=Cm(3); section.gutter=Cm(0); section.header_distance=Cm(.5); section.footer_distance=Cm(2)
    return doc

def generate(directory):
    target=Path(directory); target.mkdir(parents=True,exist_ok=True); manifest=[]
    for spec in CASES:
        doc=_institutional_document()
        if spec.get("table"):
            doc.add_table(rows=1,cols=1).cell(0,0).text=spec["text"]
        elif spec.get("format_case")=="border":
            table=doc.add_table(rows=1,cols=1); table.cell(0,0).text=spec["text"]
            for run in table.cell(0,0).paragraphs[0].runs:run.font.name="Calibri"; run.font.size=Pt(8)
            p=table.cell(0,0).paragraphs[0]; p.paragraph_format.line_spacing=1; p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(0)
            borders=OxmlElement("w:tblBorders"); inside=OxmlElement("w:insideV"); inside.set(qn("w:val"),"single"); borders.append(inside); table._tbl.tblPr.append(borders)
        else:
            p=doc.add_paragraph(spec["text"])
            if spec.get("format_case")=="font":p.runs[0].font.size=Pt(12)
            if spec.get("format_case")=="alignment":p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        path=target/(spec["case_id"]+".docx"); doc.save(path)
        manifest.append({**spec,"check_format":bool(spec.get("format_case")),"document":path.name,"description":spec.get("description",spec["case_id"])})
    payload={"schema_version":"editorial-validation-bank-1.0","synthetic":True,"cases":manifest}
    (target/"manifest.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return payload

if __name__=="__main__":
    payload=generate(Path(__file__).resolve().parents[1]/"data"/"editorial_validation_bank")
    print(f"Banco editorial sintético generado: {len(payload['cases'])} casos")
