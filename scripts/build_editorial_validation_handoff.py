"""Construye un paquete autocontenido para validación humana editorial."""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from revisor_crf.editorial import load_editorial_catalog


DECISIONS=("Sí, es un error","No, es falso positivo","Es una excepción autorizada","No puedo determinarlo")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def instructions_text(case_count):
    return f"""# Validación humana del revisor editorial ASF

## Objetivo

Clasificar una muestra de {case_count} hallazgos producidos por el motor determinista. Esta actividad no entrena modelos, no modifica las CRF y no requiere usar Ollama.

## Cómo trabajar

1. Abra `INICIAR_AQUI.html` para ver la lista de casos.
2. Abra un caso: la parte que debe revisar aparece resaltada en amarillo.
3. Regrese a `validacion_editorial.xlsx`, hoja **Casos**, y elija una respuesta en la columna amarilla **Decisión**.
4. Escriba una observación sólo cuando ayude a explicar la decisión.
5. No corrija ni edite los documentos originales.
6. Al terminar, devuelva únicamente el Excel completado.

## Decisiones permitidas

- **Sí, es un error:** el señalamiento es correcto y útil.
- **No, es falso positivo:** el texto, cifra o formato está bien y el sistema no debió señalarlo.
- **Es una excepción autorizada:** la diferencia existe, pero está permitida institucionalmente.
- **No puedo determinarlo:** falta contexto o es necesario abrir el DOCX original.

## Seguridad

Los HTML pueden contener fragmentos de las CRF. Comparta esta carpeta sólo mediante OneDrive, SharePoint o Teams institucional. No utilice servicios personales o públicos.
"""


def _highlight(text, span):
    text=str(text or ""); span=str(span or "")
    if span and span in text:
        before,after=text.split(span,1)
        return escape(before)+f"<mark>{escape(span)}</mark>"+escape(after)
    return escape(text or "No hay fragmento textual; revise la ubicación y el valor efectivo.")


def _effective_text(value):
    if value is None:return "No aplica"
    return json.dumps(value,ensure_ascii=False,indent=2)


def _render_case(item,rule,index,total):
    finding=item["finding"]
    location=finding.get("location") or {}
    previous=f'caso-{index-1:02d}.html' if index>1 else '../INICIAR_AQUI.html'
    following=f'caso-{index+1:02d}.html' if index<total else '../INICIAR_AQUI.html'
    return f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Caso {index}</title><style>body{{font:17px/1.55 system-ui;margin:auto;max-width:900px;padding:28px;color:#17232d;background:#f3f6f7}}main{{background:white;border:1px solid #d5dfe3;border-radius:12px;padding:26px}}h1{{margin-top:0}}.step{{background:#e8f4f6;border-left:5px solid #087187;padding:14px}}.text{{font-size:20px;background:#f7f8f9;padding:20px;border-radius:8px;overflow-wrap:anywhere}}mark{{background:#ffe36e;padding:2px 4px;border-radius:3px}}.finding{{background:#fff6dc;border:1px solid #ead28a;padding:16px;border-radius:8px}}.choices li{{margin:10px 0}}dt{{font-weight:700;margin-top:12px}}dd{{margin-left:0}}nav{{display:flex;justify-content:space-between;margin-top:22px}}a{{color:#075f73;font-weight:700}}pre{{white-space:pre-wrap;background:#f5f6f7;padding:12px}}</style></head><body><main><p>Caso {index} de {total}</p><h1>¿El sistema señaló correctamente este caso?</h1><p class="step"><strong>Su tarea:</strong> lea el texto y lo detectado. Después seleccione una decisión en la fila {index+1} del Excel.</p><h2>Texto o elemento a revisar</h2><div class="text">{_highlight(finding.get("original_text"),finding.get("problem_span"))}</div><h2>Lo que detectó el sistema</h2><div class="finding"><strong>{escape(str(finding.get("message") or item.get("message") or ""))}</strong><p>Parte señalada: <mark>{escape(str(finding.get("problem_span") or "No disponible"))}</mark></p></div><h2>¿Qué debe decidir?</h2><ul class="choices"><li><strong>Sí, es un error:</strong> el señalamiento es correcto.</li><li><strong>No, es falso positivo:</strong> el contenido está bien.</li><li><strong>Es una excepción autorizada:</strong> la diferencia está permitida.</li><li><strong>No puedo determinarlo:</strong> necesita consultar el DOCX o un criterio adicional.</li></ul><h2>Dónde aparece</h2><dl><dt>Documento</dt><dd>{escape(item['file'])}</dd><dt>Sección</dt><dd>{escape(str(location.get('section') or 'No identificada'))}</dd><dt>Tabla/celda</dt><dd>{escape(str(location.get('cell_reference') or location.get('table_id') or 'No aplica'))}</dd><dt>Referencia interna</dt><dd><code>{escape(str(location.get('paragraph_id') or item.get('paragraph_id') or ''))}</code></dd></dl><details><summary>Ver fundamento y datos técnicos</summary><dl><dt>Regla</dt><dd>{escape(item['rule_id'])}: {escape(str(rule.get('title') or 'Regla externa'))}</dd><dt>Fundamento</dt><dd>{escape(str(rule.get('source_section') or finding.get('source_section') or ''))}, p. {escape(str(rule.get('source_page') or finding.get('source_page') or '—'))}</dd><dt>Valor efectivo</dt><dd><pre>{escape(_effective_text(finding.get('effective_value')))}</pre></dd></dl></details><nav><a href="{previous}">← Anterior</a><a href="../INICIAR_AQUI.html">Lista de casos</a><a href="{following}">Siguiente →</a></nav></main></body></html>'''


def _render_index(items):
    rows="".join(f'<tr><td>{i}</td><td>{escape(item["file"])}</td><td>{escape(str(item["finding"].get("problem_span") or "—"))}</td><td>{escape(str(item["finding"].get("message") or item.get("message") or ""))}</td><td><a href="casos/caso-{i:02d}.html">Abrir caso</a></td></tr>' for i,item in enumerate(items,1))
    return f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Validación editorial</title><style>body{{font:16px/1.45 system-ui;margin:auto;max-width:1200px;padding:28px;color:#17232d}}.start{{background:#e8f4f6;border-left:5px solid #087187;padding:16px}}table{{width:100%;border-collapse:collapse}}th,td{{padding:10px;border-bottom:1px solid #d5dfe3;text-align:left;vertical-align:top}}th{{background:#edf2f4}}a{{color:#075f73;font-weight:700}}</style></head><body><h1>Validación de casos editoriales</h1><div class="start"><strong>Pasos:</strong> abra un caso, revise la parte amarilla y registre su decisión en la misma fila de <code>validacion_editorial.xlsx</code>. No necesita interpretar códigos técnicos.</div><p>{len(items)} casos por revisar.</p><table><thead><tr><th>Caso</th><th>Documento</th><th>Parte señalada</th><th>Qué detectó</th><th></th></tr></thead><tbody>{rows}</tbody></table></body></html>'''


def build_workbook(items,rules,target):
    wb=Workbook(); ws=wb.active; ws.title="Casos"
    headers=("Caso","Abrir caso","Documento","Texto o elemento","Parte señalada","Lo que detectó el sistema",
             "Decisión","Motivo u observaciones","Nombre del revisor","Fecha")
    ws.append(headers); editable=set(range(7,11))
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78")
        cell.alignment=Alignment(wrap_text=True,vertical="top")
    reports=target/"casos"; reports.mkdir(parents=True,exist_ok=True)
    created=[]
    for index,item in enumerate(items,1):
        rule=rules.get(item["rule_id"],{})
        finding=item["finding"]; name=f"caso-{index:02d}.html"; destination=reports/name
        destination.write_text(_render_case(item,rule,index,len(items)),encoding="utf-8"); created.append(destination)
        ws.append((index,"ABRIR",item["file"],finding.get("original_text") or "Sin fragmento textual",
                   finding.get("problem_span") or "No disponible",finding.get("message") or item.get("message"),"","","",""))
        row=ws.max_row; link=ws.cell(row,2); link.hyperlink=f"casos/{name}"; link.style="Hyperlink"
        for column in editable:ws.cell(row,column).fill=PatternFill("solid",fgColor="FFF2CC")
    formula='"'+",".join(DECISIONS)+'"'; validation=DataValidation(type="list",formula1=formula,allow_blank=False)
    validation.error="Seleccione una de las cuatro decisiones."; validation.errorTitle="Valor no permitido"
    ws.add_data_validation(validation); validation.add(f"G2:G{len(items)+1}")
    widths=(8,14,42,70,28,60,32,55,25,16)
    for index,width in enumerate(widths,1):ws.column_dimensions[ws.cell(1,index).column_letter].width=width
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:J{len(items)+1}"
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical="top")
    guide=wb.create_sheet("Instrucciones")
    guide.append(("Paso","Instrucción")); guide.append(("1","Abra INICIAR_AQUI.html o use ABRIR en la hoja Casos."))
    guide.append(("2","Revise exclusivamente el texto resaltado y la explicación del sistema."))
    guide.append(("3","Seleccione una opción en Decisión. Las observaciones son opcionales salvo en excepciones."))
    guide.append(("4","Devuelva únicamente este Excel completado; no modifique los HTML."))
    for cell in guide[1]:cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78")
    guide.column_dimensions["A"].width=24; guide.column_dimensions["B"].width=100
    for row in guide.iter_rows():
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical="top")
    wb.save(target/"validacion_editorial.xlsx")
    index=target/"INICIAR_AQUI.html"; index.write_text(_render_index(items),encoding="utf-8")
    return [index,*created]


def build_package(pilot_root,target):
    pilot_root=Path(pilot_root).resolve(); target=Path(target).resolve(); target.mkdir(parents=True,exist_ok=True)
    summary=json.loads((pilot_root/"summary.json").read_text(encoding="utf-8")); items=[]
    for item in summary["validation_sample"]:
        source=(pilot_root/item["report"]).resolve()
        if not source.is_file() or pilot_root not in source.parents:raise ValueError(f"Reporte inválido: {source}")
        review_path=source.with_suffix(".json")
        review=json.loads(review_path.read_text(encoding="utf-8"))
        candidates=(review.get("editorial_review") or {}).get("findings",[])
        exact=[finding for finding in candidates
               if finding.get("id")==item["finding_id"]
               and finding.get("rule_id")==item["rule_id"]
               and (finding.get("location") or {}).get("paragraph_id")==item["paragraph_id"]
               and finding.get("message")==item["message"]]
        if len(exact)!=1:
            raise ValueError(f"No se pudo identificar de forma unívoca el caso {item['finding_id']} en {review_path}")
        items.append({**item,"finding":exact[0]})
    rules={x["rule_id"]:x for x in load_editorial_catalog()["rules"]}
    created=build_workbook(items,rules,target)
    (target/"INSTRUCCIONES.md").write_text(instructions_text(len(items)),encoding="utf-8")
    files=[target/"INSTRUCCIONES.md",target/"validacion_editorial.xlsx",*created]
    manifest={"schema_version":"editorial-validation-handoff-2.0","generated_at":datetime.now(timezone.utc).isoformat(),
              "source_pilot":str(pilot_root),"cases":len(items),"contains_original_docx":False,
              "files":[{"path":x.relative_to(target).as_posix(),"sha256":digest(x),"bytes":x.stat().st_size} for x in sorted(files)]}
    (target/"MANIFEST.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return manifest


def main(argv=None):
    parser=argparse.ArgumentParser(description="Prepara una carpeta para delegar la validación editorial.")
    parser.add_argument("--pilot",default="results/editorial_pilot_81")
    parser.add_argument("--output",default="results/paquete_validacion_editorial")
    args=parser.parse_args(argv)
    try:manifest=build_package(args.pilot,args.output)
    except (OSError,ValueError,KeyError,json.JSONDecodeError) as exc:print(f"Error: {exc}"); return 1
    print(f"Paquete creado en {Path(args.output).resolve()}: {manifest['cases']} casos, {len(manifest['files'])} archivos.")
    return 0


if __name__=="__main__":raise SystemExit(main())
