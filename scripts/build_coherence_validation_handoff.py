"""Prepara una muestra real para validar comprensión y coherencia de observaciones CRF."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from datetime import datetime, timezone
from html import escape
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


OPTIONS = {
    "comprension": ("Sí - claramente", "Se entiende parcialmente", "No se entiende", "No puedo determinarlo"),
    "coherencia": ("Sí", "No", "No puedo determinarlo"),
    "responsable": ("Sí", "No", "No aplica"),
    "conclusion": ("Sí", "No", "No puede determinarse"),
    "cifras": ("Sí", "No", "No aparecen cifras", "Necesito consultar cuadro o DOCX"),
    "problema": ("Ninguno", "Redacción confusa", "Responsable ambiguo", "Información insuficiente",
                 "Secuencia ilógica", "Conclusión no sustentada", "Contradicción",
                 "Inconsistencia de cifras", "Repetición innecesaria", "Otro"),
}


def _plain(text):
    normalized=unicodedata.normalize("NFD",str(text or ""))
    return "".join(char for char in normalized if unicodedata.category(char)!="Mn").upper()


def _digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def extract_observations(path):
    """Extrae sólo resultados marcados expresamente CON OBSERVACIÓN: SÍ (X)."""
    document=Document(path); paragraphs=[p.text.strip() for p in document.paragraphs]
    headers=[]
    for index,text in enumerate(paragraphs):
        normalized=_plain(text)
        if "NUM. DEL RESULTADO" not in normalized and "NUM DEL RESULTADO" not in normalized:
            continue
        number_match=re.search(r"RESULTADO\s*:?\s*(\d+)",normalized)
        observed=bool(re.search(r"\bSI\s*\(\s*X\s*\)",normalized))
        headers.append((index,number_match.group(1) if number_match else "No identificado",observed))
    results=[]
    for position,(start,number,observed) in enumerate(headers):
        if not observed:continue
        end=headers[position+1][0] if position+1<len(headers) else len(paragraphs)
        scope=paragraphs[start+1:end]
        description=next((i for i,text in enumerate(scope) if "DESCRIPCION DEL RESULTADO" in _plain(text)),None)
        body=scope[description+1:] if description is not None else scope
        useful=[]
        for text in body:
            if not text:continue
            normalized=_plain(text)
            if normalized.startswith(("ELABORO:","REVISO:","AUTORIZO:")):break
            if text.lstrip().startswith("____________________"):break
            if useful and re.match(r"^\d+\.\s+",text) and text==text.upper():break
            useful.append(text)
        content="\n\n".join(useful).strip()
        if len(content)<120:continue
        results.append({"document":Path(path).name,"document_path":str(Path(path).resolve()),
                        "document_sha256":_digest(path),"result_number":number,"text":content,
                        "paragraphs":len(useful),"characters":len(content),
                        "contains_numbers":bool(re.search(r"\d",content)),
                        "source_range":{"start_paragraph":start+1,"end_paragraph":end}})
    return results


def select_sample(candidates,maximum=40):
    def bucket(item):
        length="short" if item["characters"]<700 else "medium" if item["characters"]<1600 else "long"
        return length+("_numbers" if item["contains_numbers"] else "_no_numbers")
    groups={}
    for item in candidates:
        item={**item,"selection_bucket":bucket(item)}
        groups.setdefault(item["selection_bucket"],[]).append(item)
    for values in groups.values():
        values.sort(key=lambda item:hashlib.sha256(f'{item["document"]}|{item["result_number"]}'.encode()).hexdigest())
    selected=[]; used=set(); used_documents=set()
    while len(selected)<maximum:
        progressed=False
        for name in sorted(groups):
            choice=next((item for item in groups[name]
                         if (item["document_sha256"],item["result_number"]) not in used
                         and item["document_sha256"] not in used_documents),None)
            if choice:
                groups[name].remove(choice); used.add((choice["document_sha256"],choice["result_number"]));
                used_documents.add(choice["document_sha256"]); selected.append(choice); progressed=True
                if len(selected)>=maximum:break
        if not progressed:break
    if len(selected)<maximum:
        remaining=sorted((item for values in groups.values() for item in values),
                         key=lambda item:hashlib.sha256(f'{item["document_sha256"]}|{item["result_number"]}'.encode()).hexdigest())
        for item in remaining:
            key=(item["document_sha256"],item["result_number"])
            if key not in used:
                used.add(key); selected.append(item)
                if len(selected)>=maximum:break
    return selected


def _case_html(item,index,total):
    paragraphs="".join(f"<p>{escape(part)}</p>" for part in item["text"].split("\n\n"))
    previous=f'caso-{index-1:02d}.html' if index>1 else '../INICIAR_AQUI.html'
    following=f'caso-{index+1:02d}.html' if index<total else '../INICIAR_AQUI.html'
    return f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Caso de coherencia {index}</title><style>body{{font:17px/1.6 system-ui;margin:auto;max-width:960px;padding:28px;color:#17232d;background:#f2f5f6}}main{{background:white;border:1px solid #d6dfe5;border-radius:12px;padding:28px}}.task{{background:#e4f3f5;border-left:5px solid #087187;padding:15px}}.observation{{font-family:Georgia,serif;font-size:19px;background:#fafafa;border:1px solid #d6dfe5;padding:22px;border-radius:8px}}.questions{{background:#fff7dc;padding:18px;border-radius:8px}}li{{margin:9px 0}}nav{{display:flex;justify-content:space-between;margin-top:24px}}a{{color:#075f73;font-weight:700}}</style></head><body><main><p>Caso {index} de {total}</p><h1>Resultado {escape(item['result_number'])}</h1><p class="task"><strong>Su tarea:</strong> lea la observación completa y responda las preguntas de la fila {index+1} en el Excel. Evalúe si una persona que no participó en la auditoría puede comprender qué ocurrió y por qué se concluyó así.</p><h2>Observación</h2><div class="observation">{paragraphs}</div><h2>Preguntas que debe contestar en el Excel</h2><ol class="questions"><li>¿Se entiende claramente la observación?</li><li>¿Las ideas siguen una secuencia coherente?</li><li>¿Se identifica quién realizó u omitió la acción?</li><li>¿La conclusión se desprende de los hechos narrados?</li><li>¿Las cifras parecen concordantes o debe consultarse el cuadro/DOCX?</li><li>Si existe un problema, copie el fragmento exacto y explíquelo brevemente.</li></ol><details><summary>Ubicación y datos de control</summary><p><strong>Documento:</strong> {escape(item['document'])}</p><p><strong>Resultado:</strong> {escape(item['result_number'])}</p><p><strong>SHA-256:</strong> <code>{item['document_sha256']}</code></p></details><nav><a href="{previous}">← Anterior</a><a href="../INICIAR_AQUI.html">Lista</a><a href="{following}">Siguiente →</a></nav></main></body></html>'''


def _index_html(items):
    rows="".join(f'<tr><td>{i}</td><td>{escape(item["document"])}</td><td>{escape(item["result_number"])}</td><td>{item["paragraphs"]}</td><td>{item["characters"]}</td><td><a href="casos/caso-{i:02d}.html">Abrir</a></td></tr>' for i,item in enumerate(items,1))
    return f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Validación de coherencia</title><style>body{{font:16px/1.5 system-ui;margin:auto;max-width:1150px;padding:28px;color:#17232d}}.start{{background:#e4f3f5;border-left:5px solid #087187;padding:16px}}table{{width:100%;border-collapse:collapse}}th,td{{padding:10px;border-bottom:1px solid #d6dfe5;text-align:left}}th{{background:#eef2f4}}a{{color:#075f73;font-weight:700}}</style></head><body><h1>Validación humana de coherencia y comprensión</h1><div class="start">Abra cada caso, lea la observación completa y conteste la misma fila en <code>validacion_coherencia.xlsx</code>. No corrija el DOCX ni necesita conocer códigos del sistema.</div><p>Muestra: {len(items)} observaciones marcadas como “con observación”.</p><table><thead><tr><th>Caso</th><th>Documento</th><th>Resultado</th><th>Párrafos</th><th>Caracteres</th><th></th></tr></thead><tbody>{rows}</tbody></table></body></html>'''


def _validation(sheet,column,options,last_row):
    validation=DataValidation(type="list",formula1='"'+','.join(options)+'"',allow_blank=False)
    validation.error="Seleccione una opción de la lista."; validation.errorTitle="Valor no permitido"
    sheet.add_data_validation(validation); letter=sheet.cell(1,column).column_letter
    validation.add(f"{letter}2:{letter}{last_row}")


def _workbook(items,path):
    wb=Workbook(); ws=wb.active; ws.title="Casos"
    headers=("Caso","Abrir caso","Documento","Resultado","Observación completa",
             "1. Comprensión","2. Coherencia","3. Responsable claro","4. Conclusión sustentada",
             "5. Concordancia de cifras","6. Problema principal","7. Fragmento problemático",
             "8. Explicación breve","Reescritura opcional","Nombre del revisor","Fecha")
    ws.append(headers)
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78")
        cell.alignment=Alignment(wrap_text=True,vertical="top")
    for index,item in enumerate(items,1):
        ws.append((index,"ABRIR",item["document"],item["result_number"],item["text"],"","","","","","","","","","",""))
        row=ws.max_row; ws.cell(row,2).hyperlink=f"casos/caso-{index:02d}.html"; ws.cell(row,2).style="Hyperlink"
        for column in range(6,17):ws.cell(row,column).fill=PatternFill("solid",fgColor="FFF2CC")
        ws.row_dimensions[row].height=100
    for column,key in zip(range(6,12),("comprension","coherencia","responsable","conclusion","cifras","problema")):
        _validation(ws,column,OPTIONS[key],len(items)+1)
    widths=(8,13,42,12,85,24,20,22,25,30,30,38,50,55,25,16)
    for column,width in enumerate(widths,1):ws.column_dimensions[ws.cell(1,column).column_letter].width=width
    ws.freeze_panes="F2"; ws.auto_filter.ref=f"A1:P{len(items)+1}"
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical="top")
    guide=wb.create_sheet("Instrucciones")
    guide.append(("Paso","Qué hacer"))
    guide.append((1,"Abra INICIAR_AQUI.html o el vínculo ABRIR de cada fila."))
    guide.append((2,"Lea la observación completa; no evalúe ortografía aislada, sino comprensión y coherencia."))
    guide.append((3,"Complete las seis listas amarillas. Si señala un problema, copie el fragmento exacto y explique por qué."))
    guide.append((4,"La reescritura es opcional. Devuelva únicamente este Excel completado."))
    guide.append(("Importante","No consulte ni comparta estos textos fuera de medios institucionales."))
    for cell in guide[1]:cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78")
    guide.column_dimensions["A"].width=18; guide.column_dimensions["B"].width=110
    for row in guide.iter_rows():
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical="top")
    wb.save(path)


def build_package(input_root,output_root,maximum=40):
    input_root=Path(input_root).resolve(); output_root=Path(output_root).resolve()
    files=sorted(input_root.rglob("*.docx"),key=lambda path:path.as_posix().casefold())
    if not files:raise ValueError(f"No se encontraron DOCX en {input_root}")
    candidates=[]; failures=[]
    for path in files:
        try:candidates.extend(extract_observations(path))
        except (OSError,ValueError) as exc:failures.append({"file":path.name,"error":str(exc)[:300]})
    items=select_sample(candidates,maximum)
    if not items:raise ValueError("No se encontraron resultados marcados expresamente con observación.")
    output_root.mkdir(parents=True,exist_ok=True); cases=output_root/"casos"; cases.mkdir(exist_ok=True)
    created=[]
    for index,item in enumerate(items,1):
        target=cases/f"caso-{index:02d}.html"; target.write_text(_case_html(item,index,len(items)),encoding="utf-8"); created.append(target)
    index_path=output_root/"INICIAR_AQUI.html"; index_path.write_text(_index_html(items),encoding="utf-8")
    workbook=output_root/"validacion_coherencia.xlsx"; _workbook(items,workbook)
    instructions=output_root/"INSTRUCCIONES.txt"
    instructions.write_text("Abra INICIAR_AQUI.html, revise cada observación y complete la misma fila en validacion_coherencia.xlsx. Devuelva solamente el Excel completado.\n",encoding="utf-8")
    manifest={"schema_version":"coherence-validation-handoff-1.0","generated_at":datetime.now(timezone.utc).isoformat(),
              "source_directory":str(input_root),"documents_scanned":len(files),"observations_found":len(candidates),
              "cases":len(items),"failures":failures,"contains_original_docx":False,
              "selection":[{k:v for k,v in item.items() if k not in {"text","document_path"}} for item in items]}
    files_out=[index_path,workbook,instructions,*created]
    manifest["files"]=[{"path":path.relative_to(output_root).as_posix(),"sha256":_digest(path),"bytes":path.stat().st_size} for path in files_out]
    (output_root/"MANIFEST.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return manifest


def main(argv=None):
    parser=argparse.ArgumentParser(description="Prepara validación humana de coherencia sobre CRF reales.")
    parser.add_argument("input"); parser.add_argument("--output",default="results/paquete_validacion_coherencia")
    parser.add_argument("--cases",type=int,default=40)
    args=parser.parse_args(argv)
    try:manifest=build_package(args.input,args.output,args.cases)
    except (OSError,ValueError) as exc:print(f"Error: {exc}"); return 1
    print(f"Paquete creado: {manifest['cases']} casos de {manifest['observations_found']} observaciones en {manifest['documents_scanned']} documentos.")
    return 0


if __name__=="__main__":raise SystemExit(main())
