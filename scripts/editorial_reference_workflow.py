"""Cierra la doble validación y construye bancos editoriales pareados.

Los libros y JSON con texto institucional deben escribirse bajo ``results/``.
El único artefacto apto para Git es el banco sintético, que nunca copia texto,
nombres de archivo ni identificadores documentales del corpus de origen.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


DECISION_KEYS = ("usable", "category", "interpretation", "improves", "protected")
CATEGORIES = (
    "Ortografía o gramática",
    "Coherencia o claridad",
    "Cifras o datos",
    "Formato o cuadro",
    "Criterio técnico o normativo",
    "Otro",
)
OPTIONS = {
    "usable": ("Sí", "No", "No puedo determinarlo"),
    "category": CATEGORIES,
    "interpretation": ("Sí", "No", "No puedo determinarlo"),
    "improves": ("Sí", "No", "No puedo determinarlo"),
    "protected": ("Sí", "No", "No cambiaron datos protegidos", "No puedo determinarlo"),
}
BLIND_HEADERS = (
    "Caso ciego", "Tipo de evidencia", "Texto anterior", "Texto final",
    "1. ¿Sirve como referencia?", "2. Categoría correcta",
    "3. ¿Interpretación correcta?", "4. ¿El texto final mejora?",
    "5. ¿Cambio de datos aceptable?", "Explicación breve",
    "Nombre del segundo revisor", "Fecha",
)
ADJ_HEADERS = (
    "Caso ciego", "Texto anterior", "Texto final", "Primera decisión",
    "Segunda decisión", "1. Decisión final: sirve", "2. Categoría final",
    "3. Interpretación final", "4. Texto final mejora",
    "5. Cambio de datos aceptable", "Justificación de adjudicación",
    "Nombre del adjudicador", "Fecha",
)


def _read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _write_json(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _text_hash(value):
    return hashlib.sha256(str(value or "").encode("utf-8")).hexdigest()


def _now():
    return datetime.now(timezone.utc).isoformat()


def _json_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def _load_sources(validation_path, corpus_path):
    validation = _read_json(validation_path)
    corpus = _read_json(corpus_path)
    if validation.get("schema_version") != "editorial-revision-validation-1.0":
        raise ValueError("La validación primaria no tiene un esquema compatible.")
    if corpus.get("schema_version") != "editorial-revision-evidence-1.0":
        raise ValueError("El corpus de evidencia no tiene un esquema compatible.")
    records = {item["id"]: item for item in corpus["records"]}
    selected = []
    for case in validation["cases"]:
        if not case.get("eligible_for_evaluation_reference"):
            continue
        record = records.get(case["evidence_id"])
        if record is None:
            raise ValueError(f"No existe la evidencia {case['evidence_id']} en el corpus.")
        selected.append((case, record))
    if not selected:
        raise ValueError("No hay referencias elegibles para la segunda revisión.")
    return validation, selected


def _blind_id(evidence_id, validation_hash, position):
    token = hashlib.sha256(f"{validation_hash}:{evidence_id}".encode()).hexdigest()[:10]
    return f"BR-{position:03d}-{token}"


def _style_sheet(sheet, widths):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _add_validations(sheet, first_row, last_row, start_column):
    for offset, key in enumerate(DECISION_KEYS):
        values = OPTIONS[key]
        formula = '"' + ",".join(values) + '"'
        validation = DataValidation(type="list", formula1=formula, allow_blank=False)
        validation.error = "Seleccione una opción de la lista."
        validation.errorTitle = "Valor no permitido"
        sheet.add_data_validation(validation)
        validation.add(f"{sheet.cell(first_row, start_column + offset).column_letter}{first_row}:"
                       f"{sheet.cell(last_row, start_column + offset).column_letter}{last_row}")


def build_blind_package(validation_path, corpus_path, workbook_path, manifest_path):
    """Crea el libro ciego y un manifiesto separado que no contiene texto."""
    validation, selected = _load_sources(validation_path, corpus_path)
    validation_hash = _file_hash(validation_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Segunda revisión"
    sheet.append(BLIND_HEADERS)
    manifest_cases = []
    for position, (case, record) in enumerate(selected, 1):
        blind_id = _blind_id(case["evidence_id"], validation_hash, position)
        fixed = (blind_id, record["evidence_type"], record.get("before"), record.get("after"))
        sheet.append((*fixed, None, None, None, None, None, None, None, None))
        manifest_cases.append({
            "blind_id": blind_id,
            "evidence_id": case["evidence_id"],
            "fixed_sha256": [_text_hash(value) for value in fixed],
        })
    _style_sheet(sheet, {"A": 22, "B": 20, "C": 70, "D": 70, "E": 24, "F": 28,
                         "G": 26, "H": 28, "I": 30, "J": 45, "K": 28, "L": 16})
    _add_validations(sheet, 2, len(selected) + 1, 5)
    red = PatternFill("solid", fgColor="FFC7CE")
    sheet.conditional_formatting.add(
        f"E2:J{len(selected)+1}",
        FormulaRule(
            formula=[
                'OR($E2="",AND($E2="Sí",COUNTBLANK($F2:$I2)>0),'
                'AND(OR($E2="No",$E2="No puedo determinarlo"),'
                'COUNTBLANK($F2:$I2)>0,COUNTBLANK($F2:$I2)<4),'
                'AND(OR($E2="No",$E2="No puedo determinarlo"),'
                'COUNTBLANK($F2:$I2)=4,$J2=""))'
            ],
            fill=red,
        ),
    )
    instructions = workbook.create_sheet("Instrucciones")
    instructions.append(["Paquete ciego de segunda revisión"])
    instructions.append(["Evalúe cada caso sin consultar la primera revisión. No edite las columnas A:D."])
    instructions.append([
        "Si la primera respuesta es Sí, complete las cinco decisiones. Si es No o No puedo "
        "determinarlo, puede dejar 2 a 5 vacías, pero debe explicar el motivo. El nombre no se importará."
    ])
    Path(workbook_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)
    manifest = {
        "schema_version": "editorial-blind-review-manifest-1.0",
        "generated_at": _now(),
        "privacy": {"contains_raw_institutional_text": False, "safe_for_version_control": False},
        "source": {"validation_sha256": validation_hash, "corpus_sha256": _file_hash(corpus_path)},
        "workbook_sha256": _file_hash(workbook_path),
        "case_count": len(manifest_cases),
        "cases": manifest_cases,
    }
    _write_json(manifest_path, manifest)
    return manifest


def _decision(values, start):
    decision = dict(zip(DECISION_KEYS, values[start:start + 5]))
    invalid = [(key, value) for key, value in decision.items() if value not in OPTIONS[key]]
    if invalid:
        raise ValueError("Hay decisiones vacías o no permitidas: " + ", ".join(key for key, _ in invalid))
    return decision


def _second_review_decision(values, start, explanation):
    """Admite respuestas terminales justificadas sin inventar decisiones no aplicables."""
    usable = values[start]
    if usable not in OPTIONS["usable"]:
        raise ValueError("La decisión sobre utilidad está vacía o no está permitida.")
    remaining = values[start + 1:start + 5]
    terminal = usable in {"No", "No puedo determinarlo"}
    if terminal and all(value in (None, "") for value in remaining):
        if not str(explanation or "").strip():
            raise ValueError("Una respuesta terminal sin decisiones 2 a 5 requiere explicación.")
        return {
            "usable": usable,
            "category": None,
            "interpretation": None,
            "improves": None,
            "protected": None,
        }
    return _decision(values, start)


def import_second_review(workbook_path, manifest_path, validation_path, corpus_path, output_path):
    """Importa la segunda revisión, rechazando cualquier cambio en la evidencia fija."""
    manifest = _read_json(manifest_path)
    if manifest.get("schema_version") != "editorial-blind-review-manifest-1.0":
        raise ValueError("El manifiesto ciego no es compatible.")
    validation, selected = _load_sources(validation_path, corpus_path)
    if manifest["source"]["validation_sha256"] != _file_hash(validation_path):
        raise ValueError("La validación primaria cambió después de generar el paquete.")
    selected_by_id = {case["evidence_id"]: (case, record) for case, record in selected}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Segunda revisión" not in workbook.sheetnames:
        raise ValueError("El libro no contiene la hoja Segunda revisión.")
    sheet = workbook["Segunda revisión"]
    if tuple(cell.value for cell in sheet[1]) != BLIND_HEADERS:
        raise ValueError("Las columnas del paquete ciego fueron modificadas.")
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if len(rows) != manifest["case_count"]:
        raise ValueError("El número de casos no coincide con el manifiesto ciego.")
    imported = []
    fixed_changes = []
    for row, item in zip(rows, manifest["cases"]):
        case, record = selected_by_id[item["evidence_id"]]
        canonical = (item["blind_id"], record["evidence_type"], record.get("before"), record.get("after"))
        for index, (expected, actual) in enumerate(zip(canonical, row[:4])):
            if expected != actual and not (expected in (None, "") and actual in (None, "")):
                fixed_changes.append({"blind_id": item["blind_id"], "column": BLIND_HEADERS[index]})
        decision = _second_review_decision(row, 4, row[9])
        imported.append({
            "blind_id": item["blind_id"], "evidence_id": item["evidence_id"],
            "decision": decision, "explanation": row[9], "review_date": _json_value(row[11]),
        })
    if fixed_changes:
        columns = ", ".join(f"{x['blind_id']} ({x['column']})" for x in fixed_changes)
        raise ValueError(f"Se detectaron cambios en evidencia fija: {columns}.")
    payload = {
        "schema_version": "editorial-second-review-1.0", "generated_at": _now(),
        "privacy": {"contains_raw_institutional_text": True, "contains_free_text": True,
                    "safe_for_version_control": False, "reviewer_names_stored": False},
        "source": {"workbook_sha256": _file_hash(workbook_path),
                   "manifest_sha256": _file_hash(manifest_path),
                   "validation_sha256": _file_hash(validation_path),
                   "corpus_sha256": _file_hash(corpus_path)},
        "summary": {
            "cases": len(imported),
            "terminal_justified": sum(
                item["decision"]["usable"] in {"No", "No puedo determinarlo"}
                and item["decision"]["category"] is None
                for item in imported
            ),
        },
        "cases": imported,
    }
    _write_json(output_path, payload)
    return payload


def _decision_text(decision):
    return " | ".join(
        f"{key}={decision[key] if decision.get(key) is not None else 'No aplica'}"
        for key in DECISION_KEYS
    )


def build_adjudication_workbook(validation_path, second_path, corpus_path, workbook_path, manifest_path):
    """Genera un libro que contiene exclusivamente los desacuerdos."""
    validation, selected = _load_sources(validation_path, corpus_path)
    second = _read_json(second_path)
    if second.get("schema_version") != "editorial-second-review-1.0":
        raise ValueError("La segunda revisión no es compatible.")
    second_by_id = {item["evidence_id"]: item for item in second["cases"]}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Adjudicación"
    sheet.append(ADJ_HEADERS)
    manifest_cases = []
    for first, record in selected:
        second_case = second_by_id.get(first["evidence_id"])
        if second_case is None:
            raise ValueError(f"Falta la segunda revisión de {first['evidence_id']}.")
        first_decision = first["human_decision"]
        second_decision = second_case["decision"]
        if all(first_decision[key] == second_decision[key] for key in DECISION_KEYS):
            continue
        fixed = (second_case["blind_id"], record.get("before"), record.get("after"),
                 _decision_text(first_decision), _decision_text(second_decision))
        sheet.append((*fixed, None, None, None, None, None, None, None, None))
        manifest_cases.append({
            "blind_id": second_case["blind_id"], "evidence_id": first["evidence_id"],
            "fixed_sha256": [_text_hash(value) for value in fixed],
        })
    _style_sheet(sheet, {"A": 22, "B": 60, "C": 60, "D": 65, "E": 65, "F": 24,
                         "G": 28, "H": 24, "I": 24, "J": 28, "K": 45, "L": 25, "M": 16})
    if manifest_cases:
        _add_validations(sheet, 2, len(manifest_cases) + 1, 6)
    Path(workbook_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)
    manifest = {
        "schema_version": "editorial-adjudication-manifest-1.0", "generated_at": _now(),
        "privacy": {"contains_raw_institutional_text": False, "safe_for_version_control": False},
        "source": {"validation_sha256": _file_hash(validation_path), "second_review_sha256": _file_hash(second_path)},
        "workbook_sha256": _file_hash(workbook_path), "case_count": len(manifest_cases), "cases": manifest_cases,
    }
    _write_json(manifest_path, manifest)
    return manifest


def import_adjudication(workbook_path, manifest_path, corpus_path, output_path):
    manifest = _read_json(manifest_path)
    if manifest.get("schema_version") != "editorial-adjudication-manifest-1.0":
        raise ValueError("El manifiesto de adjudicación no es compatible.")
    corpus = _read_json(corpus_path)
    records = {item["id"]: item for item in corpus["records"]}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Adjudicación"]
    if tuple(cell.value for cell in sheet[1]) != ADJ_HEADERS:
        raise ValueError("Las columnas de adjudicación fueron modificadas.")
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if len(rows) != manifest["case_count"]:
        raise ValueError("El número de adjudicaciones no coincide con el manifiesto.")
    imported = []
    for row, item in zip(rows, manifest["cases"]):
        record = records[item["evidence_id"]]
        canonical_hashes = [
            _text_hash(item["blind_id"]), _text_hash(record.get("before")), _text_hash(record.get("after")),
            item["fixed_sha256"][3], item["fixed_sha256"][4],
        ]
        if [_text_hash(value) for value in row[:5]] != canonical_hashes:
            raise ValueError(f"Se modificó evidencia fija en {item['blind_id']}.")
        decision = _decision(row, 5)
        if not row[10]:
            raise ValueError(f"Falta la justificación de adjudicación en {item['blind_id']}.")
        imported.append({"blind_id": item["blind_id"], "evidence_id": item["evidence_id"],
                         "decision": decision, "justification": row[10], "decision_date": _json_value(row[12])})
    payload = {
        "schema_version": "editorial-adjudication-1.0", "generated_at": _now(),
        "privacy": {"contains_raw_institutional_text": True, "contains_free_text": True,
                    "safe_for_version_control": False, "adjudicator_names_stored": False},
        "source": {"manifest_sha256": _file_hash(manifest_path),
                   "validation_sha256": manifest["source"]["validation_sha256"],
                   "second_review_sha256": manifest["source"]["second_review_sha256"],
                   "corpus_sha256": _file_hash(corpus_path)},
        "summary": {"cases": len(imported)}, "cases": imported,
    }
    _write_json(output_path, payload)
    return payload


def _eligible(decision):
    return (decision["usable"] == "Sí" and decision["interpretation"] == "Sí"
            and decision["improves"] == "Sí" and decision["protected"] != "No")


SYNTHETIC_TEMPLATES = {
    "Ortografía o gramática": ("El área entrego el informe número {n}.", "El área entregó el informe número {n}."),
    "Coherencia o claridad": ("El área revisó el expediente {n}, esto fue realizado por la misma.",
                               "El área revisó el expediente {n} y documentó el procedimiento."),
    "Cifras o datos": ("Las partidas fueron 30 y 40; el total reportado fue 80 en el caso {n}.",
                        "Las partidas fueron 30 y 40; el total reportado fue 70 en el caso {n}."),
    "Formato o cuadro": ("nota del cuadro {n}", "Nota del cuadro {n}."),
    "Criterio técnico o normativo": ("La observación del caso {n} incumplió la normativa aplicable.",
                                      "La observación del caso {n} incumplió el artículo 12 del Reglamento de Ejemplo."),
    "Otro": ("Se integró la documentación del caso {n} y también además se revisó.",
              "Se integró la documentación del caso {n} y además se revisó."),
}


def build_paired_banks(validation_path, second_path, adjudication_path, corpus_path,
                       real_output_path, synthetic_output_path):
    """Construye el banco final; el real es local y el sintético es versionable."""
    validation, selected = _load_sources(validation_path, corpus_path)
    second = _read_json(second_path)
    adjudication = _read_json(adjudication_path)
    if second.get("schema_version") != "editorial-second-review-1.0":
        raise ValueError("La segunda revisión no es compatible.")
    if adjudication.get("schema_version") != "editorial-adjudication-1.0":
        raise ValueError("La adjudicación no es compatible.")
    validation_hash = _file_hash(validation_path)
    corpus_hash = _file_hash(corpus_path)
    if (second.get("source", {}).get("validation_sha256") != validation_hash
            or second.get("source", {}).get("corpus_sha256") != corpus_hash):
        raise ValueError("La segunda revisión no corresponde a la validación y corpus actuales.")
    if (adjudication.get("source", {}).get("validation_sha256") != validation_hash
            or adjudication.get("source", {}).get("corpus_sha256") != corpus_hash
            or adjudication.get("source", {}).get("second_review_sha256") != _file_hash(second_path)):
        raise ValueError("La adjudicación no corresponde a las fuentes actuales.")
    if len({item["evidence_id"] for item in second["cases"]}) != len(second["cases"]):
        raise ValueError("La segunda revisión contiene evidencias duplicadas.")
    if len({item["evidence_id"] for item in adjudication["cases"]}) != len(adjudication["cases"]):
        raise ValueError("La adjudicación contiene evidencias duplicadas.")
    second_by_id = {item["evidence_id"]: item for item in second["cases"]}
    adjudicated_by_id = {item["evidence_id"]: item for item in adjudication["cases"]}
    eligible_ids = {first["evidence_id"] for first, _ in selected}
    if set(second_by_id) != eligible_ids:
        raise ValueError("La segunda revisión no cubre exactamente las referencias elegibles.")
    real_pairs = []
    exclusions = []
    categories = Counter()
    for first, record in selected:
        second_case = second_by_id.get(first["evidence_id"])
        if second_case is None:
            raise ValueError(f"Falta la segunda revisión de {first['evidence_id']}.")
        disagreement = any(first["human_decision"][key] != second_case["decision"][key] for key in DECISION_KEYS)
        if disagreement:
            adjudicated = adjudicated_by_id.get(first["evidence_id"])
            if adjudicated is None:
                raise ValueError(f"Falta adjudicar el desacuerdo {second_case['blind_id']}.")
            decision = adjudicated["decision"]
        else:
            decision = first["human_decision"]
        if not _eligible(decision):
            exclusions.append({"blind_id": second_case["blind_id"], "reason": "final_decision_not_eligible"})
            continue
        before, after = record.get("before"), record.get("after")
        if not before or not after:
            exclusions.append({"blind_id": second_case["blind_id"], "reason": "missing_pair_text"})
            continue
        if before == after:
            exclusions.append({"blind_id": second_case["blind_id"], "reason": "identical_pair_text"})
            continue
        pair_id = f"PAIR-{len(real_pairs)+1:03d}"
        category = decision["category"]
        categories[category] += 1
        real_pairs.append({
            "pair_id": pair_id, "blind_id": second_case["blind_id"], "category": category,
            "positive": {"text": before, "expects_finding": True},
            "negative": {"text": after, "expects_finding": False},
        })
    expected_adjudications = set()
    for first, _ in selected:
        second_case = second_by_id[first["evidence_id"]]
        if any(first["human_decision"][key] != second_case["decision"][key] for key in DECISION_KEYS):
            expected_adjudications.add(first["evidence_id"])
    if set(adjudicated_by_id) != expected_adjudications:
        raise ValueError("La adjudicación no cubre exactamente los desacuerdos detectados.")

    real_cases = []
    for pair in real_pairs:
        for variant, field in (("before", "positive"), ("after", "negative")):
            example = pair[field]
            real_cases.append({
                "case_id": f"{pair['pair_id']}-{variant}", "fragment": example["text"],
                "expected_finding": example["expects_finding"], "acceptable_spans": [],
                "source_pair_id": pair["pair_id"], "variant": variant, "category": pair["category"],
            })
    real_payload = {
        "schema_version": "editorial-paired-evaluation-bank-1.0", "generated_at": _now(),
        "privacy": {"contains_raw_institutional_text": True, "safe_for_version_control": False},
        "summary": {"pairs": len(real_pairs), "examples": len(real_pairs) * 2,
                    "excluded": len(exclusions), "by_category": dict(categories)},
        "cases": real_cases, "exclusions": exclusions,
    }
    synthetic_pairs = []
    for number, pair in enumerate(real_pairs, 1):
        before, after = SYNTHETIC_TEMPLATES[pair["category"]]
        synthetic_pairs.append({
            "pair_id": f"SYN-{number:03d}", "category": pair["category"],
            "positive": {"text": before.format(n=number), "expects_finding": True},
            "negative": {"text": after.format(n=number), "expects_finding": False},
        })
    synthetic_cases = []
    for pair in synthetic_pairs:
        for variant, field in (("before", "positive"), ("after", "negative")):
            example = pair[field]
            synthetic_cases.append({
                "case_id": f"{pair['pair_id']}-{variant}", "fragment": example["text"],
                "expected_finding": example["expects_finding"], "acceptable_spans": [],
                "source_pair_id": pair["pair_id"], "variant": variant, "category": pair["category"],
            })
    synthetic_payload = {
        "schema_version": "editorial-synthetic-paired-bank-1.0", "synthetic": True,
        "privacy": {"contains_raw_institutional_text": False, "safe_for_version_control": True},
        "summary": {"pairs": len(synthetic_pairs), "examples": len(synthetic_pairs) * 2,
                    "by_category": dict(categories)}, "cases": synthetic_cases,
    }
    _write_json(real_output_path, real_payload)
    _write_json(synthetic_output_path, synthetic_payload)
    return real_payload, synthetic_payload


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--validation", default="results/editorial_revision_validation/validated_cases.json")
    common.add_argument("--corpus", default="results/editorial_revision_corpus/corpus.json")
    prepare = sub.add_parser("prepare", parents=[common])
    prepare.add_argument("--workbook", default="results/editorial_second_review/segunda_revision_ciega.xlsx")
    prepare.add_argument("--manifest", default="results/editorial_second_review/manifest.json")
    second = sub.add_parser("import-second", parents=[common])
    second.add_argument("workbook"); second.add_argument("--manifest", required=True); second.add_argument("--output", required=True)
    adjudicate = sub.add_parser("prepare-adjudication", parents=[common])
    adjudicate.add_argument("--second", required=True); adjudicate.add_argument("--workbook", required=True); adjudicate.add_argument("--manifest", required=True)
    import_adj = sub.add_parser("import-adjudication")
    import_adj.add_argument("workbook"); import_adj.add_argument("--manifest", required=True); import_adj.add_argument("--corpus", required=True); import_adj.add_argument("--output", required=True)
    banks = sub.add_parser("build-banks", parents=[common])
    banks.add_argument("--second", required=True); banks.add_argument("--adjudication", required=True)
    banks.add_argument("--real-output", default="results/editorial_revision_validation/paired_evaluation_bank.json")
    banks.add_argument("--synthetic-output", default="tests/data/editorial_paired_synthetic.json")
    args = parser.parse_args(argv)
    if args.command == "prepare":
        result = build_blind_package(args.validation, args.corpus, args.workbook, args.manifest)
    elif args.command == "import-second":
        result = import_second_review(args.workbook, args.manifest, args.validation, args.corpus, args.output)
    elif args.command == "prepare-adjudication":
        result = build_adjudication_workbook(args.validation, args.second, args.corpus, args.workbook, args.manifest)
    elif args.command == "import-adjudication":
        result = import_adjudication(args.workbook, args.manifest, args.corpus, args.output)
    else:
        result, _ = build_paired_banks(args.validation, args.second, args.adjudication, args.corpus,
                                       args.real_output, args.synthetic_output)
    print(json.dumps(result.get("summary", {"cases": result.get("case_count")}), ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
