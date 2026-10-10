"""Pruebas automatizadas de la interfaz web interna (MVP).

Utiliza exclusivamente documentos sintéticos generados en carpetas temporales.
"""

from __future__ import annotations

import io
import time
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from revisor_web.app import create_app
from revisor_web.config import WebConfig
from revisor_web.storage import JobStorage
from revisor_web.validation import (
    CorruptDocumentError,
    FileTooLargeError,
    InvalidExtensionError,
    InvalidSignatureError,
    validate_docx_archive,
    validate_filename,
    validate_header_signature,
)


TEST_USER = "ASF\\revisor.uno"
TEST_HEADERS = {"X-Forwarded-User": TEST_USER}


@pytest.fixture
def synthetic_valid_docx(tmp_path: Path) -> Path:
    """Genera un documento DOCX sintético válido con una discrepancia para auditoría."""
    doc = Document()
    doc.add_paragraph("Auditoría Sintética de Prueba 2026")
    doc.add_paragraph("Cuadro 1. Recursos Asignados (miles de pesos)")

    table = doc.add_table(rows=4, cols=2)
    data = [
        ("Concepto", "Importe"),
        ("Rubro A", "100.0"),
        ("Rubro B", "50.0"),
        ("Total", "200.0"),  # Discrepancia: 100 + 50 != 200
    ]
    for r_idx, (col0, col1) in enumerate(data):
        table.cell(r_idx, 0).text = col0
        table.cell(r_idx, 1).text = col1

    doc.add_paragraph("El importe total ejercido fue de 150.0 miles de pesos.")
    target = tmp_path / "sintetico_prueba.docx"
    doc.save(target)
    return target


@pytest.fixture
def test_app_and_client(tmp_path: Path):
    """Crea una aplicación con configuración aislada en tmp_path."""
    jobs_dir = tmp_path / "web_jobs"
    cfg = WebConfig(
        jobs_dir=jobs_dir,
        max_file_size_mb=2,
        job_ttl_seconds=10,
        cleanup_interval_seconds=1,
        max_concurrent_jobs=2,
        max_queue_size=2,
        trusted_proxy_hosts=("testclient",),
    )
    app = create_app(cfg)
    client = TestClient(app, headers=TEST_HEADERS)
    return app, client, jobs_dir


def test_index_page(test_app_and_client):
    """Verifica que la página principal cargue y no dependa de CDNs."""
    _, client, _ = test_app_and_client
    response = client.get("/")
    assert response.status_code == 200
    assert "text/html" in response.headers["content-type"]
    html = response.text
    assert "Revisor de CRF · Servidor Interno" in html
    assert "Arrastre su archivo DOCX aquí" in html
    assert "metric-limitations" in html
    # Sin dependencias externas de CDNs
    assert "cdn.jsdelivr.net" not in html
    assert "cdnjs.cloudflare.com" not in html
    assert "fonts.googleapis.com" not in html


def test_health_check(test_app_and_client):
    """Verifica el endpoint técnico de salud."""
    _, client, _ = test_app_and_client
    response = client.get("/api/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert data["host"] == "127.0.0.1"
    assert data["max_file_size_mb"] == 2


def test_static_assets(test_app_and_client):
    """Verifica que el logo institucional y favicon se sirvan como assets públicos."""
    app, _, _ = test_app_and_client
    # Cliente sin credenciales de proxy para validar acceso público
    public_client = TestClient(app)
    res_logo = public_client.get("/static/logo_asf.png")
    assert res_logo.status_code == 200
    assert "image/png" in res_logo.headers["content-type"]
    assert len(res_logo.content) > 1000

    res_fav = public_client.get("/favicon.ico")
    assert res_fav.status_code == 200
    assert "image/x-icon" in res_fav.headers["content-type"]
    assert len(res_fav.content) > 100



def test_validation_rules():
    """Comprueba las funciones unitarias de validación técnica."""
    # 1. Extensión
    assert validate_filename("documento.docx") == "documento.docx"
    assert validate_filename("INFORME.DOCX") == "INFORME.DOCX"
    with pytest.raises(InvalidExtensionError):
        validate_filename("documento.pdf")
    with pytest.raises(InvalidExtensionError):
        validate_filename("")

    # 2. Firma mágica
    validate_header_signature(b"PK\x03\x04\x14\x00\x06\x00")
    with pytest.raises(InvalidSignatureError):
        validate_header_signature(b"%PDF-1.5 fake pdf")

    # 3. Archivo corrupto o no ZIP
    with pytest.raises(CorruptDocumentError):
        validate_docx_archive(Path("nonexistent.docx"))


def test_upload_invalid_extension(test_app_and_client):
    """Rechazo inmediato si la extensión no es .docx."""
    _, client, _ = test_app_and_client
    files = {"file": ("reporte.pdf", b"%PDF-1.4...", "application/pdf")}
    response = client.post("/api/jobs", files=files)
    assert response.status_code == 400
    assert "Solo se admiten documentos con extensión .docx" in response.json()["detail"]


def test_upload_invalid_signature(test_app_and_client):
    """Rechazo si el archivo tiene extensión .docx pero no es ZIP."""
    _, client, _ = test_app_and_client
    files = {"file": ("falso.docx", b"Esto es texto plano, no zip", "application/octet-stream")}
    response = client.post("/api/jobs", files=files)
    assert response.status_code == 400
    assert "firma ZIP" in response.json()["detail"]


def test_upload_empty_file(test_app_and_client):
    """Rechazo de archivos con 0 bytes."""
    _, client, _ = test_app_and_client
    files = {"file": ("vacio.docx", b"", "application/octet-stream")}
    response = client.post("/api/jobs", files=files)
    assert response.status_code == 400
    assert "vacío" in response.json()["detail"]


def test_upload_file_too_large(tmp_path: Path):
    """Rechazo con código 413 si excede el tamaño máximo configurado."""
    jobs_dir = tmp_path / "web_jobs"
    cfg = WebConfig(jobs_dir=jobs_dir, max_file_size_mb=1, trusted_proxy_hosts=("testclient",))
    app = create_app(cfg)
    client = TestClient(app, headers=TEST_HEADERS)

    large_payload = b"PK\x03\x04" + b"A" * (2 * 1024 * 1024)
    files = {"file": ("grande.docx", large_payload, "application/octet-stream")}
    response = client.post("/api/jobs", files=files)
    assert response.status_code == 413
    assert "excede el límite máximo" in response.json()["detail"]


def test_upload_fake_zip_missing_ooxml(test_app_and_client, tmp_path: Path):
    """Rechazo si es un ZIP pero no contiene la estructura word/document.xml."""
    _, client, _ = test_app_and_client
    zip_path = tmp_path / "fake.docx"
    with zipfile.ZipFile(zip_path, "w") as z:
        z.writestr("test.txt", "contenido cualquiera")

    files = {"file": ("fake.docx", zip_path.read_bytes(), "application/octet-stream")}
    response = client.post("/api/jobs", files=files)
    assert response.status_code == 400
    assert "estructura OOXML" in response.json()["detail"]


def test_full_job_lifecycle(test_app_and_client, synthetic_valid_docx: Path):
    """Ejecuta el ciclo completo: carga, revisión, consulta de estado y descargas."""
    _, client, jobs_dir = test_app_and_client

    original_bytes = synthetic_valid_docx.read_bytes()
    original_size = len(original_bytes)

    # 1. Enviar documento sintético
    files = {"file": ("sintetico_prueba.docx", original_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/jobs", files=files)
    assert response.status_code == 202
    data = response.json()
    job_id = data["job_id"]
    assert data["status"] == "recibido"
    assert data["original_filename"] == "sintetico_prueba.docx"
    assert data["file_size"] == original_size

    # 2. Esperar que el motor determinístico complete el procesamiento
    max_wait = 15
    start = time.time()
    job_data = None
    while time.time() - start < max_wait:
        poll_resp = client.get(f"/api/jobs/{job_id}")
        assert poll_resp.status_code == 200
        job_data = poll_resp.json()
        if job_data["status"] in {"terminado", "fallido"}:
            break
        time.sleep(0.3)

    assert job_data is not None
    assert job_data["status"] == "terminado", f"El trabajo falló: {job_data.get('error')}"

    # 3. Comprobar métricas del resumen
    summary = job_data["summary"]
    assert summary["alerts_count"] >= 1, "Debe detectar la discrepancia sintética (100+50!=200)"
    assert summary["schema_version"] == "review-1.2"
    assert summary["limitations_count"] == 0
    assert summary["duration_seconds"] > 0
    assert "NUM-TOTAL-001" in summary["by_rule"] or "NUM-CROSS-001" in summary["by_rule"]

    # 4. Descargar entregable JSON
    json_resp = client.get(f"/api/jobs/{job_id}/download/json")
    assert json_resp.status_code == 200
    assert "application/json" in json_resp.headers["content-type"]
    assert f"revision_{job_id[:8]}.json" in json_resp.headers["content-disposition"]
    json_content = json_resp.json()
    assert json_content["schema_version"] == "review-1.2"
    assert json_content["limitations"] == []
    assert len(json_content["alerts"]) >= 1
    # Comprobar que el JSON contiene commented_copy igual que la CLI
    assert "commented_copy" in json_content
    assert json_content["commented_copy"]["file"] == "commented.docx"
    assert json_content["commented_copy"]["comments_added"] >= 1

    # 5. Descargar entregable HTML
    html_resp = client.get(f"/api/jobs/{job_id}/download/html")
    assert html_resp.status_code == 200
    assert "text/html" in html_resp.headers["content-type"]
    assert f"revision_{job_id[:8]}.html" in html_resp.headers["content-disposition"]
    assert "<!doctype html>" in html_resp.text.lower()

    # 6. Visualización directa de HTML en navegador con aislamiento CSP reforzado
    view_resp = client.get(f"/api/jobs/{job_id}/view/html")
    assert view_resp.status_code == 200
    assert "<!doctype html>" in view_resp.text.lower()
    csp_header = view_resp.headers.get("content-security-policy", "")
    assert "connect-src 'none'" in csp_header
    assert "form-action 'none'" in csp_header
    assert "base-uri 'none'" in csp_header

    # 7. Descargar entregable DOCX comentado
    docx_resp = client.get(f"/api/jobs/{job_id}/download/docx")
    assert docx_resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument" in docx_resp.headers["content-type"]
    assert f"revision_{job_id[:8]}_comentada.docx" in docx_resp.headers["content-disposition"]

    # Verificar que el DOCX comentado es un ZIP válido y contiene comments.xml
    with zipfile.ZipFile(io.BytesIO(docx_resp.content)) as archive:
        names = set(archive.namelist())
        assert "word/document.xml" in names
        assert "word/comments.xml" in names, "Debe contener comentarios anclados de las alertas"

    # 8. Criterio de aceptación: el archivo sintético original NO fue modificado
    assert synthetic_valid_docx.read_bytes() == original_bytes


def test_download_before_finished_returns_409(test_app_and_client):
    """Comprueba que no se permitan descargas antes de que el trabajo termine."""
    _, client, jobs_dir = test_app_and_client
    storage = JobStorage(jobs_dir)
    meta = storage.create_job("test.docx", 100, owner=TEST_USER.casefold())
    job_id = meta["job_id"]

    resp = client.get(f"/api/jobs/{job_id}/download/json")
    assert resp.status_code == 409
    assert "aún no ha terminado" in resp.json()["detail"]


def test_ttl_cleanup(tmp_path: Path):
    """Verifica que los trabajos expirados se eliminen del disco, excepto si están en procesando."""
    jobs_dir = tmp_path / "web_jobs"
    storage = JobStorage(jobs_dir)

    # 1. Trabajo expirado terminado -> debe eliminarse
    meta_done = storage.create_job("terminado.docx", 500)
    job_id_done = meta_done["job_id"]
    storage.update_job(job_id_done, status="terminado")
    meta_path_done = jobs_dir / job_id_done / "meta.json"
    import json
    data_done = json.loads(meta_path_done.read_text(encoding="utf-8"))
    data_done["created_timestamp"] = time.time() - 100
    meta_path_done.write_text(json.dumps(data_done), encoding="utf-8")

    # 2. Trabajo expirado en 'procesando' -> NUNCA debe eliminarse
    meta_proc = storage.create_job("en_proceso.docx", 500)
    job_id_proc = meta_proc["job_id"]
    storage.update_job(job_id_proc, status="procesando")
    meta_path_proc = jobs_dir / job_id_proc / "meta.json"
    data_proc = json.loads(meta_path_proc.read_text(encoding="utf-8"))
    data_proc["created_timestamp"] = time.time() - 100
    meta_path_proc.write_text(json.dumps(data_proc), encoding="utf-8")

    # Limpiar con TTL de 30 segundos
    deleted = storage.cleanup_expired_jobs(ttl_seconds=30)
    assert job_id_done in deleted
    assert not (jobs_dir / job_id_done).exists()

    # El trabajo en proceso se mantiene intacto
    assert job_id_proc not in deleted
    assert (jobs_dir / job_id_proc).exists()
    assert storage.get_job(job_id_proc)["status"] == "procesando"


def test_server_saturation_rejection(tmp_path: Path):
    """Verifica el rechazo HTTP 503 cuando se excede la concurrencia y cola."""
    jobs_dir = tmp_path / "web_jobs"
    cfg = WebConfig(
        jobs_dir=jobs_dir,
        max_concurrent_jobs=1,
        max_queue_size=0,
        trusted_proxy_hosts=("testclient",),
    )
    app = create_app(cfg)
    client = TestClient(app, headers=TEST_HEADERS)

    # Forzar estado de saturación en el servicio
    app.state.service._active_jobs = 1

    files = {"file": ("sintetico.docx", b"PK\x03\x04...", "application/octet-stream")}
    response = client.post("/api/jobs", files=files)
    assert response.status_code == 503
    assert "saturado" in response.json()["detail"]
    assert response.headers.get("retry-after") == "10"


def test_zero_document_text_in_logs(test_app_and_client, synthetic_valid_docx: Path, caplog):
    """Verifica la política de CERO texto confidencial ni nombres de archivo en los registros."""
    import logging
    _, client, _ = test_app_and_client

    identifiable_filename = "documento_secreto_identificable_asf_9999.docx"

    with caplog.at_level(logging.INFO, logger="revisor_web"):
        files = {"file": (identifiable_filename, synthetic_valid_docx.read_bytes(), "application/octet-stream")}
        response = client.post("/api/jobs", files=files)
        assert response.status_code == 202
        job_id = response.json()["job_id"]

        for _ in range(30):
            job_data = client.get(f"/api/jobs/{job_id}").json()
            if job_data["status"] == "terminado":
                break
            time.sleep(0.2)

    all_logs = " ".join(record.message for record in caplog.records)
    # Ninguna frase textual del documento sintético debe figurar en logs
    assert "Auditoría Sintética" not in all_logs
    assert "Cuadro 1" not in all_logs
    assert "Recursos Asignados" not in all_logs
    assert "Rubro A" not in all_logs
    assert "El importe total ejercido fue de" not in all_logs
    # Tampoco el nombre de archivo del usuario debe figurar en los logs
    assert "documento_secreto_identificable_asf_9999" not in all_logs
    # Pero sí debe registrarse la trazabilidad técnica por UUID y tamaño
    assert job_id in all_logs
    assert "completado con éxito" in all_logs


def test_regression_invalid_upload_does_not_delete_prior_valid_job(
    test_app_and_client, synthetic_valid_docx: Path
):
    """Prueba de regresión: una carga fallida posterior no debe eliminar trabajos válidos previos."""
    _, client, jobs_dir = test_app_and_client

    # 1. Crear un trabajo válido y esperar que termine
    files_valid = {"file": ("valido.docx", synthetic_valid_docx.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    resp_valid = client.post("/api/jobs", files=files_valid)
    assert resp_valid.status_code == 202
    valid_job_id = resp_valid.json()["job_id"]

    for _ in range(30):
        poll = client.get(f"/api/jobs/{valid_job_id}").json()
        if poll["status"] == "terminado":
            break
        time.sleep(0.2)

    assert client.get(f"/api/jobs/{valid_job_id}").json()["status"] == "terminado"

    # Confirmar que existen sus artefactos
    assert (jobs_dir / valid_job_id / "report.json").exists()
    assert (jobs_dir / valid_job_id / "report.html").exists()
    assert (jobs_dir / valid_job_id / "commented.docx").exists()

    # 2. Cargar un DOCX inválido (que fallará en la validación OOXML interna)
    files_invalid = {"file": ("invalido.docx", b"PK\x03\x04fake_content_without_word_parts", "application/octet-stream")}
    resp_invalid = client.post("/api/jobs", files=files_invalid)
    assert resp_invalid.status_code == 400

    # 3. Comprobar que el trabajo válido anterior y todos sus artefactos siguen intactos
    poll_after = client.get(f"/api/jobs/{valid_job_id}")
    assert poll_after.status_code == 200
    assert poll_after.json()["status"] == "terminado"

    # Los tres artefactos del trabajo válido deben seguir existiendo y descargándose
    assert client.get(f"/api/jobs/{valid_job_id}/download/json").status_code == 200
    assert client.get(f"/api/jobs/{valid_job_id}/download/html").status_code == 200
    assert client.get(f"/api/jobs/{valid_job_id}/download/docx").status_code == 200


def test_storage_delete_job_strict_uuid(tmp_path: Path):
    """Comprueba que delete_job valide estrictamente el UUID y confine el borrado."""
    storage = JobStorage(tmp_path / "web_jobs")

    # Trabajo válido
    meta = storage.create_job("archivo.docx", 100)
    job_id = meta["job_id"]
    job_dir = storage.base_dir / job_id
    assert job_dir.exists()

    # Borrado dirigido
    assert storage.delete_job(job_id) is True
    assert not job_dir.exists()

    # Intento de borrado de ID no existente
    assert storage.delete_job("00000000000000000000000000000000") is False

    # Intento de inyección de ruta / escape de directorio
    assert storage.delete_job("../") is False
    assert storage.delete_job("..") is False
    assert storage.delete_job("invalido_no_hex") is False


def test_authentication_is_required_except_health(tmp_path: Path):
    """La salud es publica, pero la interfaz y API exigen identidad del proxy."""
    cfg = WebConfig(jobs_dir=tmp_path / "jobs", trusted_proxy_hosts=("testclient",))
    client = TestClient(create_app(cfg))

    assert client.get("/api/health").status_code == 200
    unauthorized = client.get("/")
    assert unauthorized.status_code == 401
    assert unauthorized.headers["www-authenticate"] == "Negotiate"


def test_local_demo_needs_no_proxy_but_cannot_bind_to_network(tmp_path: Path):
    """La presentación local es cómoda sin debilitar la configuración institucional."""
    cfg = WebConfig(
        jobs_dir=tmp_path / "demo-jobs",
        host="127.0.0.1",
        demo_mode=True,
        demo_identity="demo.local",
    )
    client = TestClient(create_app(cfg), client=("127.0.0.1", 50000))
    response = client.get("/")
    assert response.status_code == 200
    assert "Demostración local" in response.text
    assert response.headers["x-revisor-demo"] == "local-only"

    external_client = TestClient(create_app(cfg), client=("192.0.2.10", 50000))
    assert external_client.get("/").status_code == 403

    with pytest.raises(ValueError, match="loopback"):
        create_app(WebConfig(jobs_dir=tmp_path / "bad", host="0.0.0.0", demo_mode=True))


def test_forwarded_identity_is_rejected_from_untrusted_peer(tmp_path: Path):
    """Una cabecera inventada por un cliente no sustituye a IIS/WIA."""
    cfg = WebConfig(jobs_dir=tmp_path / "jobs", trusted_proxy_hosts=("127.0.0.1",))
    client = TestClient(create_app(cfg), headers=TEST_HEADERS)

    response = client.get("/")
    assert response.status_code == 403
    assert "proxy confiable" in response.json()["detail"]


def test_job_ownership_hides_status_view_and_download(test_app_and_client, synthetic_valid_docx: Path):
    """Un segundo usuario no puede descubrir ni recuperar artefactos ajenos."""
    _, owner_client, _ = test_app_and_client
    response = owner_client.post(
        "/api/jobs",
        files={"file": ("propio.docx", synthetic_valid_docx.read_bytes(), "application/octet-stream")},
    )
    job_id = response.json()["job_id"]
    other_headers = {"X-Forwarded-User": "ASF\\revisor.dos"}

    for route in (
        f"/api/jobs/{job_id}",
        f"/api/jobs/{job_id}/view/html",
        f"/api/jobs/{job_id}/download/json",
    ):
        denied = owner_client.get(route, headers=other_headers)
        assert denied.status_code == 404
        assert denied.json()["detail"] == "Trabajo no encontrado."


def test_security_headers_are_applied(test_app_and_client):
    """Las respuestas, incluidos errores, llevan defensas HTTP basicas."""
    _, client, _ = test_app_and_client
    response = client.get("/")
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "SAMEORIGIN"
    assert response.headers["referrer-policy"] == "no-referrer"
    assert response.headers["cache-control"] == "no-store"
    assert "object-src 'none'" in response.headers["content-security-policy"]


def test_startup_recovers_interrupted_jobs(tmp_path: Path):
    """Un reinicio convierte estados transitorios en fallos recuperables."""
    jobs_dir = tmp_path / "jobs"
    storage = JobStorage(jobs_dir)
    received = storage.create_job("uno.docx", 10, owner=TEST_USER.casefold())
    processing = storage.create_job("dos.docx", 10, owner=TEST_USER.casefold())
    storage.update_job(processing["job_id"], status="procesando")
    cfg = WebConfig(jobs_dir=jobs_dir, trusted_proxy_hosts=("testclient",))

    with TestClient(create_app(cfg), headers=TEST_HEADERS) as client:
        assert client.get("/api/health").status_code == 200

    for job_id in (received["job_id"], processing["job_id"]):
        meta = storage.get_job(job_id)
        assert meta["status"] == "fallido"
        assert "reinicio" in meta["error"]


def test_finding_decision_is_validated_and_owned(test_app_and_client, synthetic_valid_docx: Path):
    """Las decisiones solo aceptan estados previstos y hallazgos reales del propietario."""
    _, client, jobs_dir = test_app_and_client
    created = client.post(
        "/api/jobs",
        files={"file": ("decision.docx", synthetic_valid_docx.read_bytes(), "application/octet-stream")},
    )
    job_id = created.json()["job_id"]
    for _ in range(50):
        meta = client.get(f"/api/jobs/{job_id}").json()
        if meta["status"] in {"terminado", "fallido"}:
            break
        time.sleep(0.1)
    report = __import__("json").loads((jobs_dir / job_id / "report.json").read_text(encoding="utf-8"))
    finding_id = report["alerts"][0]["id"]

    accepted = client.post(
        f"/api/jobs/{job_id}/findings/{finding_id}/decision",
        json={"status": "confirmed", "comment": "Confirmado por cotejo."},
    )
    assert accepted.status_code == 200
    assert accepted.json()["reviewer"] == TEST_USER.casefold()
    assert accepted.json()["status"] == "confirmed"

    invalid_status = client.post(
        f"/api/jobs/{job_id}/findings/{finding_id}/decision",
        json={"status": "approved"},
    )
    assert invalid_status.status_code == 422
    too_long = client.post(
        f"/api/jobs/{job_id}/findings/{finding_id}/decision",
        json={"status": "confirmed", "comment": "x" * 501},
    )
    assert too_long.status_code == 422
    missing = client.post(
        f"/api/jobs/{job_id}/findings/no-existe/decision",
        json={"status": "rejected"},
    )
    assert missing.status_code == 404
    denied = client.post(
        f"/api/jobs/{job_id}/findings/{finding_id}/decision",
        headers={"X-Forwarded-User": "ASF\\revisor.dos"},
        json={"status": "rejected"},
    )
    assert denied.status_code == 404


def test_pilot_telemetry_is_sanitized_and_computes_metrics(test_app_and_client, caplog):
    """La bitacora permite medir el piloto sin registrar la identidad en claro."""
    import logging

    app, client, _ = test_app_and_client
    with caplog.at_level(logging.INFO, logger="revisor_web.audit"):
        response = client.get("/")
        assert response.status_code == 200
        app.state.telemetry.record(
            "job_completed",
            job_id="a" * 32,
            identity=TEST_USER.casefold(),
            status="terminado",
            duration_seconds=2.5,
        )
        app.state.telemetry.record("job_saturated", identity=TEST_USER.casefold())
        app.state.telemetry.record("job_purged", count=3)

    snapshot = app.state.telemetry.snapshot()
    assert snapshot["success_rate"] == 1.0
    assert snapshot["processing_p95_seconds"] == 2.5
    assert snapshot["saturation_rejections"] == 1
    assert snapshot["purged_jobs"] == 3
    logs = " ".join(record.message for record in caplog.records)
    assert TEST_USER.casefold() not in logs
    assert '"user_hash":"' in logs


def test_corrupt_decision_store_is_not_silently_overwritten(tmp_path: Path):
    """Una bitácora ilegible exige intervención y conserva sus bytes originales."""
    storage = JobStorage(tmp_path / "jobs")
    meta = storage.create_job("prueba.docx", 10, owner=TEST_USER.casefold())
    path = storage.base_dir / meta["job_id"] / "decisions.json"
    original = b"{registro-invalido"
    path.write_bytes(original)

    with pytest.raises(ValueError, match="no es legible"):
        storage.save_decision(
            meta["job_id"], "finding-1",
            {"status": "confirmed", "comment": None, "reviewer": TEST_USER.casefold()},
        )
    assert path.read_bytes() == original
