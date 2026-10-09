import json

import pytest
from openpyxl import load_workbook

from build_revision_evidence_handoff import _workbook, select_sample
from import_revision_validation import import_validation


def _record(identifier, document, category):
    return {
        "id": identifier,
        "evidence_type": "revision_only",
        "document": document,
        "document_sha256": identifier * 64,
        "part": "word/document.xml",
        "xml_paths": [f"/{identifier}"],
        "result_number": "1",
        "comment_id": None,
        "comment": None,
        "anchor_original": None,
        "anchor_final": None,
        "before": "Texto anterior.",
        "after": "Texto final.",
        "has_associated_revision": True,
        "suggested_category": category,
        "protection_check": {"safe": True, "removed": [], "added": []},
    }


def _fixture(tmp_path):
    records = [
        _record("a", "uno.docx", "Otro"),
        _record("b", "dos.docx", "Coherencia o claridad"),
    ]
    corpus = tmp_path / "corpus.json"
    corpus.write_text(json.dumps({
        "schema_version": "editorial-revision-evidence-1.0",
        "records": records,
    }), encoding="utf-8")
    template = tmp_path / "template.xlsx"
    _workbook(select_sample(records, len(records)), template)
    submitted = tmp_path / "submitted.xlsx"
    workbook = load_workbook(template)
    sheet = workbook["Casos"]
    for row in range(2, 4):
        sheet.cell(row, 10).value = "Sí"
        sheet.cell(row, 11).value = "Coherencia o claridad"
        sheet.cell(row, 12).value = "Sí"
        sheet.cell(row, 13).value = "Sí"
        sheet.cell(row, 14).value = "No cambiaron datos protegidos"
        sheet.cell(row, 16).value = "Revisor omitido del JSON"
    sheet.cell(2, 8).value = "Texto final corregido."
    workbook.save(submitted)
    return corpus, template, submitted


def test_imports_decisions_and_reports_immutable_edits_without_reviewer_name(tmp_path):
    corpus, template, submitted = _fixture(tmp_path)
    output = tmp_path / "validated.json"
    result = import_validation(submitted, corpus, output, template)
    assert result["summary"]["cases"] == 2
    assert result["summary"]["eligible_evaluation_references"] == 2
    assert result["summary"]["immutable_field_changes"] == 1
    assert result["immutable_field_changes"][0]["column"] == "Texto final"
    assert "Revisor omitido" not in output.read_text(encoding="utf-8")


def test_rejects_incomplete_decisions(tmp_path):
    corpus, template, submitted = _fixture(tmp_path)
    workbook = load_workbook(submitted)
    workbook["Casos"].cell(2, 10).value = None
    workbook.save(submitted)
    with pytest.raises(ValueError, match="decisiones vacías"):
        import_validation(submitted, corpus, tmp_path / "validated.json", template)
