import json
from dataclasses import replace
from decimal import Decimal
from zipfile import ZipFile

import pytest
from docx import Document

from revisor_crf.html_report import render_html
from revisor_crf.commented_docx import export_commented_docx
from revisor_crf.review import review_document
from revisor_crf.semantic import (
    AmountType,
    SemanticRelation,
    SemanticConfig,
    StaticSemanticProvider,
    build_semantic_cases,
    public_case,
    response_schema,
    system_prompt,
)
from revisor_crf.semantic_evaluation import evaluate_provider, select_model


class AgreeingProvider:
    def __init__(self, relation="same_measure", confidence=0.99):
        self.relation = relation
        self.confidence = confidence
        self.calls = []

    def metadata(self):
        return {"provider": "static", "model": "semantic-test", "digest": "abc123"}

    def analyze(self, cases, phase):
        self.calls.append((phase, [case["case_id"] for case in cases]))
        return [{
            "case_id": case["case_id"], "relation": self.relation,
            "table_measure": case["deterministic"]["table_measure"],
            "narrative_measure": case["deterministic"]["narrative_measure"],
            "confidence": self.confidence,
            "table_evidence": case["table"]["row"],
            "narrative_evidence": case["narrative"]["text"],
        } for case in cases]


class FailingProvider:
    def metadata(self):
        return {"provider": "static", "model": "unavailable", "digest": None}

    def analyze(self, cases, phase):
        raise RuntimeError("semantic_unavailable")


def semantic_document(tmp_path):
    document = Document()
    document.add_heading("Resultado 1", level=1)
    document.add_paragraph("La obra del mercado tuvo un total pagado de 120 pesos.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concepto"
    table.rows[0].cells[1].text = "Total pagado"
    row = table.add_row().cells
    row[0].text = "Construcción del mercado municipal zona norte"
    row[1].text = "100 pesos"
    path = tmp_path / "semantic.docx"
    document.save(path)
    return path


def config(tmp_path, **changes):
    base = SemanticConfig(cache_path=str(tmp_path / "semantic.sqlite3"))
    return replace(base, **changes)


def semantic_alerts(report):
    return [item for item in report["alerts"] if item["rule_id"] == "NUM-CROSS-ML-001"]


def test_prompt_defines_measure_as_amount_type_even_when_values_differ():
    prompt = system_prompt("associate")
    assert "tipo de importe coinciden" in prompt
    assert "aunque los valores numéricos sean distintos" in prompt
    assert "No compares ni exijas igualdad" in prompt
    assert "confianza 0.99 solo cuando" in prompt
    assert "no agregues comillas" in prompt


def test_hybrid_review_requires_two_matching_passes_and_is_cached(tmp_path):
    path = semantic_document(tmp_path)
    cfg = tmp_path / "semantic.json"
    cfg.write_text(json.dumps({"cache_path": str(tmp_path / "isolated.sqlite3")}), encoding="utf-8")
    provider = AgreeingProvider()
    first = review_document(path, semantic=True, semantic_provider=provider,
                            semantic_config_path=cfg)
    alert, = semantic_alerts(first)
    assert alert["origin"] == "hybrid"
    assert alert["operation"]["declarado"] == Decimal("100")
    assert alert["operation"]["calculado"] == Decimal("120")
    assert first["schema_version"] == "review-1.1"
    assert first["semantic_review"]["status"] == "complete"
    assert [call[0] for call in provider.calls] == ["associate", "verify"]
    commented = tmp_path / "hybrid-commented.docx"
    metadata = export_commented_docx(path, first, commented)
    assert metadata["comments_added"] == 1
    with ZipFile(commented) as archive:
        comments = archive.read("word/comments.xml").decode("utf-8")
    assert "hallazgo híbrido" in comments and "confianza 0.99" in comments
    calls = len(provider.calls)
    second = review_document(path, semantic=True, semantic_provider=provider,
                             semantic_config_path=cfg)
    assert len(provider.calls) == calls
    assert second["semantic_review"]["cache"]["hits"] == 2


def test_low_confidence_relation_is_html_only_suggestion(tmp_path):
    cfg = tmp_path / "semantic.json"
    cfg.write_text(json.dumps({"cache_path": str(tmp_path / "suggestions.sqlite3")}), encoding="utf-8")
    path = semantic_document(tmp_path)
    report = review_document(
        path, semantic=True,
        semantic_provider=AgreeingProvider("same_concept", 0.80), semantic_config_path=cfg,
    )
    assert semantic_alerts(report) == []
    suggestion, = report["semantic_review"]["suggestions"]
    assert suggestion["relation"] == "same_concept"
    html = render_html(report)
    assert "Relaciones semánticas que requieren revisión humana" in html
    assert "review-decisions-1.1" in html
    metadata = export_commented_docx(path, report, tmp_path / "suggestion-only.docx")
    assert metadata["comments_added"] == 0


def test_model_cannot_invent_measure_for_tuxtla_style_narrative(tmp_path):
    document = Document()
    document.add_heading("Resultado 1", level=1)
    document.add_paragraph(
        "De las operaciones seleccionadas, 147 pesos corresponden al Fondo de Aportaciones "
        "para el Fortalecimiento de los Municipios (FORTAMUN); 33 pesos al Fondo de "
        "Aportaciones para la Infraestructura Social Municipal (FAISMUN)."
    )
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Concepto", "Importe presupuestal", "Total pagado")):
        cell.text = text
    for label, budget, paid in (("Subtotal FORTAMUN", "147 pesos", "147 pesos"),
                                ("Subtotal FAISMUN", "33 pesos", "32 pesos")):
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = label, budget, paid
    path = tmp_path / "tuxtla-regression.docx"
    document.save(path)
    cfg = tmp_path / "semantic.json"
    cfg.write_text(json.dumps({"cache_path": str(tmp_path / "tuxtla.sqlite3")}), encoding="utf-8")
    provider = AgreeingProvider()
    report = review_document(path, semantic=True, semantic_provider=provider,
                             semantic_config_path=cfg)
    assert semantic_alerts(report) == []
    assert report["semantic_review"]["candidate_cases"] == 0
    assert provider.calls == []
    html = render_html(report)
    assert "Revisión semántica completada sin casos elegibles" in html
    assert "No se enviaron casos al modelo" in html


def test_amount_and_unit_row_is_not_a_semantic_concept(tmp_path):
    document = Document()
    document.add_paragraph("El valor de 25,360 UMA equivale a 2,753,335.20 pesos.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concepto"
    table.rows[0].cells[1].text = "Monto máximo por contrato"
    cells = table.add_row().cells
    cells[0].text = "25,360 UMA"
    cells[1].text = "21,000,000.00 pesos"
    path = tmp_path / "uma-row.docx"
    document.save(path)
    provider = AgreeingProvider()
    cfg = tmp_path / "semantic.json"
    cfg.write_text(json.dumps({"cache_path": str(tmp_path / "uma.sqlite3")}), encoding="utf-8")
    report = review_document(path, semantic=True, semantic_provider=provider,
                             semantic_config_path=cfg)
    assert report["semantic_review"]["candidate_cases"] == 0
    assert provider.calls == []


def test_unavailable_model_preserves_deterministic_report(tmp_path):
    cfg = tmp_path / "semantic.json"
    cfg.write_text(json.dumps({"cache_path": str(tmp_path / "failure.sqlite3")}), encoding="utf-8")
    report = review_document(
        semantic_document(tmp_path), semantic=True,
        semantic_provider=FailingProvider(), semantic_config_path=cfg,
    )
    assert report["semantic_review"]["status"] == "unavailable"
    assert semantic_alerts(report) == []
    assert any(item.get("element") == "semantic_review" for item in report["warnings"])


def test_invalid_evidence_is_rejected_without_alert(tmp_path):
    class InventingProvider(AgreeingProvider):
        def analyze(self, cases, phase):
            decisions = super().analyze(cases, phase)
            for decision in decisions:
                decision["narrative_evidence"] = "texto que no existe"
            return decisions

    cfg = tmp_path / "semantic.json"
    cfg.write_text(json.dumps({"cache_path": str(tmp_path / "invalid.sqlite3")}), encoding="utf-8")
    report = review_document(
        semantic_document(tmp_path), semantic=True,
        semantic_provider=InventingProvider(), semantic_config_path=cfg,
    )
    assert semantic_alerts(report) == []
    assert report["semantic_review"]["invalid_responses"] == 1


def test_config_restricts_local_profile_to_loopback(tmp_path):
    path = tmp_path / "config.json"
    path.write_text(json.dumps({"base_url": "http://192.168.1.20:11434/api"}), encoding="utf-8")
    with pytest.raises(ValueError, match="perfil local"):
        SemanticConfig.load(path)


def test_contract_schema_and_public_cases_do_not_expose_internal_objects():
    schema = response_schema(["case-1"])
    assert schema["properties"]["decisions"]["minItems"] == 1
    case = {"case_id": "case-1", "table": {}, "narrative": {}, "deterministic": {},
            "_table_number": {"private": True}}
    assert "_table_number" not in public_case(case)
    assert SemanticRelation.SAME_MEASURE.value == "same_measure"
    assert AmountType.PAID.value == "pagado"


def test_model_selection_prioritizes_precision_then_speed():
    results = [
        {"model": "slow", "metrics": {"precision_same_measure": 0.99, "accuracy": 0.96,
                                        "duration_seconds": 500, "invalid_responses": 0}},
        {"model": "fast", "metrics": {"precision_same_measure": 0.99, "accuracy": 0.96,
                                        "duration_seconds": 300, "invalid_responses": 0}},
        {"model": "imprecise", "metrics": {"precision_same_measure": 0.97, "accuracy": 0.99,
                                             "duration_seconds": 10, "invalid_responses": 0}},
    ]
    assert select_model(results)["model"] == "fast"


def test_evaluation_keeps_repeated_structural_ids_separate_by_document():
    def case(row):
        return {
            "case_id": "same-structural-id",
            "table": {"text": "100 pesos", "row": row, "column": "Total pagado"},
            "narrative": {"text": "El total pagado fue de 120 pesos."},
            "deterministic": {"table_measure": "pagado", "narrative_measure": "pagado"},
        }

    class EvaluationProvider:
        def analyze(self, cases, _phase):
            return [{
                "case_id": item["case_id"],
                "relation": "same_measure" if item["table"]["row"] == "FORTAMUN" else "unrelated",
                "table_measure": "pagado", "narrative_measure": "pagado", "confidence": 0.99,
                "table_evidence": item["table"]["text"],
                "narrative_evidence": item["narrative"]["text"],
            } for item in cases]

    records = [
        {"document_sha256": "a" * 64, "case": case("FORTAMUN"),
         "expected_relation": "same_measure"},
        {"document_sha256": "b" * 64, "case": case("FAISMUN"),
         "expected_relation": "unrelated"},
    ]
    metrics = evaluate_provider(EvaluationProvider(), records, batch_size=2)
    assert metrics["records"] == 2
    assert metrics["accuracy"] == 1.0
    assert metrics["document_p95_seconds"] >= 0
