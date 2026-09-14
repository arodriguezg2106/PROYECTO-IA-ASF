from collections import Counter
from lxml import etree

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part

from audit_crf import inventory, output_elements, structure_check
from format_stats import analyze_formats
from revisor_crf import read_docx


def p(text):
    element = OxmlElement("w:p")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    element.append(run)
    return element


def wrapper(tag, *children):
    element = OxmlElement(tag)
    element.extend(children)
    return element


def extract(document, tmp_path):
    path = tmp_path / "extended.docx"
    document.save(path)
    result = read_docx(path, mode="all")
    entries = inventory(path)
    elements = list(output_elements(result))
    structural = structure_check(path, elements)
    assert structural == {"duplicados_elementos": [], "errores_orden_origen": []}
    keys = [(e["source"]["part"], e["source"]["xml_path"]) for e in elements]
    assert len(keys) == len(set(keys))
    assert len({e["id"] for e in elements}) == len(elements)
    assert sum(e["type"] == "paragraph" for e in elements) == len(entries)
    # Compare independent XML hashes, not just counts.
    import hashlib
    texts = {(e["source"]["part"], e["source"]["xml_path"]): e["text"]
             for e in elements if e["type"] == "paragraph"}
    for entry in entries:
        assert hashlib.sha256(texts[(entry["parte"], entry["ruta_xml"])].encode()).hexdigest() == entry["sha256_texto"]
    return result, path


def test_nested_controls_tables_and_custom_xml_order(tmp_path):
    d = Document()
    d.add_paragraph("Antes")
    table = d.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Celda"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Anidada"
    # A content control wraps both a paragraph and a table, including nested controls.
    d.element.body.insert(1, wrapper("w:sdt", wrapper("w:sdtContent", p("Control"),
                              wrapper("w:sdt", wrapper("w:sdtContent", table._tbl)))))
    d.element.body.insert(2, wrapper("w:customXml", p("Otro contenedor")))
    d.add_paragraph("Después")
    result, _ = extract(d, tmp_path)
    assert [b["type"] for b in result["blocks"]] == ["paragraph", "paragraph", "table", "paragraph", "paragraph"]
    assert [b["text"] for b in result["blocks"] if b["type"] == "paragraph"] == ["Antes", "Control", "Otro contenedor", "Después"]


def test_inline_and_block_revisions_and_fields(tmp_path):
    d = Document()
    paragraph = d.add_paragraph("Inicio ")
    inserted = wrapper("w:ins", wrapper("w:r", OxmlElement("w:t")))
    inserted[0][0].text = "nuevo"
    deleted = wrapper("w:del", wrapper("w:r", OxmlElement("w:delText")))
    deleted[0][0].text = "anterior"
    paragraph._p.extend([inserted, deleted])
    field = wrapper("w:fldSimple", wrapper("w:r", OxmlElement("w:t")))
    field[0][0].text = "42"
    paragraph._p.append(wrapper("w:sdt", wrapper("w:sdtContent", field)))
    instruction = wrapper("w:r", OxmlElement("w:instrText"))
    instruction[0].text = " PAGE "
    paragraph._p.append(instruction)
    d.element.body.insert(1, wrapper("w:moveFrom", p("Origen")))
    d.element.body.insert(2, wrapper("w:moveTo", p("Destino")))
    result, _ = extract(d, tmp_path)
    assert result["blocks"][0]["text"] == "Inicio nuevoanterior42"
    fragments = result["blocks"][0]["content"]
    assert fragments[1]["revisions"][0]["type"] == "ins"
    assert fragments[2]["revisions"][0]["type"] == "del"
    assert fragments[-1]["type"] == "field_instruction"
    assert result["blocks"][1]["revisions"][0]["type"] == "moveFrom"


def test_textbox_anchoring_and_no_double_text(tmp_path):
    d = Document()
    paragraph = d.add_paragraph("Antes")
    box = wrapper("w:txbxContent", p("Texto de forma"))
    paragraph._p.append(wrapper("w:r", wrapper("w:pict", box)))
    paragraph.add_run("Después")
    d.add_paragraph("Final")
    result, _ = extract(d, tmp_path)
    host = result["blocks"][0]
    assert host["text"] == "AntesDespués"
    assert host["blocks"][0]["text"] == "Texto de forma"
    assert [c["type"] for c in host["content"]] == ["text", "block", "text"]
    assert host["content"][1]["id"] == host["blocks"][0]["id"]


def test_notes_comments_and_separators(tmp_path):
    d = Document()
    d.add_paragraph("Cuerpo")
    for plural, singular, rel in (("footnotes", "footnote", RT.FOOTNOTES),
                                  ("endnotes", "endnote", RT.ENDNOTES), ("comments", "comment", RT.COMMENTS)):
        root = parse_xml(f'<w:{plural} {nsdecls("w")}/>')
        entry = wrapper("w:" + singular, p("Nota sintética"))
        entry.set(qn("w:id"), "1")
        root.append(entry)
        if singular != "comment":
            empty = wrapper("w:" + singular, wrapper("w:p", wrapper("w:r", OxmlElement("w:separator"))))
            empty.set(qn("w:type"), "separator")
            empty.set(qn("w:id"), "-1")
            root.append(empty)
        part = Part(PackURI(f"/word/{plural}.xml"),
                    f"application/vnd.openxmlformats-officedocument.wordprocessingml.{plural}+xml",
                    etree.tostring(root), d.part.package)
        d.part.relate_to(part, rel)
    result, _ = extract(d, tmp_path)
    assert Counter(s["kind"] for s in result["supplementary_parts"]) == {"footnote": 2, "endnote": 2, "comment": 1}
    assert len(result["blocks"]) == 1


def test_wrapped_rows_and_cells(tmp_path):
    d = Document()
    table = d.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Dato"
    row = table._tbl.tr_lst[0]
    cell = row.tc_lst[0]
    row.append(wrapper("w:sdt", wrapper("w:sdtContent", cell)))
    table._tbl.append(wrapper("w:ins", row))
    result, _ = extract(d, tmp_path)
    assert result["blocks"][0]["rows"][0]["cells"][0]["text"] == "Dato"


def test_tab_stops_are_not_text_and_style_inheritance(tmp_path):
    d = Document()
    style = d.styles.add_style("Inherited heading", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = d.styles["Heading 2"]
    paragraph = d.add_paragraph("Título sintético", style=style)
    from docx.shared import Pt
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(36))
    paragraph.runs[0].font.size = Pt(13)
    paragraph.runs[0].bold = True
    result, filename = extract(d, tmp_path)
    assert result["blocks"][0]["text"] == "Título sintético"
    assert result["blocks"][0]["heading_level"] == 2
    stats = analyze_formats(filename)
    assert stats["niveles_esquema_resueltos"] == {"1": 1}
    assert stats["combinaciones"][0]["tamanos_pt"] == [13.0]
    assert stats["combinaciones"][0]["negritas"] == ["true"]


def test_nested_footer_controls_and_shared_part(tmp_path):
    from docx.enum.section import WD_SECTION_START
    d = Document()
    d.add_paragraph("Cuerpo")
    footer = d.sections[0].footer
    footer._element.append(wrapper("w:sdt", wrapper("w:sdtContent",
                           wrapper("w:sdt", wrapper("w:sdtContent", p("Pie controlado"))))))
    d.add_section(WD_SECTION_START.NEW_PAGE)
    result, _ = extract(d, tmp_path)
    first = result["sections"][0]["headers_and_footers"][0]
    inherited = result["sections"][1]["headers_and_footers"][0]
    assert [b["text"] for b in first["blocks"]] == ["", "Pie controlado"]
    assert inherited["blocks"] == []
    assert inherited["content_ref"] == first["id"]
