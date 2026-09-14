from decimal import Decimal
import pytest
from docx import Document
from revisor_crf import read_docx
from revisor_crf.numbers import detect_numbers, extract_numbers


def extract(doc, tmp_path):
    path = tmp_path / "synthetic.docx"
    doc.save(path)
    return extract_numbers(read_docx(path))["cifras"]


@pytest.mark.parametrize("source", ["column_header", "row_header", "table_title", "table_note"])
def test_structural_units(source, tmp_path):
    d = Document()
    if source == "table_title":
        t = d.add_table(rows=3, cols=2)
        t.cell(0, 0).merge(t.cell(0, 1)).text = "Cuadro sintético\n(Pesos)"
        t.cell(1, 0).text = "Concepto"
        t.cell(1, 1).text = "Importe"
        t.cell(2, 0).text = "Rubro"
        t.cell(2, 1).text = "12"
    else:
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Concepto"
        t.cell(0, 1).text = "Importe en pesos" if source == "column_header" else "Valor"
        t.cell(1, 0).text = "Importe en pesos" if source == "row_header" else "Rubro"
        t.cell(1, 1).text = "12"
        if source == "table_note":
            d.add_paragraph("Nota: cifras en pesos")
    n, = extract(d, tmp_path)
    assert n["unit_source"] == source
    assert n["valor_normalizado"] == Decimal(12)
    assert n["confidence"] == ("medium" if source == "table_note" else "high")


def test_general_row_unit_certainty(tmp_path):
    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Número de municipios"
    t.cell(0, 1).text = "12"
    n, = extract(d, tmp_path)
    assert (n["unit_source"], n["confidence"], n["tipo"]) == ("row_header", "high", "cantidad")


def test_narrative_is_not_a_unit_declaration(tmp_path):
    d = Document()
    d.add_paragraph("Los municipios entregaron la documentación solicitada.")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "12"
    n, = extract(d, tmp_path)
    assert n["unit_source"] == "none" and n["confidence"] == "low"
    assert not n["unit_candidates"][0]["applicable"]


def test_priority_and_conflict():
    sources = ["explicit", "column_header", "row_header", "table_title", "table_note", "previous_paragraph"]
    for i, source in enumerate(sources):
        candidates = [{"unit": "pesos", "source": s, "element_id": s, "unambiguous": True} for s in reversed(sources[i:])]
        n, = detect_numbers("12", {"table_unit_candidates": candidates})
        assert n["unit_source"] == source
    n, = detect_numbers("12 pesos", {"encabezado_columna": "miles de pesos", "column_header_unambiguous": True})
    assert n["confidence"] == "low" and n["classification"] == "ambiguous"
    assert not n["usable_for_arithmetic"]


@pytest.mark.parametrize("text", ["Auditoría núm. 001234", "Número de auditoría: 001234", "Auditoría AB-001234"])
def test_audit_number_separate_category(text):
    n, = detect_numbers(text)
    assert n["tipo"] == n["classification"] == "numero_de_auditoria"
    assert n["valor_original"].endswith("001234")
    assert n["usable_for_identifier_matching"] and not n["usable_for_arithmetic"]


@pytest.mark.parametrize("text", ["001234", "Se realizaron 12 auditorías", "Ejercicio 2024", "Folio 001234"])
def test_numbers_not_promoted_to_audit(text):
    n, = detect_numbers(text)
    assert n["tipo"] != "numero_de_auditoria"


def test_uncertain_heading_never_high():
    n, = detect_numbers("12", {"encabezado_columna": "pesos", "column_header_unambiguous": False})
    assert n["confidence"] == "medium"
    assert n["reason"]


def test_audit_column_label_supports_compound_identifier():
    n, = detect_numbers("AB-001234", {"encabezado_columna": "Número de auditoría"})
    assert n["tipo"] == "numero_de_auditoria" and n["usable_for_identifier_matching"]
    assert n["texto_original"] == n["valor_original"] == "AB-001234"
    n, = detect_numbers("12 pesos", {"encabezado_columna": "Importe de auditoría"})
    assert n["tipo"] == "importe"
