import hashlib
import json
from decimal import Decimal
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement

from revisor_crf.cli import main
from revisor_crf.commented_docx import export_commented_docx, verify_commented_docx
from revisor_crf.html_report import render_html
from revisor_crf.review import review_document


def save_document(tmp_path, build, name="review.docx"):
    document = Document()
    build(document)
    path = tmp_path / name
    document.save(path)
    return path


def amount_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concepto"
    table.rows[0].cells[1].text = "Importe"
    for label, amount in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = amount
    return table


def alerts(report, rule_id):
    return [alert for alert in report["alerts"] if alert["rule_id"] == rule_id]


def test_incorrect_total_has_exact_operation_and_evidence(tmp_path):
    def build(document):
        document.add_paragraph("NÚM. DEL RESULTADO: 4 CON OBSERVACIÓN SÍ ( ) NO ( X )")
        amount_table(document, [
            ("Obra norte", "100 pesos"), ("Obra sur", "200 pesos"), ("Total", "250 pesos")])
    path = save_document(tmp_path, build)
    report = review_document(path)
    alert, = alerts(report, "NUM-TOTAL-001")
    assert alert["severity"] == "high"
    assert alert["operation"]["declarado"] == Decimal("250")
    assert alert["operation"]["calculado"] == Decimal("300")
    assert alert["operation"]["diferencia"] == Decimal("-50")
    assert len(alert["evidence"]) == 3
    assert all(item["ubicacion"]["elemento_id"] for item in alert["evidence"])
    assert alert["where_to_review"]["resultado"] == "4"
    assert alert["where_to_review"]["tabla_ordinal"] == 1
    assert alert["where_to_review"]["fila"] == 4
    assert alert["where_to_review"]["encabezado_columna"] == "Importe"
    assert alert["where_to_review"]["texto_busqueda"]


def test_correct_total_and_multiple_subtotal_blocks(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Obra uno", "40 pesos"), ("Obra dos", "60 pesos"), ("Subtotal", "100 pesos"),
        ("Obra tres", "25 pesos"), ("Obra cuatro", "75 pesos"), ("Total observado", "100 pesos")]))
    report = review_document(path)
    assert alerts(report, "NUM-TOTAL-001") == []
    assert report["coverage"]["NUM-TOTAL-001"]["evaluated"] == 2


def test_named_subtotals_feed_grand_total_without_explicit_unit(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("FORTAMUN", ""),
        ("Energía eléctrica", "16,851,276.55"),
        ("Vehículos", "990,800.02"),
        ("Subtotal FORTAMUN", "17,842,076.57"),
        ("FAISMUN", ""),
        ("Edificación", "2,336,597.35"),
        ("División de terrenos y construcción de obras de urbanización", "1,559,559.22"),
        ("Subtotal FAISMUN", "3,896,156.57"),
        ("PFM", ""),
        ("Sentencias", "7,910,707.99"),
        ("Subtotal PFM", "7,910,707.99"),
        ("Total", "29,648,941.13"),
    ]))
    report = review_document(path)
    assert alerts(report, "NUM-TOTAL-001") == []
    assert report["coverage"]["NUM-TOTAL-001"] == {
        "evaluated": 4, "excluded": 0, "alerts": 0,
    }


def test_grand_total_uses_named_subtotals_without_double_counting_details(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Fondo A", ""), ("Detalle A", "10"), ("Subtotal Fondo A", "10"),
        ("Fondo B", ""), ("Detalle B", "20"), ("Subtotal Fondo B", "20"),
        ("Total", "31"),
    ]))
    alert, = alerts(review_document(path), "NUM-TOTAL-001")
    assert alert["operation"]["calculado"] == Decimal("30")
    assert alert["operation"]["block_strategy"] == "named_subtotals"
    assert [item["valor"] for item in alert["evidence"] if item["rol"] == "sumando"] == [
        Decimal("10"), Decimal("20"),
    ]


def test_monetary_table_title_wins_over_units_word_in_row_description(tmp_path):
    def build(document):
        table = document.add_table(rows=2, cols=2)
        title = table.cell(0, 0).merge(table.cell(0, 1))
        title.text = "TIPO DE ADJUDICACIÓN DE LAS ADQUISICIONES REVISADAS\n(Pesos)"
        table.cell(1, 0).text = "Objeto del contrato"
        table.cell(1, 1).text = "Monto ejercido"
        detail = table.add_row().cells
        detail[0].text = "Adquisición de dos unidades RAM 1200 Modelo 2025"
        detail[1].text = "990,800.02"
        total = table.add_row().cells
        total[0].text = "Total"
        total[1].text = "990,800.02"

    report = review_document(save_document(tmp_path, build))
    assert alerts(report, "NUM-TOTAL-001") == []
    assert report["coverage"]["NUM-TOTAL-001"] == {
        "evaluated": 1, "excluded": 0, "alerts": 0,
    }


def test_incompatible_units_make_total_not_evaluable(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Obra norte", "100 pesos"), ("Contratos", "2 obras"), ("Total", "100 pesos")]))
    report = review_document(path)
    alert, = alerts(report, "NUM-TOTAL-001")
    assert alert["severity"] == "medium"
    assert alert["operation"]["type"] == "not_evaluable"
    assert alert["not_evaluable_reason"]


def test_total_label_accepts_trailing_punctuation(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Obra", "10 pesos"), ("Importe total:", "20 pesos")]))
    alert, = alerts(review_document(path), "NUM-TOTAL-001")
    assert alert["operation"]["type"] == "sum"


def test_negative_value_is_flagged_and_not_silently_summed(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Ajuste", "(20) pesos"), ("Obra", "100 pesos"), ("Total", "80 pesos")]))
    report = review_document(path)
    negative, = alerts(report, "DATA-NEG-001")
    assert negative["severity"] == "medium"
    total, = alerts(report, "NUM-TOTAL-001")
    assert total["operation"]["type"] == "not_evaluable"


def test_table_narrative_mismatch_requires_concept_scope_and_unit(tmp_path):
    def build(document):
        document.add_heading("Resultado 1", level=1)
        amount_table(document, [("Obra Alfa", "100 pesos"), ("Obra Beta", "200 pesos")])
        document.add_paragraph("La Obra Alfa reportó un importe de 120 pesos.")
        document.add_paragraph("Una materia distinta reportó 999 pesos.")
    report = review_document(save_document(tmp_path, build))
    alert, = alerts(report, "NUM-CROSS-001")
    assert alert["operation"]["declarado"] == Decimal("100")
    assert alert["operation"]["calculado"] == Decimal("120")
    assert {item["rol"] for item in alert["evidence"]} == {"cifra_tabular", "cifra_narrativa"}
    assert report["coverage"]["NUM-CROSS-001"]["evaluated"] == 1


def test_equal_table_narrative_value_produces_no_alert(tmp_path):
    def build(document):
        document.add_heading("Resultado 1", level=1)
        amount_table(document, [("Obra Alfa", "100 pesos")])
        document.add_paragraph("La Obra Alfa reportó 100 pesos.")
    report = review_document(save_document(tmp_path, build))
    assert alerts(report, "NUM-CROSS-001") == []
    assert report["coverage"]["NUM-CROSS-001"]["evaluated"] == 1


def test_table_narrative_does_not_cross_explicit_result_numbers(tmp_path):
    def build(document):
        document.add_paragraph("NÚM. DEL RESULTADO: 2 CON OBSERVACIÓN SÍ ( ) NO ( X )")
        amount_table(document, [("Adquisición RAM", "100 pesos")])
        document.add_paragraph("NÚM. DEL RESULTADO: 7 CON OBSERVACIÓN SÍ ( ) NO ( X )")
        document.add_paragraph("La Adquisición RAM reportó un importe de 120 pesos.")

    report = review_document(save_document(tmp_path, build))
    assert alerts(report, "NUM-CROSS-001") == []
    assert report["coverage"]["NUM-CROSS-001"]["evaluated"] == 0


def test_ambiguous_narrative_relationship_is_medium_alert(tmp_path):
    def build(document):
        document.add_heading("Resultado 1", level=1)
        amount_table(document, [("Obra Alfa", "100 pesos")])
        document.add_paragraph("La Obra Alfa reportó 100 pesos.")
        document.add_paragraph("Después, la Obra Alfa registró 110 pesos.")
    report = review_document(save_document(tmp_path, build))
    alert, = alerts(report, "NUM-CROSS-001")
    assert alert["severity"] == "medium"
    assert alert["operation"]["type"] == "not_evaluable"


def test_concept_match_with_incompatible_unit_is_not_evaluable(tmp_path):
    def build(document):
        document.add_heading("Resultado 1", level=1)
        amount_table(document, [("Obra Alfa", "100 pesos")])
        document.add_paragraph("La Obra Alfa reportó 100 obras.")
    alert, = alerts(review_document(save_document(tmp_path, build)), "NUM-CROSS-001")
    assert alert["severity"] == "medium"
    assert alert["operation"]["type"] == "not_evaluable"


def test_observation_conclusion_conflict_and_consistency(tmp_path):
    def conflict(document):
        document.add_heading("Resultado 1", level=1)
        document.add_paragraph("La revisión terminó con observación.")
        document.add_heading("Conclusión", level=2)
        document.add_paragraph("Conclusión: sin observaciones.")
    report = review_document(save_document(tmp_path, conflict, "conflict.docx"))
    alert, = alerts(report, "DOC-OBS-001")
    assert alert["severity"] == "medium"
    assert len(alert["evidence"]) == 2

    def consistent(document):
        document.add_heading("Resultado 2", level=1)
        document.add_paragraph("La revisión terminó sin observaciones.")
        document.add_heading("Conclusión", level=2)
        document.add_paragraph("Conclusión: sin observación.")
    assert alerts(review_document(save_document(tmp_path, consistent, "consistent.docx")), "DOC-OBS-001") == []


def test_observation_scope_matches_separate_result_heading(tmp_path):
    def build(document):
        document.add_heading("Resultado 3", level=1)
        document.add_paragraph("La revisión terminó con observación.")
        document.add_heading("Conclusión del Resultado 3", level=1)
        document.add_paragraph("Sin observaciones.")
    alert, = alerts(review_document(save_document(tmp_path, build)), "DOC-OBS-001")
    assert alert["operation"]["type"] == "status_comparison"


def test_observation_checkbox_in_table_matches_explicit_conclusion(tmp_path):
    def build(document):
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "NÚM. DEL RESULTADO: 7 CON OBSERVACIÓN SÍ ( X ) NO ( )"
        document.add_paragraph("En virtud de lo anterior, se concluye que no se identificaron irregularidades.")
    alert, = alerts(review_document(save_document(tmp_path, build)), "DOC-OBS-001")
    assert alert["operation"]["declarado"] == "positive"
    assert alert["operation"]["calculado"] == "negative"
    assert alert["evidence"][0]["ubicacion"]["tabla"] is not None


def test_observation_checkbox_consistent_conclusion_has_no_alert(tmp_path):
    def build(document):
        document.add_paragraph("NÚM. DEL RESULTADO: 2 CON OBSERVACIÓN SÍ ( ) NO ( X )")
        document.add_paragraph("Se concluye que las obras no presentan irregularidades.")
    report = review_document(save_document(tmp_path, build))
    assert alerts(report, "DOC-OBS-001") == []
    assert report["coverage"]["DOC-OBS-001"]["evaluated"] == 1


def test_alert_ids_are_deterministic_and_source_is_unchanged(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Obra uno", "10 pesos"), ("Obra dos", "20 pesos"), ("Suma", "40 pesos")]))
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    first, second = review_document(path), review_document(path)
    assert [item["id"] for item in first["alerts"]] == [item["id"] for item in second["alerts"]]
    assert first["source"]["sha256"] == before
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_all_mode_keeps_final_and_original_totals_separate(tmp_path):
    def build(document):
        table = amount_table(document, [("Obra uno", "10 pesos"), ("Obra dos", "20 pesos"), ("Total", " pesos")])
        paragraph = table.cell(3, 1).paragraphs[0]
        deleted = OxmlElement("w:del")
        deleted_run = OxmlElement("w:r")
        deleted_text = OxmlElement("w:delText")
        deleted_text.text = "30"
        deleted_run.append(deleted_text)
        deleted.append(deleted_run)
        inserted = OxmlElement("w:ins")
        inserted_run = OxmlElement("w:r")
        inserted_text = OxmlElement("w:t")
        inserted_text.text = "40"
        inserted_run.append(inserted_text)
        inserted.append(inserted_run)
        paragraph._p.insert(0, deleted)
        paragraph._p.insert(1, inserted)
    path = save_document(tmp_path, build)
    assert alerts(review_document(path, mode="original"), "NUM-TOTAL-001") == []
    final_alert, = alerts(review_document(path, mode="final"), "NUM-TOTAL-001")
    assert final_alert["operation"]["declarado"] == Decimal("40")
    all_alerts = alerts(review_document(path, mode="all"), "NUM-TOTAL-001")
    high = [item for item in all_alerts if item["severity"] == "high"]
    assert len(high) == 1 and high[0]["operation"]["declarado"] == Decimal("40")


def test_html_is_offline_escaped_and_contains_decision_workflow(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("<script>alert(1)</script>", "10 pesos"), ("Obra", "20 pesos"), ("Total", "40 pesos")]))
    html = render_html(review_document(path))
    assert "http://" not in html and "https://" not in html
    assert "review-decisions-1.0" in html
    assert "localStorage" in html and "Importar decisiones" in html
    assert "Dónde revisar" in html and "Copiar ubicación" in html
    assert "Copiar texto para buscar" in html and "Qué debe verificar" in html
    assert "calc-table" in html and "Valor normalizado" in html
    assert r'.join("\n")' in html
    assert "<script>alert(1)</script>" not in html


def test_commented_copy_has_true_word_comments_and_preserves_original(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Obra uno", "10 pesos"), ("Obra dos", "20 pesos"), ("Total", "40 pesos")]))
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    report = review_document(path)
    target = tmp_path / "review_commented.docx"
    metadata = export_commented_docx(path, report, target)
    assert metadata["comments_added"] == 1
    assert metadata["alerts_anchored"] == 1
    assert metadata["alerts_skipped"] == []
    assert metadata["structural_verification"]["valid"] is True
    assert verify_commented_docx(target)["valid"] is True
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
    with ZipFile(target) as archive:
        comments = archive.read("word/comments.xml").decode("utf-8")
        document = archive.read("word/document.xml").decode("utf-8")
    assert "Posible hallazgo" in comments and report["alerts"][0]["id"] in comments
    assert "commentRangeStart" in document and "commentReference" in document


def test_commented_copy_groups_alerts_anchored_to_same_paragraph(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Obra uno", "10 pesos"), ("Obra dos", "20 pesos"), ("Total", "40 pesos")]))
    report = review_document(path)
    duplicate = dict(report["alerts"][0])
    duplicate["id"] = duplicate["id"] + "-second"
    report["alerts"].append(duplicate)
    metadata = export_commented_docx(path, report, tmp_path / "grouped.docx")
    assert metadata["comments_added"] == 1
    assert metadata["alerts_anchored"] == 2


def test_cli_review_outputs_serializable_json_and_html(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Obra uno", "10 pesos"), ("Obra dos", "20 pesos"), ("Total", "40 pesos")]))
    output, html, commented = tmp_path / "review.json", tmp_path / "review.html", tmp_path / "commented.docx"
    assert main([str(path), "--review", "-o", str(output), "--html", str(html),
                 "--commented-docx", str(commented)]) == 0
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["schema_version"] == "review-1.0"
    assert payload["source"]["unchanged"] is True
    assert payload["commented_copy"]["comments_added"] == 1
    assert html.read_text(encoding="utf-8").startswith("<!doctype html>")
    assert commented.is_file()


def test_custom_catalog_and_invalid_cli_combinations(tmp_path):
    path = save_document(tmp_path, lambda d: amount_table(d, [
        ("Obra", "10 pesos"), ("Gran acumulado", "20 pesos")]))
    catalog = tmp_path / "rules.json"
    catalog.write_text(json.dumps({"version": "test", "total_labels": ["^Gran acumulado$"]}), encoding="utf-8")
    report = review_document(path, catalog_path=catalog)
    assert alerts(report, "NUM-TOTAL-001")
    assert report["engine"]["catalog_version"] == "test"
    assert main([str(path), "--numbers", "--review", "-o", str(tmp_path / "bad.json")]) == 1
