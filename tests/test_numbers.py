from decimal import Decimal

import pytest
from docx import Document
from docx.oxml import OxmlElement

from revisor_crf import read_docx, export_json
from revisor_crf.numbers import detect_numbers, extract_numbers, decimal_value


@pytest.mark.parametrize("text,value,kind,multiplier", [
    ("Importe: 1,234.56 pesos", "1234.56", "importe", "1"),
    ("1.25 miles de pesos", "1.25", "importe", "1000"),
    ("2 millones de pesos", "2", "importe", "1000000"),
    ("2.30 mdp", "2.30", "importe", "1000000"),
    ("12.5%", "12.5", "porcentaje", "1"),
    ("7 por ciento", "7", "porcentaje", "1"),
    ("artículo 33", "33", "artículo", "1"),
    ("Ejercicio fiscal 2024", "2024", "año", "1"),
    ("Se realizaron 15 obras", "15", "cantidad", "1"),
    ("MXN 0.10", "0.10", "importe", "1"),
])
def test_classification_and_decimal(text, value, kind, multiplier):
    numbers = detect_numbers(text)
    assert len(numbers) == 1
    n = numbers[0]
    assert n["tipo"] == kind
    assert n["valor_decimal"] == Decimal(value)
    assert isinstance(n["valor_decimal"], Decimal)
    assert n["multiplicador"] == Decimal(multiplier)
    assert n["valor_normalizado"] == Decimal(value) * Decimal(multiplier)


@pytest.mark.parametrize("text,iso", [("31/12/2024", "2024-12-31"), ("2024-12-31", "2024-12-31"),
                                      ("3 de enero de 2025", "2025-01-03")])
def test_dates_are_one_nonmonetary_token(text, iso):
    numbers = detect_numbers(text)
    assert len(numbers) == 1
    assert numbers[0]["tipo"] == "fecha"
    assert numbers[0]["fecha_iso"] == iso
    assert numbers[0]["valor_decimal"] is None


@pytest.mark.parametrize("text,raw,subtype", [("Folio AB-001/2024", "AB-001/2024", "folio"),
                                           ("Auditoría núm. 001346", "001346", "numero_auditoria"),
                                           ("Folio AB123", "AB123", "folio")])
def test_identifiers_preserve_zeros(text, raw, subtype):
    n, = detect_numbers(text)
    assert n["tipo"] == ("numero_de_auditoria" if subtype == "numero_auditoria" else "folio")
    assert n["subtipo"] == subtype
    assert n["valor_original"] == raw
    assert n["valor_normalizado"] is None


@pytest.mark.parametrize("raw,expected", [("1,234.50", "1234.50"), ("1.234,50", "1234.50"),
                                          ("12,50", "12.50"), ("1\u00a0234.50", "1234.50"),
                                          ("−2.50", "-2.50")])
def test_numeric_notations(raw, expected):
    assert decimal_value(raw)[0] == Decimal(expected)


def test_footnote_marker_is_not_appended_to_money_decimal():
    n, = detect_numbers("28,530.722/", {
        "encabezado_columna": "Diferencias (pesos)",
        "column_header_unambiguous": True,
    })
    assert n["texto_original"] == "28,530.72"
    assert n["valor_decimal"] == Decimal("28530.72")


def test_ambiguous_and_unclassified():
    n, = detect_numbers("123")
    assert n["tipo"] == "cantidad" and n["ambigua"]
    n, = detect_numbers("12-14")
    assert n["tipo"] == "otro" and n["no_clasificada"] and n["ambigua"]
    n, = detect_numbers("31/02/2024")
    assert n["tipo"] == "fecha" and n["ambigua"]
    assert decimal_value("12,34,56")[0] is None


def test_exact_large_money_and_json(tmp_path):
    raw = "123456789012345678901234567890.12"
    n, = detect_numbers(raw + " millones de pesos")
    assert str(n["valor_normalizado"]) == "123456789012345678901234567890120000.00"
    output = tmp_path / "numbers.json"
    export_json(n, output)
    import json
    assert json.loads(output.read_text(encoding="utf-8"))["valor_decimal"] == raw


def change(p, tag, text):
    wrapper = OxmlElement("w:" + tag)
    run = OxmlElement("w:r")
    value = OxmlElement("w:delText" if tag == "del" else "w:t")
    value.text = text
    run.append(value)
    wrapper.append(run)
    p._p.append(wrapper)


def test_replacement_does_not_invent_contradictory_number(tmp_path):
    d = Document()
    p = d.add_paragraph("Importe: ")
    change(p, "del", "100")
    change(p, "ins", "200")
    p.add_run(" pesos")
    path = tmp_path / "revision.docx"
    d.save(path)
    final = read_docx(path)
    original = read_docx(path, mode="original")
    all_result = read_docx(path, mode="all")
    assert final["blocks"][0]["text"] == "Importe: 200 pesos"
    assert original["blocks"][0]["text"] == "Importe: 100 pesos"
    assert final["blocks"][0]["text_all"] == "Importe: 100200 pesos"
    assert [n["valor_decimal"] for n in extract_numbers(final)["cifras"]] == [Decimal("200")]
    assert [n["valor_decimal"] for n in extract_numbers(original)["cifras"]] == [Decimal("100")]
    numbers = extract_numbers(all_result)["cifras"]
    assert {(n["valor_decimal"], n["revision_status"]) for n in numbers} == {(Decimal("200"), "inserted"), (Decimal("100"), "deleted")}
    assert len(numbers) == 2
    assert all(len(r["blocks"]) == 1 for r in (final, original, all_result))


def test_normal_split_runs_deduplicated_in_all(tmp_path):
    d = Document()
    p = d.add_paragraph("1,")
    p.add_run("234").bold = True
    p.add_run(".50 pesos")
    path = tmp_path / "runs.docx"
    d.save(path)
    n, = extract_numbers(read_docx(path, mode="all"))["cifras"]
    assert n["valor_decimal"] == Decimal("1234.50")
    assert n["vistas"] == ["final", "original"]


def test_table_context_and_nested_coordinates(tmp_path):
    d = Document()
    d.add_paragraph("Cifras en miles de pesos")
    table = d.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Concepto", "Importe", "Año")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("Obra sintética", "12.50", "2024")):
        cell.text = text
    nested = table.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "7 personas"
    path = tmp_path / "table.docx"
    d.save(path)
    numbers = extract_numbers(read_docx(path))["cifras"]
    money = next(n for n in numbers if n["valor_original"] == "12.50")
    assert money["valor_normalizado"] == Decimal("12500.00")
    assert money["encabezado_columna"] == "Importe"
    assert money["encabezado_fila"] == "Obra sintética"
    assert (money["fila"], money["columna"]) == (2, 2)
    year = next(n for n in numbers if n["valor_original"] == "2024")
    assert year["tipo"] == "año" and year["multiplicador"] == Decimal(1)
    inner = next(n for n in numbers if n["valor_original"] == "7")
    assert inner["tabla"] != money["tabla"]


def test_moved_text_and_excluded_deleted_row(tmp_path):
    d = Document()
    host = d.add_paragraph()
    change(host, "moveFrom", "100 pesos ")
    change(host, "moveTo", "200 pesos")
    table = d.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "900 pesos"
    props = table._tbl.tr_lst[0].get_or_add_trPr()
    props.append(OxmlElement("w:del"))
    path = tmp_path / "moves.docx"
    d.save(path)
    assert [n["valor_decimal"] for n in extract_numbers(read_docx(path))["cifras"]] == [Decimal(200)]
    original = extract_numbers(read_docx(path, mode="original"))["cifras"]
    assert [n["revision_status"] for n in original] == ["moved_from", "deleted"]
    assert [n["valor_decimal"] for n in original] == [Decimal(100), Decimal(900)]


def test_unknown_revisions_are_preserved_but_not_in_final(tmp_path):
    d = Document()
    p = d.add_paragraph()
    change(p, "ins", "50 pesos")
    ins = p._p[-1]
    outer = OxmlElement("w:del")
    p._p.append(outer)
    outer.append(ins)
    path = tmp_path / "unknown.docx"
    d.save(path)
    assert extract_numbers(read_docx(path))["cifras"] == []
    n, = extract_numbers(read_docx(path, mode="all"))["cifras"]
    assert n["revision_status"] == "unknown" and n["ambigua"]


def test_invalid_reading_mode(tmp_path):
    with pytest.raises(ValueError, match="Modo"):
        read_docx(tmp_path / "absent.docx", mode="invalid")


@pytest.mark.parametrize("text,expected", [("(1,234.50) pesos", "-1234.50"),
                                          ("$ (1,234.50)", "-1234.50"),
                                          ("Miles de pesos: 2.5", "2500.0")])
def test_accounting_sign_and_prefix_units(text, expected):
    n, = detect_numbers(text)
    assert n["valor_normalizado"] == Decimal(expected)
    assert n["tipo"] == "importe"


def test_revision_ranges_across_paragraphs(tmp_path):
    from docx.oxml.ns import qn
    d = Document()
    first = d.add_paragraph()
    start = OxmlElement("w:moveFromRangeStart")
    start.set(qn("w:id"), "8")
    first._p.append(start)
    first.add_run("10 pesos")
    last = d.add_paragraph("20 pesos")
    end = OxmlElement("w:moveFromRangeEnd")
    end.set(qn("w:id"), "8")
    last._p.append(end)
    last.add_run(" y 30 pesos")
    path = tmp_path / "ranges.docx"
    d.save(path)
    final = extract_numbers(read_docx(path))["cifras"]
    assert [n["valor_decimal"] for n in final] == [Decimal(30)]
    original = extract_numbers(read_docx(path, mode="original"))["cifras"]
    assert [n["revision_status"] for n in original] == ["moved_from", "moved_from", "normal"]


def test_deleted_heading_does_not_contaminate_final_context(tmp_path):
    d = Document()
    title = d.add_heading("Título eliminado", level=1)
    wrapper = OxmlElement("w:del")
    d.element.body.insert(0, wrapper)
    wrapper.append(title._p)
    d.add_paragraph("100 pesos")
    path = tmp_path / "heading.docx"
    d.save(path)
    result = read_docx(path)
    assert result["blocks"][0]["text"] == ""
    assert result["blocks"][0]["text_all"] == "Título eliminado"
    assert result["blocks"][1]["headings"] == []


def test_header_unit_revisions_are_projected_separately(tmp_path):
    d = Document()
    t = d.add_table(rows=2, cols=1)
    header = t.cell(0, 0).paragraphs[0]
    change(header, "del", "Importe en pesos")
    change(header, "ins", "Importe en miles de pesos")
    t.cell(1, 0).text = "2"
    path = tmp_path / "headers.docx"
    d.save(path)
    n, = extract_numbers(read_docx(path))["cifras"]
    assert n["valor_normalizado"] == Decimal(2000)
    n, = extract_numbers(read_docx(path, mode="original"))["cifras"]
    assert n["valor_normalizado"] == Decimal(2)
    values = extract_numbers(read_docx(path, mode="all"))["cifras"]
    assert {n["valor_normalizado"] for n in values} == {Decimal(2), Decimal(2000)}


def test_cli_numeric_export_is_bounded(tmp_path):
    from revisor_crf.cli import main
    import json
    d = Document()
    d.add_paragraph("Contexto sintético extenso. " * 30 + "100 pesos" + " Resto sintético." * 30)
    path = tmp_path / "cli.docx"
    output = tmp_path / "cli.json"
    d.save(path)
    assert main([str(path), "--numbers", "--mode", "final", "-o", str(output)]) == 0
    result = json.loads(output.read_text(encoding="utf-8"))
    assert "blocks" not in result
    assert len(result["cifras"][0]["contexto_inmediato"]) <= 150
