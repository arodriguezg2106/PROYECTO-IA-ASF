import hashlib

from docx import Document
from openpyxl import load_workbook

from build_coherence_validation_handoff import build_package, extract_observations


def _document(path, observed=True):
    document=Document()
    document.add_paragraph(f"NÚM. DEL RESULTADO: 1 CON OBSERVACIÓN SI  ({'X' if observed else ' '}) NO ({' ' if observed else 'X'})")
    document.add_paragraph("DESCRIPCIÓN DEL RESULTADO:")
    document.add_paragraph("La entidad omitió presentar la documentación comprobatoria solicitada por 100 pesos.")
    document.add_paragraph("Por lo anterior, incumplió la disposición aplicable.")
    document.add_paragraph("NÚM. DEL RESULTADO: 2 CON OBSERVACIÓN SI  ( ) NO (X)")
    document.add_paragraph("DESCRIPCIÓN DEL RESULTADO:")
    document.add_paragraph("El resultado no tuvo observaciones.")
    document.save(path)


def test_coherence_package_extracts_only_observed_results_and_preserves_sources(tmp_path):
    source=tmp_path/"source"; source.mkdir(); path=source/"crf.docx"; _document(path)
    before=hashlib.sha256(path.read_bytes()).hexdigest()
    observations=extract_observations(path)
    assert len(observations)==1 and observations[0]["result_number"]=="1"
    assert "documentación comprobatoria" in observations[0]["text"]
    target=tmp_path/"package"; manifest=build_package(source,target,40)
    assert manifest["cases"]==1 and manifest["contains_original_docx"] is False
    assert hashlib.sha256(path.read_bytes()).hexdigest()==before
    html=(target/"casos"/"caso-01.html").read_text(encoding="utf-8")
    assert "¿Se entiende claramente" in html and "documentación comprobatoria" in html
    workbook=load_workbook(target/"validacion_coherencia.xlsx")
    sheet=workbook["Casos"]
    assert sheet.max_column==16 and len(sheet.data_validations.dataValidation)==6
    assert sheet["B2"].hyperlink.target=="casos/caso-01.html"
