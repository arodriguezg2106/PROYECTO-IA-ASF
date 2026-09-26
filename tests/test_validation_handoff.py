import json

from openpyxl import load_workbook

from build_editorial_validation_handoff import build_package


def test_handoff_shows_only_exact_case_and_simple_decision(tmp_path):
    pilot=tmp_path/"pilot"; document=pilot/"documents"/"doc"; document.mkdir(parents=True)
    (document/"review.html").write_text("REPORTE COMPLETO QUE NO DEBE COPIARSE",encoding="utf-8")
    base={"id":"ED-repetido","rule_id":"ASF-ED-I-05-YO","category":"wording",
          "severity":"warning","detection_method":"deterministic","evaluation_status":"confirmed",
          "location":{"paragraph_id":"body/b0001","section":"Resultado"},
          "original_text":"Se contrataron bienes y/o servicios.","problem_span":"y/o",
          "message":"Debe elegirse ‘y’ u ‘o’ de acuerdo con el sentido.",
          "source_section":"I.5.5","source_page":31,"effective_value":None}
    other={**base,"message":"Otro hallazgo con el mismo identificador."}
    (document/"review.json").write_text(json.dumps({"editorial_review":{"findings":[other,base]}},ensure_ascii=False),encoding="utf-8")
    sample={"finding_id":base["id"],"rule_id":base["rule_id"],"severity":"warning",
            "evaluation_status":"confirmed","detection_method":"deterministic",
            "paragraph_id":"body/b0001","message":base["message"],"file":"ejemplo.docx",
            "output_key":"doc","report":"documents/doc/review.html"}
    (pilot/"summary.json").write_text(json.dumps({"validation_sample":[sample]},ensure_ascii=False),encoding="utf-8")
    target=tmp_path/"package"; manifest=build_package(pilot,target)
    case=(target/"casos"/"caso-01.html").read_text(encoding="utf-8")
    assert "<mark>y/o</mark>" in case and base["message"] in case
    assert "Otro hallazgo" not in case and "REPORTE COMPLETO" not in case
    workbook=load_workbook(target/"validacion_editorial.xlsx")
    sheet=workbook["Casos"]
    assert sheet.max_column==10 and sheet["B2"].hyperlink.target=="casos/caso-01.html"
    assert sheet["G1"].value=="Decisión" and len(sheet.data_validations.dataValidation)==1
    assert manifest["schema_version"]=="editorial-validation-handoff-2.0"
