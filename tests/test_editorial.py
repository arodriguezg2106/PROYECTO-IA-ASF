import hashlib
import json
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from revisor_crf.editorial import (
    EDITORIAL_PROMPT_SHA256,
    EDITORIAL_PROMPT_VERSION,
    SYSTEM_PROMPT,
    EditorialConfig,
    _ignore_spelling_match,
    _parse_response,
    protected_elements,
    verify_protected,
)
from revisor_crf.editorial_models import authorize_model, load_model_registry, evaluation_passes
from revisor_crf.editorial_context_bank import load_context_bank
from revisor_crf.editorial import retrieve_rules
from revisor_crf.editorial_format import inspect_format
from revisor_crf.visual import NullVisualVerifier
from revisor_crf.editorial_report import render_markdown
from revisor_crf.html_report import render_html
from revisor_crf.review import review_document
from revisor_crf.cli import main as cli_main
from revisor_crf.commented_docx import export_commented_docx
from revisor_crf.language_tool import CachedLanguageToolProvider, LanguageToolProvider
from revisor_crf.ollama_research import (
    EVALUATION_SHA256 as RESEARCH_EVALUATION_SHA256,
    PROMPT_SHA256 as RESEARCH_PROMPT_SHA256,
)
from revisor_crf.sarif_report import render_sarif
from revisor_crf.editorial_learning import apply_priorities, build_dataset, learning_records, load_json, redact_text

def make_doc(tmp_path, text, *, style=True):
    doc=Document()
    if style:
        normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(11); normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph(text); path=tmp_path/"editorial.docx"; doc.save(path); return path

def findings(tmp_path,text):
    return review_document(make_doc(tmp_path,text),editorial=True)["editorial_review"]["findings"]

def test_percentages_separators_mixed_amount_and_fraction(tmp_path):
    rows=findings(tmp_path,"Avance 10 %; faltante $3'131,988.77; se hicieron 5 mil 277 y 1/10.")
    ids={x["rule_id"] for x in rows}
    assert {"ASF-ED-II-07-10","ASF-ED-II-07-05","ASF-ED-II-07-06","ASF-ED-II-07-08"} <= ids

def test_acronym_plural_units_and_wording(tmp_path):
    rows=findings(tmp_path,"Las UAA's debieron aperturar el registro de 10 kgs.")
    ids={x["rule_id"] for x in rows}
    assert {"ASF-ED-II-08-04C","ASF-ED-II-08-03","ASF-ED-I-06-06"} <= ids

def test_protected_elements_reject_changes():
    original="ASF revisó 10.0% el 01/02/2025, auditoría 123."
    assert protected_elements(original)
    assert verify_protected(original,original.replace("10.0%","11.0%"))["safe"] is False
    assert verify_protected(original,original.replace("revisó","analizó"))["safe"] is True

def test_ollama_json_validation_and_controlled_repair():
    row={"rule_id":"R1","category":"grammar","severity":"warning","problem_span":"x","message":"m","suggested_text":None,"confidence":.9}
    assert _parse_response("texto "+json.dumps({"findings":[row]})+" fin",{"R1"}) == [row]
    try: _parse_response("sin json",{"R1"})
    except ValueError as exc: assert str(exc)=="editorial_invalid_json"
    else: assert False

class FailingProvider:
    def review(self,request,schema): raise RuntimeError("ollama_unavailable")


class RecordingProvider:
    def __init__(self): self.requests=[]
    def review(self,request,schema):
        self.requests.append((request,schema))
        return json.dumps({"findings":[]})


class HighRecallProvider:
    def review(self, request, schema):
        text = request["fragment"]
        rule = request["rules"][0]
        return json.dumps({"findings": [{
            "rule_id": rule["rule_id"],
            "category": rule["category"],
            "severity": "suggestion",
            "problem_span": text,
            "message": "Sugerencia experimental que requiere revisión humana.",
            "suggested_text": None,
            "confidence": 0.7,
        }]})


def test_prompt_v2_is_conservative_and_hashed():
    assert EDITORIAL_PROMPT_VERSION == "editorial-context-2.1"
    assert '"findings": []' in SYSTEM_PROMPT
    assert "correct_examples" in SYSTEM_PROMPT
    assert "sugerencias meramente estilísticas" in SYSTEM_PROMPT
    assert "Un gerundio no es un error por sí mismo" in SYSTEM_PROMPT
    assert len(EDITORIAL_PROMPT_SHA256) == 64


def test_high_recall_prompt_and_evidence_are_pinned():
    assert RESEARCH_PROMPT_SHA256 == "3fa1a7b14e1dbb14493bd11e1ca1434392f748bb2c3a37c08499efdb9e44e2e1"
    assert RESEARCH_EVALUATION_SHA256 == "43662d3e6d19ab91f4eb7a2273489ea5306611f14bf5f4e97e0d930bfed8fb7a"


def test_ollama_receives_exactly_one_rule_per_request(tmp_path):
    path=make_doc(tmp_path,"La Secretaría (ASF) informó que se revisó el expediente completo.")
    config_path=tmp_path/"config.json"
    config_path.write_text(json.dumps({"editorial":{
        "ollama_enabled":True,"experimental_model":True,"model":"gemma3:12b",
        "check_format":False,
    }}),encoding="utf-8")
    provider=RecordingProvider()
    review_document(path,editorial=True,editorial_provider=provider,
                    editorial_config_path=config_path)
    assert len(provider.requests) == 2
    for request,schema in provider.requests:
        assert len(request["rules"]) == 1
        assert schema["properties"]["findings"]["items"]["properties"]["rule_id"]["enum"] == [request["rules"][0]["rule_id"]]

def test_ollama_failure_does_not_stop_review(tmp_path):
    config=EditorialConfig(ollama_enabled=True)
    result=review_document(make_doc(tmp_path,"Se utilizó y/o."),editorial=True,editorial_provider=FailingProvider(), editorial_config_path=None)
    # review_document uses its loaded config; exercise provider through a temporary config file.
    config_path=tmp_path/"config.json"; config_path.write_text(json.dumps({"editorial":{"ollama_enabled":True,"experimental_model":True}}),encoding="utf-8")
    result=review_document(tmp_path/"editorial.docx",editorial=True,editorial_provider=FailingProvider(),editorial_config_path=config_path)
    assert result["editorial_review"]["ollama"]["status"]=="partial"
    assert result["editorial_review"]["findings"]


def test_high_recall_research_mode_is_capped_and_labelled(tmp_path):
    doc = Document()
    for index in range(8):
        doc.add_paragraph(
            f"La unidad sintética número {index} comunicó que su informe contenía observaciones."
        )
    path = tmp_path / "research.docx"
    doc.save(path)
    config = EditorialConfig(
        ollama_enabled=True,
        experimental_model=True,
        model="gemma3:12b",
        check_format=False,
        confidence_threshold=0.7,
        ollama_prompt_mode="research_high_recall",
        ollama_max_suggestions=5,
    )
    report = review_document(
        path, editorial=True, editorial_provider=HighRecallProvider(),
        editorial_config_override=config,
    )
    review = report["editorial_review"]
    suggestions = [
        item for item in review["findings"] if item["detection_method"] == "ollama"
    ]
    assert len(suggestions) == 5
    assert all(item["requires_human_review"] for item in suggestions)
    assert review["ollama"]["research_mode"] is True
    assert review["ollama"]["max_suggestions"] == 5
    html = render_html(report)
    assert "Sugerencias experimentales de IA · alta cobertura" in html
    assert "puede producir falsos positivos" in html
    commented = export_commented_docx(path, report, tmp_path / "research-commented.docx")
    assert commented["comments_added"] == 0


def test_ollama_initialization_failure_is_fail_open(tmp_path, monkeypatch):
    import revisor_crf.editorial as editorial_module

    def unavailable(_config):
        raise RuntimeError("ollama_unavailable")

    monkeypatch.setattr(editorial_module, "OllamaEditorialProvider", unavailable)
    path = make_doc(tmp_path, "La unidad utilizó el registro y/o la documentación disponible.")
    config = EditorialConfig(
        ollama_enabled=True,
        experimental_model=True,
        model="gemma3:12b",
        check_format=False,
        ollama_prompt_mode="research_high_recall",
        ollama_max_suggestions=5,
    )
    report = review_document(path, editorial=True, editorial_config_override=config)
    assert report["editorial_review"]["ollama"]["status"] == "partial"
    assert any(
        item["rule_id"] == "ASF-ED-I-05-YO"
        for item in report["editorial_review"]["findings"]
    )

def test_inherited_word_style_and_reports(tmp_path):
    report=review_document(make_doc(tmp_path,"Texto correcto heredado."),editorial=True)
    assert not [x for x in report["editorial_review"]["findings"] if x["rule_id"]=="ASF-ED-III-03-02"]
    assert "Revisión editorial ASF" in render_html(report)
    assert "Revisión editorial ASF" in render_markdown(report)

def test_previous_review_contract_remains_when_editorial_disabled(tmp_path):
    report=review_document(make_doc(tmp_path,"Texto."))
    assert report["schema_version"]=="review-1.2"
    assert "editorial_review" not in report

def test_effective_format_precedence_page_margins_and_evidence(tmp_path):
    doc=Document(); base=doc.styles.add_style("BaseASF",1); base.font.name="Calibri"; base.font.size=Pt(11)
    child=doc.styles.add_style("HeredadoASF",1); child.base_style=base
    p=doc.add_paragraph("Este párrafo suficientemente largo hereda correctamente fuente y tamaño institucionales.",style=child)
    p.runs[0].font.size=Pt(12)
    section=doc.sections[0]; section.top_margin=Cm(2); section.bottom_margin=Cm(3); section.left_margin=Cm(4); section.right_margin=Cm(3); section.gutter=Cm(0); section.header_distance=Cm(.5); section.footer_distance=Cm(2)
    path=tmp_path/"format.docx"; doc.save(path)
    inspected=inspect_format(path); fmt=inspected["paragraphs"][0]["format"]
    assert fmt["font_name"]=={"value":"Calibri","source":"style:BaseASF"}
    assert fmt["font_size"]=={"value":12.0,"source":"direct"}
    report=review_document(path,editorial=True)["editorial_review"]
    finding=next(x for x in report["findings"] if x["rule_id"]=="ASF-ED-III-03-02")
    assert finding["effective_value"]["font_size"]["value"]==12.0
    assert finding["value_source"]=="mixed" and finding["structural_evidence"]

def test_doc_defaults_font_name_is_not_parsed_as_numeric_size(tmp_path):
    """Una fuente OOXML predeterminada textual no debe abortar la revisión."""
    doc=Document(); normal=doc.styles["Normal"]; normal.font.name=None; normal.font.size=None
    defaults=doc.styles.element.find(qn("w:docDefaults"))
    rpr_default=defaults.find(qn("w:rPrDefault"))
    rpr=rpr_default.find(qn("w:rPr"))
    fonts=rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts=OxmlElement("w:rFonts"); rpr.append(fonts)
    fonts.set(qn("w:ascii"),"Times New Roman")
    size=rpr.find(qn("w:sz"))
    if size is None:
        size=OxmlElement("w:sz"); rpr.append(size)
    size.set(qn("w:val"),"22")
    doc.add_paragraph("Este párrafo de prueba utiliza las propiedades predeterminadas del documento.")
    path=tmp_path/"doc-defaults.docx"; doc.save(path)

    fmt=inspect_format(path)["paragraphs"][0]["format"]
    assert fmt["font_name"]=={"value":"Times New Roman","source":"docDefaults"}
    assert fmt["font_size"]=={"value":11.0,"source":"docDefaults"}
    report=review_document(path,editorial=True)
    assert isinstance(report["editorial_review"]["findings"],list)

def test_acronym_first_mention_and_deduplicated_glossary(tmp_path):
    report=review_document(make_doc(tmp_path,"La Secretaría de Hacienda y Crédito Público (SHCP) informó a la ASF. Después, la SHCP respondió."),editorial=True)["editorial_review"]
    glossary={x["acronym"]:x for x in report["acronym_glossary"]}
    assert glossary["SHCP"]["denomination"].endswith("Crédito Público")
    assert glossary["ASF"]["status"]=="unknown"
    assert any(x["rule_id"]=="ASF-ED-II-08-04B" and x["problem_span"]=="ASF" for x in report["findings"])

def test_uppercase_labels_are_not_treated_as_acronyms(tmp_path):
    report=review_document(make_doc(tmp_path,"OPERACIONES DE INFORMACIÓN FINANCIERA DEL MUNICIPIO"),editorial=True)["editorial_review"]
    assert not [x for x in report["findings"] if x["rule_id"]=="ASF-ED-II-08-04B"]
    report=review_document(make_doc(tmp_path,"La ABC remitió el informe."),editorial=True)["editorial_review"]
    assert any(x["rule_id"]=="ASF-ED-II-08-04B" and x["problem_span"]=="ABC" for x in report["findings"])

def test_numeric_decimal_consistency_reuses_inventory(tmp_path):
    rows=findings(tmp_path,"Se revisaron 23,240.0 miles de pesos y 65 miles de pesos en el mismo concepto.")
    item=next(x for x in rows if x["rule_id"]=="ASF-ED-II-07-14")
    assert item["value_source"]=="numbers-1.3"
    assert len(item["structural_evidence"])==2

def test_table_vertical_borders_and_visual_null(tmp_path):
    doc=Document(); table=doc.add_table(rows=1,cols=1); table.cell(0,0).text="10.0"
    borders=OxmlElement("w:tblBorders"); inside=OxmlElement("w:insideV"); inside.set(qn("w:val"),"single"); borders.append(inside); table._tbl.tblPr.append(borders)
    path=tmp_path/"table.docx"; doc.save(path)
    report=review_document(path,editorial=True)["editorial_review"]
    assert any(x["rule_id"]=="ASF-ED-II-09-02" for x in report["findings"])
    assert report["visual_verification"]["status"]=="not_configured"

def test_lexical_retrieval_limits_context_rules():
    rules=[{"rule_id":"P","category":"percentages"},{"rule_id":"G","category":"grammar"},{"rule_id":"A","category":"acronyms"}]
    assert {x["rule_id"] for x in retrieve_rules("El avance fue 10.0%.",rules)}=={"P","G"}

def test_safe_context_detectors_and_counterexamples(tmp_path):
    positive=findings(tmp_path,"Se informó a la entidad debía entregar el informe; la auditoría estaría identificando deficiencias.")
    assert {"ASF-ED-I-04-QUE","ASF-ED-I-04-TENSE"}<={x["rule_id"] for x in positive}
    assert all(x["evaluation_status"]=="possible" for x in positive if x["rule_id"].startswith("ASF-ED-I-04-"))
    negative=findings(tmp_path,"Se informó a la entidad que debía entregar el informe; la auditoría identificó deficiencias.")
    assert not {x["rule_id"] for x in negative}&{"ASF-ED-I-04-QUE","ASF-ED-I-04-TENSE"}

def test_model_gate_and_criteria():
    registry=load_model_registry()
    assert authorize_model("qwen3:8b",experimental=True).status=="rejected"
    try:authorize_model("qwen3:8b")
    except ValueError as exc:assert "modo experimental" in str(exc)
    else:assert False
    metrics={"precision":.90,"recall":.85,"accuracy":.90,"prediction_agreement":.95,"invalid_responses":0,"unsafe_suggestions":0}
    assert evaluation_passes(metrics,registry["approval_criteria"],repetitions=3)
    assert not evaluation_passes(metrics,registry["approval_criteria"],repetitions=2)

def test_context_bank_has_balanced_splits():
    bank=load_context_bank("tests/data/editorial_context_cases.json")
    assert len(bank["cases"])==100
    assert bank["splits"]=={"development":40,"calibration":30,"test":30}
    assert 49<=sum(x["expected_finding"] for x in bank["cases"])<=51

def test_editorial_smoke_exports_all_formats_without_mutating_input(tmp_path):
    source=make_doc(tmp_path,"La entidad entregó 10.0% de la documentación solicitada.")
    before=hashlib.sha256(source.read_bytes()).hexdigest()
    json_path=tmp_path/"review.json"; html_path=tmp_path/"review.html"; markdown_path=tmp_path/"review.md"; sarif_path=tmp_path/"review.sarif"
    assert cli_main([str(source),"--review","--editorial-review","-o",str(json_path),"--html",str(html_path),"--markdown",str(markdown_path),"--sarif",str(sarif_path)])==0
    assert all(path.exists() and path.stat().st_size for path in (json_path,html_path,markdown_path,sarif_path))
    assert hashlib.sha256(source.read_bytes()).hexdigest()==before

class StaticLanguageTool:
    def check(self, text, *, rule_ids=()):
        if "entrego" not in text:
            return []
        start=text.index("entrego")
        return [{"offset":start,"length":7,"message":"Falta una tilde.",
                 "replacements":[{"value":"entregó"}],
                 "context":{"text":text},
                 "rule":{"id":"ACCENT_RULE","category":{"id":"GRAMMAR"}}},
                {"offset":3,"length":7,"message":"Posible error ortográfico.",
                 "replacements":[],"context":{"text":text},
                 "rule":{"id":"MORFOLOGIK_RULE_ES","category":{"id":"TYPOS"}}}]

def test_language_tool_is_experimental_local_and_reported(tmp_path):
    path=make_doc(tmp_path,"La entidad entrego el informe completo durante la revisión.")
    report=review_document(path,editorial=True,editorial_language_tool=True,
                           editorial_experimental_provider=True,
                           language_tool_provider=StaticLanguageTool())["editorial_review"]
    item=next(x for x in report["findings"] if x["detection_method"]=="language_tool")
    assert item["problem_span"]=="entrego" and item["suggested_text"]=="entregó"
    assert item["requires_human_review"] is True and item["evaluation_status"]=="possible"
    assert [x for x in report["findings"] if x.get("external_rule_id")=="MORFOLOGIK_RULE_ES"]
    assert report["language_tool"]["status"]=="complete"
    assert report["summary"]["distinct_groups"] <= report["summary"]["findings"]
    assert report["summary"]["substantive_findings"] + report["summary"]["format_occurrences"] == report["summary"]["findings"]

def test_spelling_filter_only_suppresses_known_terms_and_acronyms():
    match={"offset":0,"length":3,"rule":{"category":{"id":"TYPOS"}}}
    assert _ignore_spelling_match(match,"ASF")
    assert _ignore_spelling_match(match,"ABC",known_acronyms={"ABC"})
    assert _ignore_spelling_match({**match,"length":7},"Auditor",accepted_terms={"Auditor"})
    assert not _ignore_spelling_match({**match,"length":7},"entrego")

def test_language_tool_rejects_non_loopback_and_requires_gate(tmp_path):
    try: LanguageToolProvider("https://api.languagetool.org/v2")
    except ValueError as exc: assert "loopback" in str(exc)
    else: assert False

def test_editorial_rejects_two_linguistic_providers(tmp_path):
    path=make_doc(tmp_path,"La entidad entregó el informe completo durante la revisión.")
    config=EditorialConfig(ollama_enabled=True,language_tool_enabled=True,
                           language_tool_experimental=True)
    try: review_document(path,editorial=True,editorial_config_override=config)
    except ValueError as exc: assert "un proveedor lingüístico" in str(exc)
    else: assert False

class CountingLanguageTool:
    endpoint="http://127.0.0.1:8081/v2/check"; language="es"
    def __init__(self): self.calls=0
    def check(self,text,*,rule_ids=()):
        self.calls+=1
        return [{"offset":0,"length":4,"message":"Prueba","shortMessage":"",
                 "replacements":[{"value":"Texto"}],"context":{"text":text},
                 "rule":{"id":"R1","description":"Regla","issueType":"grammar",
                         "category":{"id":"GRAMMAR","name":"Gramática"}}}]

def test_language_tool_cache_uses_hash_and_never_stores_text(tmp_path):
    backend=CountingLanguageTool(); cache_path=tmp_path/"lt.sqlite3"
    cached=CachedLanguageToolProvider(backend,cache_path,namespace="test")
    text="Texto institucional confidencial 12345"
    assert cached.check(text)==cached.check(text)
    assert backend.calls==1 and cached.stats["hits"]==1 and cached.stats["misses"]==1
    assert text.encode() not in cache_path.read_bytes()
    assert b'"context"' not in cache_path.read_bytes()
    try: review_document(make_doc(tmp_path,"Texto."),editorial=True,
                         editorial_language_tool=True,
                         language_tool_provider=StaticLanguageTool())
    except ValueError as exc: assert "perfil aprobado" in str(exc).lower()
    else: assert False

def test_editorial_html_groups_repeated_occurrences(tmp_path):
    report=review_document(make_doc(tmp_path,"Se utilizó y/o, y después y/o."),editorial=True)
    html=render_html(report)
    assert "observaciones sustantivas" in html
    assert "Observaciones de formato agrupadas" in html
    assert "Validación prioritaria de tablas y cifras" in html
    assert "editorial-rule" in html and "2 ocurrencias" in html
    assert "review-decisions-1.2" in html and "editorial_finding" in html
    assert "editorial-decision-status" in html and "Falso positivo" in html

def test_sarif_contains_deterministic_and_editorial_findings(tmp_path):
    report=review_document(make_doc(tmp_path,"Se utilizó y/o."),editorial=True)
    sarif=render_sarif(report)
    assert sarif["version"]=="2.1.0"
    results=sarif["runs"][0]["results"]
    assert any(x["ruleId"]=="ASF-ED-I-05-YO" for x in results)
    assert sarif["runs"][0]["artifacts"][0]["hashes"]["sha-256"]==report["source"]["sha256"]

def test_learning_dataset_uses_human_labels_redaction_and_document_split(tmp_path):
    report=review_document(make_doc(tmp_path,"La ASF revisó 10.0% el 01/02/2025 y utilizó y/o."),editorial=True)
    findings=report["editorial_review"]["findings"]
    decisions={"schema_version":"review-decisions-1.2","source_sha256":report["source"]["sha256"],
               "decisions":[{"item_id":item["id"],"item_type":"editorial_finding",
                             "status":"confirmed" if index==0 else "rejected","comment":"Validado"}
                            for index,item in enumerate(findings[:2])]}
    records=learning_records(report,decisions)
    assert len(records)==2 and {x["label"] for x in records}=={"valid","false_positive"}
    assert len({x["split"] for x in records})==1
    serialized=json.dumps(records,ensure_ascii=False)
    assert "10.0" not in serialized and "01/02/2025" not in serialized
    assert "La ASF" not in serialized and "context_redacted" not in serialized
    assert records[0]["features"]["context_protected_type_counts"]
    dataset=build_dataset([(report,decisions)])
    assert dataset["contains_raw_document_text"] is False
    assert dataset["summary"]["records"]==2
    assert dataset["summary"]["training_readiness"]["ready"] is False

def test_learning_ignores_undecided_and_rejects_wrong_document(tmp_path):
    report=review_document(make_doc(tmp_path,"Se utilizó y/o."),editorial=True)
    item=report["editorial_review"]["findings"][0]
    pending={"schema_version":"review-decisions-1.2","source_sha256":report["source"]["sha256"],
             "decisions":[{"item_id":item["id"],"status":"pending","comment":""}]}
    assert learning_records(report,pending)==[]
    pending["source_sha256"]="otro"
    try: learning_records(report,pending)
    except ValueError as exc: assert "corresponden" in str(exc)
    else: assert False

class StaticPrioritizer:
    model_id="test-prioritizer"; model_version="0"
    def predict_valid_probability(self,rows): return [.75 for _ in rows]

def test_prioritizer_can_only_append_experimental_score(tmp_path):
    finding=review_document(make_doc(tmp_path,"Se utilizó y/o."),editorial=True)["editorial_review"]["findings"][0]
    scored=apply_priorities([finding],StaticPrioritizer())[0]
    assert scored["ml_priority"]["probability_valid"]==.75
    assert scored["ml_priority"]["requires_human_review"] is True
    for field in ("severity","evaluation_status","message","suggested_text"):
        assert scored[field]==finding[field]

def test_redaction_removes_protected_values():
    redacted=redact_text("La ASF revisó 25.0% el 01/02/2025.")
    assert "ASF" not in redacted and "25.0" not in redacted and "01/02/2025" not in redacted

def test_learning_json_loader_accepts_browser_and_powershell_utf8(tmp_path):
    path=tmp_path/"decisions.json"
    path.write_bytes(b"\xef\xbb\xbf"+json.dumps({"schema_version":"review-decisions-1.2"}).encode())
    assert load_json(path)["schema_version"]=="review-decisions-1.2"
