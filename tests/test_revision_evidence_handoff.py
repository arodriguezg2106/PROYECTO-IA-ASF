import hashlib

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook

from build_revision_evidence_handoff import build, extract_evidence, select_sample


def _revision(parent,kind,text):
    container=OxmlElement(f"w:{kind}"); run=OxmlElement("w:r")
    node=OxmlElement("w:delText" if kind=="del" else "w:t"); node.set(qn("xml:space"),"preserve"); node.text=text
    run.append(node); container.append(run); parent._p.append(container)


def test_extracts_anchored_comment_and_before_after_without_mutating_docx(tmp_path):
    source=tmp_path/"source"; source.mkdir(); path=source/"sample.docx"
    document=Document(); paragraph=document.add_paragraph("NÚM. DEL RESULTADO: 1 ")
    marked=paragraph.add_run("Texto ambiguo")
    document.add_comment(marked,text="Aclare quién realizó la acción.",author="Revisor",initials="R")
    changed=document.add_paragraph(); changed.add_run("La entidad "); _revision(changed,"del","realizaría"); _revision(changed,"ins","realizó"); changed.add_run(" la acción.")
    document.save(path); before=hashlib.sha256(path.read_bytes()).hexdigest()
    records=extract_evidence(path)
    comment=next(row for row in records if row["evidence_type"]=="comment")
    revision=next(row for row in records if row["evidence_type"]=="revision_only")
    assert comment["anchor_final"]=="Texto ambiguo" and comment["suggested_category"]=="Coherencia o claridad"
    assert "realizaría" in revision["before"] and "realizó" in revision["after"]
    corpus,manifest=build(source,tmp_path/"corpus",tmp_path/"package",2)
    assert corpus["summary"]["comments"]==1 and corpus["summary"]["revision_only"]==1
    assert manifest["cases"]==2 and hashlib.sha256(path.read_bytes()).hexdigest()==before
    workbook=load_workbook(tmp_path/"package"/"validacion_cambios_editoriales.xlsx")
    assert len(workbook["Casos"].data_validations.dataValidation)==5
    assert workbook["Casos"]["B2"].hyperlink.target=="casos/caso-01.html"


def test_sample_excludes_comments_without_visible_anchor_or_context():
    empty={"id":"empty","evidence_type":"comment","comment":"Revise.","before":"","after":"",
           "anchor_original":"","anchor_final":"","document_sha256":"a","has_associated_revision":False,
           "suggested_category":"Otro"}
    assert select_sample([empty],1)==[]
