import json
import zipfile

import pytest
from docx import Document

from evaluate_languagetool import evaluate
from languagetool_validation_bank import generate
from revisor_crf.editorial import EditorialConfig
from revisor_crf.commented_docx import export_commented_docx
from revisor_crf.html_report import render_html
from revisor_crf.language_tool_profile import load_language_tool_profile
from revisor_crf.review import review_document
from revisor_web.config import WebConfig


def _profile(path, *, status="approved_for_demo", rules=("R1",)):
    payload = {
        "schema_version": "languagetool-profile-1.0",
        "profile_version": "test-1",
        "status": status,
        "provider": {"name": "LanguageTool", "version": "6.6", "language": "es"},
        "approved_rule_ids": list(rules),
        "accepted_terms": ["ASF"],
        "excluded_categories": ["CASING"],
        "source": {"banks": []},
        "metrics": {"precision": 1.0},
        "privacy": {"contains_raw_text": False, "safe_for_version_control": True},
    }
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


class StaticProvider:
    def __init__(self):
        self.calls = []

    def check(self, text, *, rule_ids=()):
        self.calls.append((text, tuple(rule_ids)))
        if "entrego" not in text:
            return []
        offset = text.index("entrego")
        return [{"offset": offset, "length": 7, "message": "Falta una tilde.",
                 "replacements": [{"value": "entregó"}],
                 "rule": {"id": "R1", "category": {"id": "GRAMMAR"}}}]


def test_profile_requires_exact_schema_and_approved_allowlist(tmp_path):
    approved = load_language_tool_profile(_profile(tmp_path / "approved.json"))
    assert approved.approved and approved.approved_rule_ids == ("R1",)
    with pytest.raises(ValueError, match="perfil aprobado"):
        load_language_tool_profile(_profile(tmp_path / "candidate.json", status="candidate"))
    with pytest.raises(ValueError, match="LanguageTool 6.6"):
        path = _profile(tmp_path / "wrong.json")
        data = json.loads(path.read_text())
        data["provider"]["version"] = "otro"
        path.write_text(json.dumps(data))
        load_language_tool_profile(path)


def test_synthetic_bank_has_thirty_balanced_pairs(tmp_path):
    payload = generate(tmp_path / "bank.json")
    assert payload["summary"] == {"pairs": 30, "cases": 60}
    assert sum(item["expected_finding"] for item in payload["cases"]) == 30
    assert len({item["pair_id"] for item in payload["cases"]}) == 30
    assert payload["privacy"]["safe_for_version_control"] is True


def test_approved_profile_filters_blocks_and_uses_allowlist(tmp_path):
    profile = _profile(tmp_path / "approved.json")
    doc = Document()
    doc.add_heading("La entidad entrego información preliminar para el expediente", level=1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "La entidad entrego información en una celda de cuadro sintético."
    doc.add_paragraph("La entidad entrego el informe completo durante la revisión practicada.")
    source = tmp_path / "blocks.docx"
    doc.save(source)
    provider = StaticProvider()
    config = EditorialConfig(language_tool_enabled=True,
                             language_tool_profile_path=str(profile), check_format=False)
    full_report = review_document(source, editorial=True, editorial_config_override=config,
                                  language_tool_provider=provider)
    report = full_report["editorial_review"]
    lt = [item for item in report["findings"] if item["detection_method"] == "language_tool"]
    assert len(provider.calls) == 1 and provider.calls[0][1] == ("R1",)
    assert len(lt) == 1 and lt[0]["evaluation_status"] == "possible"
    assert report["language_tool"]["profile"]["status"] == "approved_for_demo"
    html = render_html(full_report)
    assert "Sugerencias lingüísticas de LanguageTool" in html
    assert "no generan comentarios de Word" in html
    commented = tmp_path / "commented.docx"
    result = export_commented_docx(source, full_report, commented)
    assert result["comments_added"] == 0
    with zipfile.ZipFile(commented) as archive:
        assert "word/comments.xml" not in archive.namelist()


def test_evaluator_is_stable_and_never_returns_raw_text():
    provider = StaticProvider()
    cases = [
        {"case_id": "bad", "fragment": "La entidad entrego el informe completo durante la revisión.",
         "expected_finding": True, "synthetic": True},
        {"case_id": "good", "fragment": "La entidad entregó el informe completo durante la revisión.",
         "expected_finding": False, "synthetic": True},
    ]
    result = evaluate(cases, provider, repetitions=2, rule_ids=("R1",),
                      accepted_terms=("ASF",), excluded_categories=("CASING",))
    assert result["passed"] and result["metrics"]["agreement"] == 1.0
    assert "fragment" not in json.dumps(result)


def test_web_language_tool_requires_safe_approved_configuration(tmp_path):
    profile = _profile(tmp_path / "approved.json")
    with pytest.raises(ValueError, match="revisión editorial"):
        WebConfig(jobs_dir=tmp_path / "a", language_tool_enabled=True,
                  language_tool_profile=profile, max_concurrent_jobs=1)
    with pytest.raises(ValueError, match="MAX_CONCURRENT_JOBS=1"):
        WebConfig(jobs_dir=tmp_path / "b", editorial_enabled=True,
                  language_tool_enabled=True, language_tool_profile=profile,
                  max_concurrent_jobs=2)
    with pytest.raises(ValueError, match="loopback"):
        WebConfig(jobs_dir=tmp_path / "c", editorial_enabled=True,
                  language_tool_enabled=True, language_tool_profile=profile,
                  language_tool_url="http://192.0.2.1:8081/v2", max_concurrent_jobs=1)
    config = WebConfig(jobs_dir=tmp_path / "d", editorial_enabled=True,
                       language_tool_enabled=True, language_tool_profile=profile,
                       max_concurrent_jobs=1)
    assert config.language_tool_enabled
