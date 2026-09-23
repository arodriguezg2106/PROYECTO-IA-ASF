import shutil
from decimal import Decimal

from docx import Document
from revisor_crf import read_docx
from revisor_crf.duplicates import analyze_duplicates
from revisor_crf.numbers import detect_numbers, extract_numbers


def test_duplicate_files_and_structurally_distinct_text(tmp_path):
    d = Document()
    d.add_paragraph("Texto puramente sintético")
    a, b, c = [tmp_path / name for name in ("a.docx", "b.docx", "c.docx")]
    d.save(a)
    shutil.copyfile(a, b)
    d.paragraphs[0].style = "Heading 1"
    d.save(c)
    report = analyze_duplicates([c, b, a])
    records = {r["documento_id"]: r for r in report["documentos"]}
    assert records["a.docx"]["duplicate_group_id"] == records["b.docx"]["duplicate_group_id"]
    assert records["a.docx"]["included_in_aggregates"]
    assert not records["b.docx"]["included_in_aggregates"]
    assert records["c.docx"]["included_in_aggregates"]
    assert records["a.docx"]["sha256_normalized_text"] == records["c.docx"]["sha256_normalized_text"]
    assert records["a.docx"]["sha256_element_sequence"] != records["c.docx"]["sha256_element_sequence"]
    pair = next(p for p in report["pares"] if p["documento_b"] == "b.docx")
    assert pair["similitud_secuencia_pct"] == 100 and pair["sha256_file_igual"]
    assert a.exists() and b.exists() and c.exists()


def test_duplicate_content_with_different_package_properties(tmp_path):
    d = Document()
    d.add_paragraph("Texto inventado")
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    d.core_properties.author = "Autor sintético A"
    d.save(a)
    d.core_properties.author = "Autor sintético B"
    d.save(b)
    report = analyze_duplicates([a, b])
    assert not report["pares"][0]["sha256_file_igual"]
    assert report["pares"][0]["identicos"]


def test_eligibility_by_domain():
    cases = [("100 pesos", "usable_for_arithmetic"), ("Ejercicio 2024", "usable_for_temporal_analysis"),
             ("31/12/2024", "usable_for_temporal_analysis"), ("Auditoría 00123", "usable_for_identifier_matching"),
             ("Folio AB-001", "usable_for_identifier_matching")]
    for text, flag in cases:
        n, = detect_numbers(text)
        assert n[flag]
        if flag != "usable_for_arithmetic":
            assert not n["usable_for_arithmetic"]
    for text in ("artículo 33", "fracción IV", "fracción 4"):
        n, = detect_numbers(text)
        assert not n["usable_for_arithmetic"]
        assert not n["usable_for_identifier_matching"]


def test_conflicting_explicit_and_inherited_units():
    n, = detect_numbers("2 miles de pesos", {"encabezado_columna": "Importe en pesos", "column_header_unambiguous": True})
    assert n["classification"] == "ambiguous" and n["confidence"] == "low"
    assert n["valor_normalizado"] is None
    assert not n["usable_for_arithmetic"]
    assert len(n["unit_candidates"]) == 2


def test_uma_is_not_inherited_as_pesos():
    n, = detect_numbers("25,360 UMA", {
        "table_unit_candidates": [
            {"unit": "pesos", "source": "table_title", "element_id": "title", "unambiguous": True}
        ]
    })
    assert n["classification"] == "ambiguous" and n["confidence"] == "low"
    assert n["valor_normalizado"] is None
    assert {candidate["unit"].lower() for candidate in n["unit_candidates"]} == {"uma", "pesos"}


def test_title_note_previous_paragraph_inheritance(tmp_path):
    for before, after, source, confidence in (("Tabla sintética: miles de pesos", "", "table_title", "high"),
                                            ("", "Nota: miles de pesos", "table_note", "medium"),
                                            ("Los datos siguientes se expresan en miles de pesos.", "", "previous_paragraph", "medium")):
        d = Document()
        if before:
            d.add_paragraph(before)
        table = d.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "10"
        table.cell(1, 0).text = "20"
        if after:
            d.add_paragraph(after)
        path = tmp_path / (source + ".docx")
        d.save(path)
        numbers = extract_numbers(read_docx(path))["cifras"]
        assert len(numbers) == 2
        assert all(n["unit_source"] == source and n["confidence"] == confidence for n in numbers)
        assert all(n["usable_for_arithmetic"] and n["column_arithmetic_eligible"] for n in numbers)
        assert numbers[0]["valor_normalizado"] == Decimal("10000")


def test_mixed_scales_and_reference_in_column_block_arithmetic(tmp_path):
    for name, values in (("scales", ["10 pesos", "20 miles de pesos"]), ("references", ["10 pesos", "artículo 33"]),
                         ("unitless", ["10", "20"])):
        d = Document()
        t = d.add_table(rows=2, cols=1)
        for i, value in enumerate(values):
            t.cell(i, 0).text = value
        path = tmp_path / (name + ".docx")
        d.save(path)
        numbers = extract_numbers(read_docx(path))["cifras"]
        assert len(numbers) == 2
        assert all(not n["usable_for_arithmetic"] for n in numbers)


def test_cross_reference_requires_concept_context():
    n, = detect_numbers("100 pesos")
    assert not n["usable_for_cross_reference"]
    n, = detect_numbers("100 pesos", {"encabezado_fila": "Obra sintética"})
    assert n["usable_for_cross_reference"]


def test_general_inherited_unit_is_not_a_percentage():
    n, = detect_numbers("10", {"encabezado_columna": "personas", "column_header_unambiguous": True})
    assert n["tipo"] == "cantidad"
    assert n["confidence"] == "high"
    assert n["valor_normalizado"] == Decimal(10)
    assert n["unit_source"] == "column_header"
