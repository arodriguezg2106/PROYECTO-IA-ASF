import hashlib
import json

import pytest
from openpyxl import load_workbook

from editorial_reference_workflow import (
    build_adjudication_workbook,
    build_blind_package,
    build_paired_banks,
    import_adjudication,
    import_second_review,
)


BASE_DECISION = {
    "usable": "Sí",
    "category": "Ortografía o gramática",
    "interpretation": "Sí",
    "improves": "Sí",
    "protected": "No cambiaron datos protegidos",
}


def _record(identifier, before, after):
    return {
        "id": identifier, "evidence_type": "revision_only", "before": before, "after": after,
        "document": f"documento-{identifier}.docx", "document_sha256": identifier * 64,
    }


def _sources(tmp_path):
    records = [
        _record("a", "El área entrego el informe.", "El área entregó el informe."),
        _record("b", "Texto confuso para revisar.", "Texto claro para revisar."),
        _record("c", "Comentario sin cambio.", "Comentario sin cambio."),
        _record("d", "No elegible.", "No se usa."),
    ]
    corpus = tmp_path / "corpus.json"
    corpus.write_text(json.dumps({
        "schema_version": "editorial-revision-evidence-1.0", "records": records,
    }, ensure_ascii=False), encoding="utf-8")
    cases = []
    for index, record in enumerate(records, 1):
        cases.append({
            "case": index, "evidence_id": record["id"], "human_decision": dict(BASE_DECISION),
            "eligible_for_evaluation_reference": record["id"] != "d",
        })
    validation = tmp_path / "validation.json"
    validation.write_text(json.dumps({
        "schema_version": "editorial-revision-validation-1.0", "cases": cases,
    }, ensure_ascii=False), encoding="utf-8")
    return validation, corpus


def _fill_second(path, disagreement=False):
    workbook = load_workbook(path)
    sheet = workbook["Segunda revisión"]
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 5).value = "Sí"
        sheet.cell(row, 6).value = "Ortografía o gramática"
        sheet.cell(row, 7).value = "Sí"
        sheet.cell(row, 8).value = "Sí"
        sheet.cell(row, 9).value = "No cambiaron datos protegidos"
        sheet.cell(row, 10).value = "Explicación"
        sheet.cell(row, 11).value = "Nombre que debe omitirse"
    if disagreement:
        sheet.cell(3, 6).value = "Coherencia o claridad"
    workbook.save(path)


def test_complete_double_review_adjudication_and_paired_banks(tmp_path):
    validation, corpus = _sources(tmp_path)
    blind_book, blind_manifest = tmp_path / "blind.xlsx", tmp_path / "blind-manifest.json"
    manifest = build_blind_package(validation, corpus, blind_book, blind_manifest)
    assert manifest["case_count"] == 3
    assert "documento-" not in blind_manifest.read_text(encoding="utf-8")

    _fill_second(blind_book, disagreement=True)
    second_path = tmp_path / "second.json"
    second = import_second_review(blind_book, blind_manifest, validation, corpus, second_path)
    assert second["summary"]["cases"] == 3
    assert second["privacy"]["safe_for_version_control"] is False
    assert "Nombre que debe omitirse" not in second_path.read_text(encoding="utf-8")

    adj_book, adj_manifest = tmp_path / "adj.xlsx", tmp_path / "adj-manifest.json"
    adj = build_adjudication_workbook(validation, second_path, corpus, adj_book, adj_manifest)
    assert adj["case_count"] == 1
    workbook = load_workbook(adj_book)
    sheet = workbook["Adjudicación"]
    sheet.cell(2, 6).value = "Sí"
    sheet.cell(2, 7).value = "Coherencia o claridad"
    sheet.cell(2, 8).value = "Sí"
    sheet.cell(2, 9).value = "Sí"
    sheet.cell(2, 10).value = "No cambiaron datos protegidos"
    sheet.cell(2, 11).value = "La segunda categoría describe mejor el cambio."
    sheet.cell(2, 12).value = "Nombre que debe omitirse"
    workbook.save(adj_book)
    adjudication_path = tmp_path / "adjudication.json"
    adjudication = import_adjudication(adj_book, adj_manifest, corpus, adjudication_path)
    assert adjudication["privacy"]["safe_for_version_control"] is False
    assert "Nombre que debe omitirse" not in adjudication_path.read_text(encoding="utf-8")

    real_path, synthetic_path = tmp_path / "real.json", tmp_path / "synthetic.json"
    real, synthetic = build_paired_banks(
        validation, second_path, adjudication_path, corpus, real_path, synthetic_path
    )
    assert real["schema_version"] == "editorial-paired-evaluation-bank-1.0"
    assert real["summary"] == {
        "pairs": 2, "examples": 4, "excluded": 1,
        "by_category": {"Ortografía o gramática": 1, "Coherencia o claridad": 1},
    }
    assert {item["variant"] for item in real["cases"]} == {"before", "after"}
    assert all(item["expected_finding"] == (item["variant"] == "before") for item in real["cases"])
    assert real["exclusions"][0]["reason"] == "identical_pair_text"
    assert synthetic["privacy"]["safe_for_version_control"] is True
    assert synthetic["summary"]["excluded"] == 1
    assert len(synthetic["source"]["real_bank_sha256"]) == 64
    synthetic_text = synthetic_path.read_text(encoding="utf-8")
    assert "Texto confuso" not in synthetic_text
    assert "documento-" not in synthetic_text


def test_second_review_detects_fixed_evidence_changes(tmp_path):
    validation, corpus = _sources(tmp_path)
    blind_book, blind_manifest = tmp_path / "blind.xlsx", tmp_path / "manifest.json"
    build_blind_package(validation, corpus, blind_book, blind_manifest)
    _fill_second(blind_book)
    workbook = load_workbook(blind_book)
    workbook["Segunda revisión"].cell(2, 3).value = "Evidencia alterada"
    workbook.save(blind_book)
    with pytest.raises(ValueError, match="evidencia fija"):
        import_second_review(blind_book, blind_manifest, validation, corpus, tmp_path / "second.json")


def test_second_review_accepts_justified_terminal_answer(tmp_path):
    validation, corpus = _sources(tmp_path)
    blind_book, blind_manifest = tmp_path / "blind.xlsx", tmp_path / "manifest.json"
    build_blind_package(validation, corpus, blind_book, blind_manifest)
    _fill_second(blind_book)
    workbook = load_workbook(blind_book)
    sheet = workbook["Segunda revisión"]
    sheet.cell(2, 5).value = "No puedo determinarlo"
    for column in range(6, 10):
        sheet.cell(2, column).value = None
    sheet.cell(2, 10).value = "La evidencia disponible no permite resolver el caso."
    workbook.save(blind_book)

    result = import_second_review(
        blind_book, blind_manifest, validation, corpus, tmp_path / "second.json"
    )
    assert result["summary"]["terminal_justified"] == 1
    assert result["cases"][0]["decision"] == {
        "usable": "No puedo determinarlo",
        "category": None,
        "interpretation": None,
        "improves": None,
        "protected": None,
    }


@pytest.mark.parametrize("explanation,partial", [(None, False), ("Explicación", True)])
def test_second_review_rejects_unjustified_or_partial_terminal_answer(
    tmp_path, explanation, partial
):
    validation, corpus = _sources(tmp_path)
    blind_book, blind_manifest = tmp_path / "blind.xlsx", tmp_path / "manifest.json"
    build_blind_package(validation, corpus, blind_book, blind_manifest)
    _fill_second(blind_book)
    workbook = load_workbook(blind_book)
    sheet = workbook["Segunda revisión"]
    sheet.cell(2, 5).value = "No puedo determinarlo"
    for column in range(6, 10):
        sheet.cell(2, column).value = None
    if partial:
        sheet.cell(2, 6).value = "Coherencia o claridad"
    sheet.cell(2, 10).value = explanation
    workbook.save(blind_book)

    with pytest.raises(ValueError, match="explicación|vacías"):
        import_second_review(
            blind_book, blind_manifest, validation, corpus, tmp_path / "second.json"
        )


def test_adjudication_accepts_justified_terminal_answer(tmp_path):
    validation, corpus = _sources(tmp_path)
    blind_book, blind_manifest = tmp_path / "blind.xlsx", tmp_path / "blind-manifest.json"
    build_blind_package(validation, corpus, blind_book, blind_manifest)
    _fill_second(blind_book, disagreement=True)
    second_path = tmp_path / "second.json"
    import_second_review(blind_book, blind_manifest, validation, corpus, second_path)
    adj_book, adj_manifest = tmp_path / "adj.xlsx", tmp_path / "adj-manifest.json"
    build_adjudication_workbook(validation, second_path, corpus, adj_book, adj_manifest)
    workbook = load_workbook(adj_book)
    sheet = workbook.active
    sheet.cell(2, 6).value = "No puedo determinarlo"
    for column in range(7, 11):
        sheet.cell(2, column).value = None
    sheet.cell(2, 11).value = "La evidencia no permite emitir una decisión final."
    workbook.save(adj_book)

    result = import_adjudication(
        adj_book, adj_manifest, corpus, tmp_path / "adjudication.json"
    )
    assert result["summary"] == {"cases": 1, "terminal_justified": 1}
    assert result["cases"][0]["decision"]["category"] is None


def test_bank_requires_adjudication_for_every_disagreement(tmp_path):
    validation, corpus = _sources(tmp_path)
    blind_book, blind_manifest = tmp_path / "blind.xlsx", tmp_path / "manifest.json"
    build_blind_package(validation, corpus, blind_book, blind_manifest)
    _fill_second(blind_book, disagreement=True)
    second_path = tmp_path / "second.json"
    import_second_review(blind_book, blind_manifest, validation, corpus, second_path)
    empty_adjudication = tmp_path / "empty-adjudication.json"
    empty_adjudication.write_text(json.dumps({
        "schema_version": "editorial-adjudication-1.0",
        "source": {
            "validation_sha256": hashlib.sha256(validation.read_bytes()).hexdigest(),
            "corpus_sha256": hashlib.sha256(corpus.read_bytes()).hexdigest(),
            "second_review_sha256": hashlib.sha256(second_path.read_bytes()).hexdigest(),
        },
        "cases": [],
    }), encoding="utf-8")
    with pytest.raises(ValueError, match="Falta adjudicar"):
        build_paired_banks(
            validation, second_path, empty_adjudication, corpus,
            tmp_path / "real.json", tmp_path / "synthetic.json",
        )
