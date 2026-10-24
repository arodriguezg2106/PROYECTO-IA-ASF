"""Documento exclusivamente sintético para demostración y pruebas."""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


DEMO_CASES = {
    "01_suma_incorrecta.docx": "Diferencia matemática comprobable",
    "02_suma_correcta_compleja.docx": "Tabla compleja correcta sin falso positivo",
    "03_operacion_no_evaluable.docx": "Limitación neutral por unidades incompatibles",
    "04_revision_editorial.docx": "Reglas editoriales ASF determinísticas",
}


def _configure_demo_document(document, title):
    """Aplica un formato estable y exclusivamente sintético a un caso de demo."""
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(3)
    section.left_margin = Cm(4)
    section.right_margin = Cm(3)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(3)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    section.header.paragraphs[0].text = "DOCUMENTO SINTÉTICO · DEMOSTRACIÓN"
    section.footer.paragraphs[0].text = "Sin información institucional"
    document.add_heading(title, level=1)
    document.add_paragraph(
        "Caso completamente inventado para demostrar y evaluar Revisor-CRF."
    )


def _amount_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Concepto"
    table.cell(0, 1).text = "Importe"
    for label, amount in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = amount
    return table


def _complex_amount_table(document):
    table = document.add_table(rows=2, cols=7)
    table.cell(0, 0).text = ""
    table.cell(0, 1).merge(table.cell(0, 6)).text = (
        "MUNICIPIO SINTÉTICO · AUDITORÍA DE RECURSOS INVENTADOS"
    )
    headers = (
        "Núm.", "Recurso", "Contrato", "Tipo de adjudicación", "Obra",
        "Monto contratado", "Monto ejercido",
    )
    for cell, text in zip(table.rows[1].cells, headers):
        cell.text = text
    rows = (
        ("1", "Fondo A", "C-1", "Invitación", "Obra A", "5,849,803.70", "5,772,399.59"),
        ("2", "Fondo A", "C-2", "Invitación", "Obra B", "2,862,601.63", "2,824,945.85"),
        ("3", "Fondo B", "C-3", "Invitación", "Obra C", "2,038,308.53", "2,038,308.53"),
        ("", "", "Total", "", "", "10,750,713.86", "10,635,653.97"),
    )
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    return table


def create_demo_cases(output_dir):
    """Genera el paquete de cuatro casos prácticos para Dirección."""
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)

    incorrect = Document()
    _configure_demo_document(incorrect, "Caso 1. Suma incorrecta")
    _amount_table(incorrect, (
        ("Rubro A", "100 pesos"),
        ("Rubro B", "50 pesos"),
        ("Total", "200 pesos"),
    ))
    incorrect.save(destination / "01_suma_incorrecta.docx")

    correct = Document()
    _configure_demo_document(correct, "Caso 2. Tabla compleja con totales correctos")
    _complex_amount_table(correct)
    correct.save(destination / "02_suma_correcta_compleja.docx")

    neutral = Document()
    _configure_demo_document(neutral, "Caso 3. Operación que no debe suponerse")
    _amount_table(neutral, (
        ("Obra norte", "100 pesos"),
        ("Contratos revisados", "2 obras"),
        ("Total", "100 pesos"),
    ))
    neutral.save(destination / "03_operacion_no_evaluable.docx")

    editorial = Document()
    _configure_demo_document(editorial, "Caso 4. Revisión editorial ASF")
    editorial.add_paragraph(
        "La unidad revisó el registro y/o la documentación. El avance fue de 10 % "
        "y se decidió aperturar el expediente conforme a la normatividad."
    )
    editorial.save(destination / "04_revision_editorial.docx")

    return {name: destination / name for name in DEMO_CASES}


def create_sample(path, *, with_discrepancy=False):
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.sections[0].header.paragraphs[0].text = "ENCABEZADO SINTÉTICO"
    document.sections[0].footer.paragraphs[0].text = "PIE SINTÉTICO"
    document.add_heading("1. Revisión sintética", level=1)
    document.add_paragraph("Texto de prueba: áéíóú, ñ y 漢字.\tDato\nSegunda línea.")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Concepto"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "Dato inventado"
    table.cell(1, 1).text = "100"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Detalle anidado"
    document.add_paragraph("Después de la tabla.")
    if with_discrepancy:
        document.add_heading("3. Cuadro demostrativo", level=1)
        discrepancy = document.add_table(rows=4, cols=2)
        discrepancy.style = "Table Grid"
        for row_index, values in enumerate((
            ("Concepto", "Importe"),
            ("Rubro A", "100.0"),
            ("Rubro B", "50.0"),
            ("Total", "200.0"),
        )):
            for column_index, value in enumerate(values):
                discrepancy.cell(row_index, column_index).text = value
    merged = document.add_table(rows=3, cols=2)
    merged.cell(0, 0).merge(merged.cell(0, 1)).text = "Combinación horizontal"
    merged.cell(1, 0).merge(merged.cell(2, 0)).text = "Combinación vertical"
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("2. Otra sección", level=1)
    document.add_paragraph("Fin del documento sintético.")
    document.save(target)
    return target


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Genera un DOCX sintético, sin datos reales.")
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--demo-suite",
        action="store_true",
        help="Genera en el directorio indicado los cuatro casos prácticos para Dirección.",
    )
    parser.add_argument(
        "--with-discrepancy",
        action="store_true",
        help="Añade 100 + 50 frente a un total declarado de 200 para la demostración.",
    )
    args = parser.parse_args()
    if args.demo_suite:
        create_demo_cases(args.output)
    else:
        create_sample(args.output, with_discrepancy=args.with_discrepancy)
